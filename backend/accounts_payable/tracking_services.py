import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import PermissionDenied
from django.db.models import Q
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from rest_framework.exceptions import ValidationError

from iam.permissions import has_effective_permission
from master_data.models import Party
from organization.models import Company
from . import models

ZERO = Decimal("0")
ROUND_DECIMAL = Decimal("0.0001")


def _to_decimal(val, default=ZERO):
    if val is None:
        return default
    try:
        return Decimal(str(val))
    except Exception:
        return default


def parse_date(val, field_name="date"):
    if not val:
        return None
    if isinstance(val, date):
        return val
    if isinstance(val, datetime):
        return val.date()
    try:
        return datetime.strptime(str(val).strip(), "%Y-%m-%d").date()
    except (ValueError, TypeError):
        raise ValidationError("Ngày theo dõi không đúng định dạng.")


def get_aging_bucket(overdue_days):
    """
    Phân loại nhóm tuổi nợ không chồng lấn theo BR-AP-TRK-37..42:
    - Chưa đến hạn: overdue_days <= 0 -> NOT_DUE
    - 1-30 ngày: 1 <= overdue_days <= 30 -> 1_30
    - 31-60 ngày: 31 <= overdue_days <= 60 -> 31_60
    - 61-90 ngày: 61 <= overdue_days <= 90 -> 61_90
    - Trên 90 ngày: overdue_days > 90 -> OVER_90
    """
    if overdue_days <= 0:
        return "NOT_DUE"
    if 1 <= overdue_days <= 30:
        return "1_30"
    if 31 <= overdue_days <= 60:
        return "31_60"
    if 61 <= overdue_days <= 90:
        return "61_90"
    return "OVER_90"


def get_aging_bucket_display(bucket_code):
    mapping = {
        "NOT_DUE": "Chưa đến hạn",
        "1_30": "1–30 ngày",
        "31_60": "31–60 ngày",
        "61_90": "61–90 ngày",
        "OVER_90": "Trên 90 ngày",
    }
    return mapping.get(bucket_code, bucket_code)


def get_payable_posting_date(payable):
    """Lấy ngày hạch toán / ngày chứng từ hiệu lực của khoản công nợ."""
    if payable.document:
        return payable.document.posting_date or payable.document.document_date
    if payable.invoice:
        return payable.invoice.posting_date or payable.invoice.invoice_date
    if payable.goods_purchase:
        return payable.goods_purchase.posting_date or payable.goods_purchase.document_date
    if payable.purchase_cost:
        return payable.purchase_cost.posting_date or payable.purchase_cost.document_date
    return payable.created_at.date()


def get_payable_source_info(payable):
    """Trả về thông tin chứng từ nguồn tạo công nợ."""
    if payable.document:
        return {
            "source_type": "DOCUMENT",
            "source_type_display": "Khoản phải trả khác" if payable.document.source_type == "OTHER" else "Chứng từ mua hàng",
            "document_number": payable.document.number,
            "document_date": payable.document.document_date,
            "posting_date": payable.document.posting_date,
            "source_document_id": str(payable.document.id),
            "invoice_number": payable.document.invoice.invoice_number if payable.document.invoice else "",
            "invoice_series": payable.document.invoice.invoice_series if payable.document.invoice else "",
            "payable_account_id": str(payable.document.payable_account_id) if payable.document.payable_account_id else None,
        }
    if payable.invoice:
        return {
            "source_type": "INVOICE",
            "source_type_display": "Hóa đơn mua hàng",
            "document_number": payable.invoice.invoice_number,
            "document_date": payable.invoice.invoice_date,
            "posting_date": payable.invoice.posting_date or payable.invoice.invoice_date,
            "source_document_id": str(payable.invoice.id),
            "invoice_number": payable.invoice.invoice_number,
            "invoice_series": payable.invoice.invoice_series,
            "payable_account_id": None,
        }
    if payable.goods_purchase:
        return {
            "source_type": "GOODS_PURCHASE",
            "source_type_display": "Mua hàng hóa / dịch vụ",
            "document_number": payable.goods_purchase.number,
            "document_date": payable.goods_purchase.document_date,
            "posting_date": payable.goods_purchase.posting_date,
            "source_document_id": str(payable.goods_purchase.id),
            "invoice_number": payable.goods_purchase.supplier_invoice.invoice_number if payable.goods_purchase.supplier_invoice else "",
            "invoice_series": payable.goods_purchase.supplier_invoice.invoice_series if payable.goods_purchase.supplier_invoice else "",
            "payable_account_id": None,
        }
    if payable.purchase_cost:
        return {
            "source_type": "PURCHASE_COST",
            "source_type_display": "Chi phí mua hàng",
            "document_number": payable.purchase_cost.number,
            "document_date": payable.purchase_cost.document_date,
            "posting_date": payable.purchase_cost.posting_date,
            "source_document_id": str(payable.purchase_cost.id),
            "invoice_number": "",
            "invoice_series": "",
            "payable_account_id": None,
        }
    return {
        "source_type": "OTHER",
        "source_type_display": "Khác",
        "document_number": str(payable.id)[:8],
        "document_date": payable.created_at.date(),
        "posting_date": payable.created_at.date(),
        "source_document_id": str(payable.id),
        "invoice_number": "",
        "invoice_series": "",
        "payable_account_id": None,
    }


def is_payable_posted(payable):
    """
    BR-AP-TRK-06..08: Chỉ chứng từ đã ghi sổ mới được tính vào số liệu công nợ.
    Loại trừ Draft, Pending, Rejected, Cancelled, Reversed.
    """
    if payable.status == models.PayableStatus.CANCELLED:
        return False
    if payable.document:
        return payable.document.status == models.DocumentStatus.POSTED
    if payable.invoice:
        return payable.invoice.status == models.DocumentStatus.POSTED
    if payable.goods_purchase:
        return payable.goods_purchase.status == "POSTED"
    if payable.purchase_cost:
        return payable.purchase_cost.status == "POSTED"
    return True


def validate_tracking_filters(company, params, user=None):
    """
    Kiểm tra và chuẩn hóa các tham số lọc cho UC-AP-02.
    Xử lý các ngoại lệ EF-01, EF-02, EF-03, EF-04.
    """
    today = timezone.localdate()

    # 1. Từ ngày & Đến ngày
    from_date_raw = params.get("from_date")
    to_date_raw = params.get("to_date")

    from_date = parse_date(from_date_raw, field_name="from_date")
    to_date = parse_date(to_date_raw, field_name="to_date")

    if to_date is None:
        to_date = today
    if from_date is None:
        # Mặc định ngày đầu kỳ hiện tại
        from_date = to_date.replace(day=1)

    # EF-04: Từ ngày không được lớn hơn Đến ngày
    if from_date > to_date:
        raise ValidationError("Từ ngày không được lớn hơn Đến ngày.")

    # 2. Nhà cung cấp
    supplier_ids = []
    if hasattr(params, "getlist"):
        supplier_ids = params.getlist("supplier_id") or params.getlist("supplier_ids")
    if not supplier_ids and params.get("supplier_id"):
        supplier_ids = [params.get("supplier_id")]
    elif not supplier_ids and params.get("supplier"):
        supplier_ids = [params.get("supplier")]

    supplier_group_id = params.get("supplier_group_id") or params.get("supplier_group")
    status_filter = (params.get("status") or params.get("debt_status") or "").upper().strip() or None
    due_status_filter = (params.get("due_status") or "").upper().strip() or None
    source_type_filter = (params.get("source_type") or "").upper().strip() or None
    payable_account_id = params.get("payable_account_id") or params.get("payable_account")
    assigned_to_id = params.get("assigned_to_id") or params.get("assigned_to") or params.get("pic")
    search = (params.get("search") or "").strip() or None

    show_zero_raw = params.get("show_zero_balance")
    show_zero_balance = False
    if show_zero_raw is not None:
        show_zero_balance = str(show_zero_raw).lower() in ("true", "1", "yes")

    return {
        "company": company,
        "from_date": from_date,
        "to_date": to_date,
        "supplier_ids": [s for s in supplier_ids if s],
        "supplier_group_id": supplier_group_id,
        "status": status_filter,
        "due_status": due_status_filter,
        "source_type": source_type_filter,
        "payable_account_id": payable_account_id,
        "assigned_to_id": assigned_to_id,
        "show_zero_balance": show_zero_balance,
        "search": search,
    }


def get_supplier_debt_summary(company, filters, user=None):
    """
    Tổng hợp công nợ theo nhà cung cấp (Section 1.2, Bước 12-24).
    Tuân thủ đầy đủ BR-AP-TRK-06 đến BR-AP-TRK-24.
    """
    from_date = filters["from_date"]
    to_date = filters["to_date"]
    show_zero_balance = filters["show_zero_balance"]

    # 1. Tập hợp nhà cung cấp trong phạm vi
    suppliers_qs = Party.objects.filter(company=company)
    if filters["supplier_ids"]:
        suppliers_qs = suppliers_qs.filter(id__in=filters["supplier_ids"])
    if filters["supplier_group_id"]:
        suppliers_qs = suppliers_qs.filter(group_id=filters["supplier_group_id"])
    if filters["search"]:
        s = filters["search"]
        suppliers_qs = suppliers_qs.filter(
            Q(code__icontains=s) | Q(name__icontains=s) | Q(tax_code__icontains=s)
        )

    suppliers = list(suppliers_qs.order_by("code"))
    results = []

    grand_total_opening_payable = Decimal("0")
    grand_total_opening_advance = Decimal("0")
    grand_total_period_increase = Decimal("0")
    grand_total_period_decrease = Decimal("0")
    grand_total_closing_payable = Decimal("0")
    grand_total_closing_advance = Decimal("0")
    grand_total_due = Decimal("0")
    grand_total_overdue = Decimal("0")

    for supplier in suppliers:
        # A. Lấy toàn bộ các khoản nợ của NCC có ngày phát sinh <= to_date
        payables = models.Payable.objects.filter(
            company=company,
            supplier=supplier,
        ).exclude(status=models.PayableStatus.CANCELLED).select_related(
            "document", "invoice", "goods_purchase", "purchase_cost", "currency"
        ).prefetch_related(
            "adjustments", "allocations__payment", "nettings"
        )

        opening_payable_balance = Decimal("0")
        period_increase = Decimal("0")
        period_decrease = Decimal("0")
        closing_payable_balance = Decimal("0")
        supplier_due_amount = Decimal("0")
        supplier_overdue_amount = Decimal("0")
        earliest_due_date = None
        max_overdue_days = 0
        unpaid_count = 0
        has_matching_payables = False

        for p in payables:
            if not is_payable_posted(p):
                continue

            doc_date = get_payable_posting_date(p)
            if doc_date > to_date:
                continue

            # Lọc theo source_type nếu có
            src_info = get_payable_source_info(p)
            if filters["source_type"] and src_info["source_type"] != filters["source_type"]:
                continue
            if filters["payable_account_id"] and src_info.get("payable_account_id") != str(filters["payable_account_id"]):
                continue

            orig = p.original_amount

            # 1. Nghiệp vụ trước from_date
            adj_inc_pre = sum(
                (a.amount for a in p.adjustments.all()
                 if a.status == "POSTED" and a.amount > 0 and a.adjustment_date < from_date),
                Decimal("0")
            )
            adj_dec_pre = sum(
                (abs(a.amount) for a in p.adjustments.all()
                 if a.status == "POSTED" and a.amount < 0 and a.adjustment_date < from_date),
                Decimal("0")
            )
            paid_pre = sum(
                (al.amount for al in p.allocations.all()
                 if al.payment.status == "CONFIRMED" and al.payment.payment_date < from_date),
                Decimal("0")
            )
            net_pre = sum(
                (n.amount for n in p.nettings.all()
                 if n.status == "POSTED" and n.netting_date < from_date),
                Decimal("0")
            )

            if doc_date < from_date:
                p_opening = max(Decimal("0"), orig + adj_inc_pre - paid_pre - adj_dec_pre - net_pre)
            else:
                p_opening = Decimal("0")

            # 2. Nghiệp vụ trong kỳ [from_date, to_date]
            adj_inc_cur = sum(
                (a.amount for a in p.adjustments.all()
                 if a.status == "POSTED" and a.amount > 0 and from_date <= a.adjustment_date <= to_date),
                Decimal("0")
            )
            adj_dec_cur = sum(
                (abs(a.amount) for a in p.adjustments.all()
                 if a.status == "POSTED" and a.amount < 0 and from_date <= a.adjustment_date <= to_date),
                Decimal("0")
            )
            paid_cur = sum(
                (al.amount for al in p.allocations.all()
                 if al.payment.status == "CONFIRMED" and from_date <= al.payment.payment_date <= to_date),
                Decimal("0")
            )
            net_cur = sum(
                (n.amount for n in p.nettings.all()
                 if n.status == "POSTED" and from_date <= n.netting_date <= to_date),
                Decimal("0")
            )

            p_inc = (orig if from_date <= doc_date <= to_date else Decimal("0")) + adj_inc_cur
            p_dec = paid_cur + adj_dec_cur + net_cur
            p_closing = max(Decimal("0"), p_opening + p_inc - p_dec)

            opening_payable_balance += p_opening
            period_increase += p_inc
            period_decrease += p_dec
            closing_payable_balance += p_closing

            # 3. Phân tích quá hạn và đến hạn tại to_date (BR-25..32)
            if p_closing > Decimal("0"):
                unpaid_count += 1
                if p.due_date:
                    if earliest_due_date is None or p.due_date < earliest_due_date:
                        earliest_due_date = p.due_date

                    if p.due_date == to_date:
                        supplier_due_amount += p_closing
                    elif p.due_date < to_date:
                        supplier_overdue_amount += p_closing
                        ov_days = (to_date - p.due_date).days
                        if ov_days > max_overdue_days:
                            max_overdue_days = ov_days

            has_matching_payables = True

        # B. Tính khoản trả trước của nhà cung cấp (BR-18, 21..24)
        payments_qs = models.Payment.objects.filter(
            company=company, supplier=supplier, status="CONFIRMED"
        )
        pay_pre = sum((pm.amount for pm in payments_qs if pm.payment_date < from_date), Decimal("0"))
        alloc_pre_rec = sum(
            (al.amount for al in models.PaymentAllocation.objects.filter(
                payment__company=company, payment__supplier=supplier,
                payment__status="CONFIRMED", payment__payment_date__lt=from_date
            )),
            Decimal("0")
        )
        alloc_pre_direct = sum((pm.allocated_amount for pm in payments_qs if pm.payment_date < from_date), Decimal("0"))
        opening_advance_balance = max(Decimal("0"), pay_pre - max(alloc_pre_rec, alloc_pre_direct))

        pay_to = sum((pm.amount for pm in payments_qs if pm.payment_date <= to_date), Decimal("0"))
        alloc_to_rec = sum(
            (al.amount for al in models.PaymentAllocation.objects.filter(
                payment__company=company, payment__supplier=supplier,
                payment__status="CONFIRMED", payment__payment_date__lte=to_date
            )),
            Decimal("0")
        )
        alloc_to_direct = sum((pm.allocated_amount for pm in payments_qs if pm.payment_date <= to_date), Decimal("0"))
        closing_advance_balance = max(Decimal("0"), pay_to - max(alloc_to_rec, alloc_to_direct))

        # C. Áp dụng các điều kiện lọc bổ sung
        # 1. Ẩn số dư bằng 0 (AF-13, BR-24)
        if not show_zero_balance:
            if (closing_payable_balance == Decimal("0") and
                    closing_advance_balance == Decimal("0") and
                    period_increase == Decimal("0") and
                    period_decrease == Decimal("0")):
                continue

        # 2. Lọc theo tình trạng đến hạn (due_status)
        if filters["due_status"]:
            if filters["due_status"] == "OVERDUE" and supplier_overdue_amount <= Decimal("0"):
                continue
            if filters["due_status"] == "DUE" and supplier_due_amount <= Decimal("0"):
                continue
            if filters["due_status"] == "NOT_DUE":
                not_due_amt = closing_payable_balance - supplier_due_amount - supplier_overdue_amount
                if not_due_amt <= Decimal("0"):
                    continue

        # 3. Lọc theo trạng thái nợ (status)
        if filters["status"]:
            if filters["status"] == "OVERDUE" and supplier_overdue_amount <= Decimal("0"):
                continue
            if filters["status"] == "PAID" and closing_payable_balance > Decimal("0"):
                continue
            if filters["status"] == "UNPAID" and (unpaid_count == 0 or supplier_overdue_amount > Decimal("0")):
                continue
            if filters["status"] == "PARTIALLY_PAID" and (closing_payable_balance <= Decimal("0") or supplier_overdue_amount > Decimal("0")):
                continue

        # D. Tổng hợp số liệu
        row = {
            "supplier_id": str(supplier.id),
            "supplier_code": supplier.code,
            "supplier_name": supplier.name,
            "tax_code": supplier.tax_code or "",
            "opening_payable_balance": str(opening_payable_balance),
            "opening_advance_balance": str(opening_advance_balance),
            "period_increase": str(period_increase),
            "period_decrease": str(period_decrease),
            "closing_payable_balance": str(closing_payable_balance),
            "closing_advance_balance": str(closing_advance_balance),
            "due_amount": str(supplier_due_amount),
            "overdue_amount": str(supplier_overdue_amount),
            "earliest_due_date": str(earliest_due_date) if earliest_due_date else None,
            "max_overdue_days": max_overdue_days,
            "unpaid_count": unpaid_count,
        }
        results.append(row)

        grand_total_opening_payable += opening_payable_balance
        grand_total_opening_advance += opening_advance_balance
        grand_total_period_increase += period_increase
        grand_total_period_decrease += period_decrease
        grand_total_closing_payable += closing_payable_balance
        grand_total_closing_advance += closing_advance_balance
        grand_total_due += supplier_due_amount
        grand_total_overdue += supplier_overdue_amount

    return {
        "from_date": str(from_date),
        "to_date": str(to_date),
        "summary_metrics": {
            "total_suppliers": len(results),
            "total_opening_payable": str(grand_total_opening_payable),
            "total_opening_advance": str(grand_total_opening_advance),
            "total_period_increase": str(grand_total_period_increase),
            "total_period_decrease": str(grand_total_period_decrease),
            "total_closing_payable": str(grand_total_closing_payable),
            "total_closing_advance": str(grand_total_closing_advance),
            "total_due_amount": str(grand_total_due),
            "total_overdue_amount": str(grand_total_overdue),
        },
        "results": results,
    }


def get_supplier_debt_detail(company, supplier_id, filters, user=None):
    """
    Chi tiết các khoản công nợ của một nhà cung cấp (Section 1.3, Bước 25-37).
    Sắp xếp mặc định theo hạn thanh toán tăng dần (BR-AP-TRK-43).
    """
    from_date = filters["from_date"]
    to_date = filters["to_date"]

    try:
        supplier = Party.objects.get(pk=supplier_id, company=company)
    except (Party.DoesNotExist, ValueError):
        raise ValidationError("Nhà cung cấp không tồn tại hoặc bạn không còn quyền truy cập.")

    payables_qs = models.Payable.objects.filter(
        company=company,
        supplier=supplier,
    ).exclude(status=models.PayableStatus.CANCELLED).select_related(
        "document", "invoice", "goods_purchase", "purchase_cost", "currency"
    ).prefetch_related(
        "adjustments", "allocations__payment", "nettings"
    )

    detail_lines = []

    opening_payable_total = Decimal("0")
    period_increase_total = Decimal("0")
    period_decrease_total = Decimal("0")
    closing_payable_total = Decimal("0")

    for p in payables_qs:
        if not is_payable_posted(p):
            continue

        doc_date = get_payable_posting_date(p)
        if doc_date > to_date:
            continue

        src_info = get_payable_source_info(p)
        if filters.get("source_type") and src_info["source_type"] != filters["source_type"]:
            continue
        if filters.get("payable_account_id") and src_info.get("payable_account_id") != str(filters["payable_account_id"]):
            continue

        orig = p.original_amount

        # Tăng, giảm trước from_date
        adj_inc_pre = sum(
            (a.amount for a in p.adjustments.all()
             if a.status == "POSTED" and a.amount > 0 and a.adjustment_date < from_date),
            Decimal("0")
        )
        adj_dec_pre = sum(
            (abs(a.amount) for a in p.adjustments.all()
             if a.status == "POSTED" and a.amount < 0 and a.adjustment_date < from_date),
            Decimal("0")
        )
        paid_pre = sum(
            (al.amount for al in p.allocations.all()
             if al.payment.status == "CONFIRMED" and al.payment.payment_date < from_date),
            Decimal("0")
        )
        net_pre = sum(
            (n.amount for n in p.nettings.all()
             if n.status == "POSTED" and n.netting_date < from_date),
            Decimal("0")
        )

        p_opening = max(Decimal("0"), orig + adj_inc_pre - paid_pre - adj_dec_pre - net_pre) if doc_date < from_date else Decimal("0")

        # Tăng, giảm trong kỳ
        adj_inc_cur = sum(
            (a.amount for a in p.adjustments.all()
             if a.status == "POSTED" and a.amount > 0 and from_date <= a.adjustment_date <= to_date),
            Decimal("0")
        )
        adj_dec_cur = sum(
            (abs(a.amount) for a in p.adjustments.all()
             if a.status == "POSTED" and a.amount < 0 and from_date <= a.adjustment_date <= to_date),
            Decimal("0")
        )
        paid_cur = sum(
            (al.amount for al in p.allocations.all()
             if al.payment.status == "CONFIRMED" and from_date <= al.payment.payment_date <= to_date),
            Decimal("0")
        )
        net_cur = sum(
            (n.amount for n in p.nettings.all()
             if n.status == "POSTED" and from_date <= n.netting_date <= to_date),
            Decimal("0")
        )

        p_inc = (orig if from_date <= doc_date <= to_date else Decimal("0")) + adj_inc_cur
        p_dec = paid_cur + adj_dec_cur + net_cur
        p_closing = max(Decimal("0"), p_opening + p_inc - p_dec)

        opening_payable_total += p_opening
        period_increase_total += p_inc
        period_decrease_total += p_dec
        closing_payable_total += p_closing

        # Lũy kế tính đến to_date cho từng khoản
        total_increase_to = adj_inc_pre + adj_inc_cur
        total_paid_to = paid_pre + paid_cur
        total_discount_to = adj_dec_pre + adj_dec_cur
        total_netted_to = net_pre + net_cur
        remaining_at_to_date = p_closing

        # Đánh giá hạn thanh toán tại to_date (BR-25..32, AF-04..06)
        days_remaining = 0
        overdue_days = 0
        due_status = "NOT_DUE"

        if p.due_date:
            if p.due_date > to_date:
                due_status = "NOT_DUE"
                days_remaining = (p.due_date - to_date).days
            elif p.due_date == to_date:
                due_status = "DUE"
                days_remaining = 0
                overdue_days = 0
            else:  # p.due_date < to_date
                if remaining_at_to_date > Decimal("0"):
                    due_status = "OVERDUE"
                    overdue_days = (to_date - p.due_date).days
                else:
                    due_status = "NOT_DUE"
                    overdue_days = 0

        aging_bucket = get_aging_bucket(overdue_days)

        # Trạng thái khoản công nợ (BR-33..36, AF-07, AF-08)
        if remaining_at_to_date == Decimal("0"):
            status = "PAID"
        else:
            if p.due_date and p.due_date < to_date:
                status = "OVERDUE"  # Ưu tiên hiển thị Quá hạn theo BR-36, AF-08
            elif (total_paid_to + total_discount_to + total_netted_to) > Decimal("0"):
                status = "PARTIALLY_PAID"
            else:
                status = "UNPAID"

        # Lọc trạng thái / tình trạng đến hạn nếu có
        if filters.get("status") and status != filters["status"]:
            continue
        if filters.get("due_status") and due_status != filters["due_status"]:
            continue

        item = {
            "payable_id": str(p.id),
            "document_number": src_info["document_number"],
            "document_date": str(src_info["document_date"]),
            "posting_date": str(src_info["posting_date"]),
            "due_date": str(p.due_date) if p.due_date else None,
            "source_type": src_info["source_type"],
            "source_type_display": src_info["source_type_display"],
            "source_document_id": src_info["source_document_id"],
            "invoice_number": src_info["invoice_number"],
            "invoice_series": src_info["invoice_series"],
            "currency_code": p.currency.code,
            "original_amount": str(orig),
            "increase_amount": str(total_increase_to),
            "paid_amount": str(total_paid_to),
            "discount_amount": str(total_discount_to),
            "netted_amount": str(total_netted_to),
            "remaining_amount": str(remaining_at_to_date),
            "days_remaining": days_remaining,
            "overdue_days": overdue_days,
            "due_status": due_status,
            "due_status_display": "Quá hạn" if due_status == "OVERDUE" else ("Đến hạn" if due_status == "DUE" else "Chưa đến hạn"),
            "aging_bucket": aging_bucket,
            "aging_bucket_display": get_aging_bucket_display(aging_bucket),
            "status": status,
            "status_display": (
                "Quá hạn" if status == "OVERDUE" else
                ("Đã thanh toán" if status == "PAID" else
                 ("Thanh toán một phần" if status == "PARTIALLY_PAID" else "Chưa thanh toán"))
            ),
            "is_overdue": overdue_days > 0 and remaining_at_to_date > Decimal("0"),
        }
        detail_lines.append(item)

    # BR-AP-TRK-43: Sắp xếp theo hạn thanh toán tăng dần (None xếp sau cùng)
    detail_lines.sort(key=lambda x: (x["due_date"] is None, x["due_date"] or "9999-99-99"))

    # Tính khoản trả trước cho NCC
    payments_qs = models.Payment.objects.filter(
        company=company, supplier=supplier, status="CONFIRMED"
    )
    pay_pre = sum((pm.amount for pm in payments_qs if pm.payment_date < from_date), Decimal("0"))
    alloc_pre_rec = sum(
        (al.amount for al in models.PaymentAllocation.objects.filter(
            payment__company=company, payment__supplier=supplier,
            payment__status="CONFIRMED", payment__payment_date__lt=from_date
        )),
        Decimal("0")
    )
    alloc_pre_direct = sum((pm.allocated_amount for pm in payments_qs if pm.payment_date < from_date), Decimal("0"))
    opening_advance_balance = max(Decimal("0"), pay_pre - max(alloc_pre_rec, alloc_pre_direct))

    pay_to = sum((pm.amount for pm in payments_qs if pm.payment_date <= to_date), Decimal("0"))
    alloc_to_rec = sum(
        (al.amount for al in models.PaymentAllocation.objects.filter(
            payment__company=company, payment__supplier=supplier,
            payment__status="CONFIRMED", payment__payment_date__lte=to_date
        )),
        Decimal("0")
    )
    alloc_to_direct = sum((pm.allocated_amount for pm in payments_qs if pm.payment_date <= to_date), Decimal("0"))
    closing_advance_balance = max(Decimal("0"), pay_to - max(alloc_to_rec, alloc_to_direct))

    return {
        "from_date": str(from_date),
        "to_date": str(to_date),
        "supplier": {
            "id": str(supplier.id),
            "code": supplier.code,
            "name": supplier.name,
            "tax_code": supplier.tax_code or "",
            "opening_payable_balance": str(opening_payable_total),
            "opening_advance_balance": str(opening_advance_balance),
            "period_increase": str(period_increase_total),
            "period_decrease": str(period_decrease_total),
            "closing_payable_balance": str(closing_payable_total),
            "closing_advance_balance": str(closing_advance_balance),
        },
        "results": detail_lines,
    }


def verify_document_view_permission(user, company, document_type, document_id):
    """
    BR-AP-TRK-03, AF-10: Kiểm tra quyền xem chi tiết từng loại chứng từ liên quan.
    """
    if not user or not user.is_authenticated:
        return False
    if getattr(user, "is_superuser", False):
        return True
    if not hasattr(user, "iam_account"):
        return False

    perm_map = {
        "AP.DOCUMENT": "AP.DOCUMENT.VIEW",
        "AP.INVOICE": "AP.INVOICE.VIEW",
        "AP.PAYMENT": "AP.PAYMENT.VIEW",
        "AP.NETTING": "AP.NETTING.VIEW",
        "AP.ADJUSTMENT": "AP.ADJUSTMENT.VIEW",
        "PURCHASE.GOODS": "PURCHASE.GOODS.VIEW",
        "PURCHASE.COST": "PURCHASE.COST.VIEW",
    }
    perm_code = perm_map.get(document_type, "AP.DOCUMENT.VIEW")
    return has_effective_permission(user.iam_account, company, perm_code)


def get_payable_timeline(company, payable_id, user=None):
    """
    Lấy lịch sử phát sinh của khoản công nợ (Section 1.4, Bước 38-40).
    Xử lý ngoại lệ EF-08 (khoản nợ bị hủy).
    """
    try:
        p = models.Payable.objects.select_related(
            "document", "invoice", "goods_purchase", "purchase_cost", "supplier", "currency"
        ).get(pk=payable_id, company=company)
    except (models.Payable.DoesNotExist, ValueError):
        raise ValidationError("Khoản công nợ không tồn tại.")

    # EF-08: Khoản công nợ đã bị hủy
    if p.status == models.PayableStatus.CANCELLED:
        raise ValidationError("Khoản công nợ không còn hiệu lực.")

    src_info = get_payable_source_info(p)
    events = []

    # 1. Sự kiện ghi nhận công nợ ban đầu
    doc_type_code = f"AP.{src_info['source_type']}" if src_info['source_type'] in ("DOCUMENT", "INVOICE") else "AP.DOCUMENT"
    can_view_src = verify_document_view_permission(user, company, doc_type_code, src_info["source_document_id"])

    events.append({
        "event_type": "RECOGNITION",
        "event_type_display": "Ghi nhận công nợ ban đầu",
        "event_date": str(src_info["posting_date"]),
        "document_type": doc_type_code,
        "document_number": src_info["document_number"],
        "document_id": src_info["source_document_id"],
        "amount": str(p.original_amount),
        "description": f"Ghi nhận nợ phải trả theo chứng từ {src_info['document_number']}",
        "can_view_document": can_view_src,
    })

    # 2. Các lần thanh toán (Payment Allocations)
    for al in p.allocations.filter(payment__status="CONFIRMED").select_related("payment").order_by("payment__payment_date"):
        can_view_pay = verify_document_view_permission(user, company, "AP.PAYMENT", str(al.payment.id))
        events.append({
            "event_type": "PAYMENT",
            "event_type_display": "Thanh toán công nợ",
            "event_date": str(al.payment.payment_date),
            "document_type": "AP.PAYMENT",
            "document_number": al.payment.external_reference or str(al.payment.id)[:8],
            "document_id": str(al.payment.id),
            "amount": str(-al.amount),
            "description": f"Thanh toán đợt theo giao dịch {al.payment.external_reference}",
            "can_view_document": can_view_pay,
        })

    # 3. Các lần điều chỉnh / Giảm trừ / Hàng trả lại
    for adj in p.adjustments.filter(status="POSTED").order_by("adjustment_date"):
        can_view_adj = verify_document_view_permission(user, company, "AP.ADJUSTMENT", str(adj.id))
        events.append({
            "event_type": "ADJUSTMENT",
            "event_type_display": "Điều chỉnh tăng công nợ" if adj.amount > 0 else "Giảm trừ / Trả lại hàng",
            "event_date": str(adj.adjustment_date),
            "document_type": "AP.ADJUSTMENT",
            "document_number": adj.posting_reference or str(adj.id)[:8],
            "document_id": str(adj.id),
            "amount": str(adj.amount),
            "description": adj.reason or "Điều chỉnh công nợ",
            "can_view_document": can_view_adj,
        })

    # 4. Bù trừ công nợ (Netting)
    for net in p.nettings.filter(status="POSTED").order_by("netting_date"):
        can_view_net = verify_document_view_permission(user, company, "AP.NETTING", str(net.id))
        events.append({
            "event_type": "NETTING",
            "event_type_display": "Bù trừ công nợ",
            "event_date": str(net.netting_date),
            "document_type": "AP.NETTING",
            "document_number": net.number,
            "document_id": str(net.id),
            "amount": str(-net.amount),
            "description": net.reason or f"Bù trừ theo đối chiếu {net.counterparty_reference}",
            "can_view_document": can_view_net,
        })

    # Sắp xếp timeline theo thứ tự ngày tăng dần
    events.sort(key=lambda x: x["event_date"])

    return {
        "payable": {
            "id": str(p.id),
            "supplier_id": str(p.supplier.id),
            "supplier_name": p.supplier.name,
            "supplier_code": p.supplier.code,
            "document_number": src_info["document_number"],
            "document_date": str(src_info["document_date"]),
            "due_date": str(p.due_date) if p.due_date else None,
            "currency_code": p.currency.code,
            "original_amount": str(p.original_amount),
            "paid_amount": str(p.paid_amount),
            "remaining_amount": str(p.balance),
            "status": p.status,
            "source_type": src_info["source_type"],
            "source_document_id": src_info["source_document_id"],
        },
        "events": events,
    }


def export_supplier_debt_summary_excel(company, data):
    """
    Xuất Excel Báo cáo Tổng hợp công nợ nhà cung cấp (AF-11, BR-05, BR-44, BR-46).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Tổng hợp công nợ NCC"

    # Styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    meta_font = Font(name="Calibri", size=10, italic=True)
    overdue_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append([company.name])
    ws.cell(1, 1).font = bold_font
    ws.append(["BÁO CÁO THEO DÕI CÔNG NỢ NHÀ CUNG CẤP"])
    ws.cell(2, 1).font = title_font
    ws.append([f"Kỳ theo dõi: Từ ngày {data['from_date']} đến ngày {data['to_date']} | Thời điểm xuất: {timezone.localtime().strftime('%d/%m/%Y %H:%M:%S')}"])
    ws.cell(3, 1).font = meta_font
    ws.append([])

    headers = [
        "STT", "Mã NCC", "Tên nhà cung cấp", "Mã số thuế",
        "Dư phải trả đầu kỳ", "Trả trước đầu kỳ",
        "Phát sinh tăng", "Phát sinh giảm",
        "Dư phải trả cuối kỳ", "Trả trước cuối kỳ",
        "Nợ đến hạn", "Nợ quá hạn", "Hạn TT gần nhất", "Số ngày quá hạn lớn nhất"
    ]
    ws.append(headers)
    header_row_idx = 5
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 1
    for r in data["results"]:
        row_data = [
            row_num,
            r["supplier_code"],
            r["supplier_name"],
            r["tax_code"],
            float(r["opening_payable_balance"]),
            float(r["opening_advance_balance"]),
            float(r["period_increase"]),
            float(r["period_decrease"]),
            float(r["closing_payable_balance"]),
            float(r["closing_advance_balance"]),
            float(r["due_amount"]),
            float(r["overdue_amount"]),
            r["earliest_due_date"] or "",
            r["max_overdue_days"],
        ]
        ws.append(row_data)
        curr_row = ws.max_row
        is_overdue = float(r["overdue_amount"]) > 0

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if is_overdue and col_idx in (12, 14):
                c.fill = overdue_fill
                c.font = bold_font

            if col_idx in (5, 6, 7, 8, 9, 10, 11, 12):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 14):
                c.alignment = Alignment(horizontal="center")
            elif col_idx in (2, 4, 13):
                c.alignment = Alignment(horizontal="center")
            else:
                c.alignment = Alignment(horizontal="left")

        row_num += 1

    # Dòng tổng cộng
    metrics = data["summary_metrics"]
    total_row = [
        "Tổng cộng", "", "", "",
        float(metrics["total_opening_payable"]),
        float(metrics["total_opening_advance"]),
        float(metrics["total_period_increase"]),
        float(metrics["total_period_decrease"]),
        float(metrics["total_closing_payable"]),
        float(metrics["total_closing_advance"]),
        float(metrics["total_due_amount"]),
        float(metrics["total_overdue_amount"]),
        "", ""
    ]
    ws.append(total_row)
    t_row = ws.max_row
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=t_row, column=col_idx)
        c.font = bold_font
        c.border = thin_border
        if col_idx in (5, 6, 7, 8, 9, 10, 11, 12):
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")

    # Column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_supplier_debt_detail_excel(company, data):
    """
    Xuất Excel Chi tiết các khoản công nợ của một nhà cung cấp (AF-12, BR-43, BR-44, BR-46).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Chi tiết công nợ NCC"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    meta_font = Font(name="Calibri", size=10, italic=True)
    overdue_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    supp = data["supplier"]
    ws.append([company.name])
    ws.cell(1, 1).font = bold_font
    ws.append([f"CHI TIẾT CÔNG NỢ NHÀ CUNG CẤP: {supp['name']} ({supp['code']})"])
    ws.cell(2, 1).font = title_font
    ws.append([f"Kỳ theo dõi: Từ ngày {data['from_date']} đến ngày {data['to_date']} | Mã số thuế: {supp['tax_code']}"])
    ws.cell(3, 1).font = meta_font
    ws.append([
        f"Dư đầu kỳ: {float(supp['opening_payable_balance']):,.2f} | "
        f"Tăng: {float(supp['period_increase']):,.2f} | "
        f"Giảm: {float(supp['period_decrease']):,.2f} | "
        f"Dư cuối kỳ: {float(supp['closing_payable_balance']):,.2f} | "
        f"Trả trước cuối kỳ: {float(supp['closing_advance_balance']):,.2f}"
    ])
    ws.cell(4, 1).font = meta_font
    ws.append([])

    headers = [
        "STT", "Số chứng từ", "Ngày chứng từ", "Ngày hạch toán",
        "Nguồn phát sinh", "Số hóa đơn", "Tiền gốc", "Tăng thêm",
        "Đã thanh toán", "Giảm trừ", "Đã bù trừ", "Còn phải trả",
        "Hạn thanh toán", "Số ngày còn lại", "Số ngày quá hạn", "Nhóm tuổi nợ", "Trạng thái"
    ]
    ws.append(headers)
    header_row_idx = 6
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_num = 1
    for r in data["results"]:
        row_data = [
            row_num,
            r["document_number"],
            r["document_date"],
            r["posting_date"],
            r["source_type_display"],
            r["invoice_number"],
            float(r["original_amount"]),
            float(r["increase_amount"]),
            float(r["paid_amount"]),
            float(r["discount_amount"]),
            float(r["netted_amount"]),
            float(r["remaining_amount"]),
            r["due_date"] or "",
            r["days_remaining"],
            r["overdue_days"],
            r["aging_bucket_display"],
            r["status_display"],
        ]
        ws.append(row_data)
        curr_row = ws.max_row
        is_overdue = r["is_overdue"]

        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.font = regular_font
            c.border = thin_border
            if is_overdue:
                c.fill = overdue_fill
                if col_idx in (12, 15, 17):
                    c.font = bold_font

            if col_idx in (7, 8, 9, 10, 11, 12):
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            elif col_idx in (1, 3, 4, 13, 14, 15, 16, 17):
                c.alignment = Alignment(horizontal="center")
            else:
                c.alignment = Alignment(horizontal="left")

        row_num += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

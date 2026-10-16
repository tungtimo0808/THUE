from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from iam.models import CompanyMembership, Permission, Role, RolePermission, UserRoleAssignment
from master_data.models import Currency, Party
from organization.models import Company
from .models import (Payable, PayableAdjustment, Payment, PaymentAllocationVoucher,
                     PaymentAllocationVoucherLine, SupplierInvoice)


class PayableReportUcAp10Tests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.start = self.today - timedelta(days=5)
        self.currency = Currency.objects.get(code="VND")
        self.company = Company.objects.create(code="AP10", name="AP10 Company", legal_name="AP10",
                                              tax_code="0100000010", accounting_currency=self.currency,
                                              legal_reporting_currency=self.currency)
        self.user = get_user_model().objects.create_superuser(username="ap10-user", email="ap10@example.com", password="test")
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.supplier = Party.objects.create(company=self.company, code="NCC10", name="Supplier 10")
        self.base = f"/api/v1/accounts-payable/reports/?company_id={self.company.id}&start_date={self.start}&end_date={self.today}"

    def payable(self, number, supplier=None, posting_day=None, amount="100"):
        supplier = supplier or self.supplier
        posting_day = posting_day or self.start
        invoice = SupplierInvoice.objects.create(company=self.company, supplier=supplier,
                                                  supplier_tax_code="0100000010", invoice_series="AA",
                                                  invoice_number=number, invoice_date=posting_day,
                                                  posting_date=posting_day, currency=self.currency,
                                                  created_by=self.user.iam_account, status="POSTED")
        return Payable.objects.create(company=self.company, supplier=supplier, invoice=invoice,
                                      currency=self.currency, exchange_rate=1, original_amount=amount,
                                      balance=amount, due_date=self.today)

    def payment(self, number, amount, day=None, supplier=None):
        return Payment.objects.create(company=self.company, supplier=supplier or self.supplier,
                                      currency=self.currency, payment_date=day or self.today,
                                      amount=amount, status="POSTED", created_by=self.user.iam_account,
                                      number=number)

    def test_opening_movements_negative_advance_and_zero_supplier_filter(self):
        payable = self.payable("INV10", posting_day=self.start - timedelta(days=2))
        self.payment("ADV10", "130", day=self.start - timedelta(days=1))
        PayableAdjustment.objects.create(company=self.company, payable=payable,
                                         adjustment_date=self.start, amount="20", reason="Tăng",
                                         status="POSTED", created_by=self.user.iam_account)
        self.payment("PAY10", "30")
        paid_supplier = Party.objects.create(company=self.company, code="PAID10", name="Paid supplier")
        self.payable("PAID-INV10", supplier=paid_supplier, amount="50")
        self.payment("PAID-PAY10", "50", supplier=paid_supplier)
        response = self.client.get(self.base)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["count"], 1)
        row = response.data["rows"][0]
        self.assertEqual(row["opening_balance"], Decimal("-30"))
        self.assertEqual(row["period_increase"], Decimal("20"))
        self.assertEqual(row["period_decrease"], Decimal("30"))
        self.assertEqual(row["closing_balance"], Decimal("-40"))
        all_rows = self.client.get(self.base + "&only_nonzero_closing=false")
        self.assertEqual(all_rows.data["count"], 2)
        self.assertEqual(all_rows.data["totals"]["closing_balance"], Decimal("-40"))

    def test_detail_source_and_validation(self):
        payable = self.payable("DETAIL10", posting_day=self.start)
        self.payment("PAYDETAIL10", "15")
        detail = self.client.get(self.base + "&report_type=DETAIL")
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertEqual([row["transaction_type"] for row in detail.data["rows"]], ["PURCHASE", "PAYMENT"])
        self.assertEqual(detail.data["rows"][-1]["balance_after"], Decimal("85"))
        transaction_id = detail.data["rows"][0]["transaction_id"]
        source = self.client.get(self.base.replace("?", "source/?") + f"&transaction_id={transaction_id}")
        self.assertEqual(source.status_code, 200, source.data)
        self.assertEqual(source.data["source_id"], str(payable.invoice_id))
        self.assertEqual(self.client.get(self.base.replace("?", "source/?") + "&transaction_id=missing").status_code, 400)
        self.assertEqual(self.client.get(self.base + "&report_type=").status_code, 400)
        self.assertEqual(self.client.get(self.base + "&page_size=30").status_code, 400)
        self.assertEqual(self.client.get(self.base + f"&end_date={self.today+timedelta(days=1)}").status_code, 400)
        drill = self.client.get(self.base.replace("?", "supplier-detail/?") + f"&supplier_id={self.supplier.id}")
        self.assertEqual(drill.status_code, 200, drill.data)
        self.assertEqual(drill.data["count"], 2)

    def test_exports_contain_all_rows_despite_pagination(self):
        self.payable("EXPORT10", posting_day=self.start)
        for i in range(21):
            self.payment(f"PAY-{i:02d}", "1")
        filtered = self.base + "&report_type=DETAIL&page_size=20"
        page = self.client.get(filtered)
        self.assertEqual(page.status_code, 200, page.data)
        self.assertEqual(page.data["count"], 22)
        self.assertEqual(len(page.data["rows"]), 20)
        self.assertEqual(page.data["total_pages"], 2)
        xlsx = self.client.get(filtered.replace("?", "export-excel/?"))
        self.assertEqual(xlsx.status_code, 200)
        sheet = load_workbook(BytesIO(xlsx.content), read_only=True).active
        self.assertEqual(sheet.max_row, 28)  # metadata + opening + header + 22 transactions + total
        pdf = self.client.get(filtered.replace("?", "export-pdf/?"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_report_does_not_include_unposted_or_cancelled_source(self):
        payable = self.payable("DRAFT10")
        payable.invoice.status = "DRAFT"
        payable.invoice.save()
        response = self.client.get(self.base)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["count"], 0)
        self.assertEqual(response.data["message"], "Không có dữ liệu công nợ phù hợp với điều kiện báo cáo.")

    def test_allocation_does_not_double_count_disbursed_payment(self):
        payable = self.payable("ALLOC10")
        payment = self.payment("PAID10", "40")
        voucher = PaymentAllocationVoucher.objects.create(
            company=self.company, number="PB10", payment=payment, allocation_date=self.today,
            description="Phân bổ", total_amount="40", created_by=self.user.iam_account)
        PaymentAllocationVoucherLine.objects.create(voucher=voucher, payable=payable, line_no=1, amount="40")
        response = self.client.get(self.base)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["totals"]["period_decrease"], Decimal("40"))
        self.assertEqual(response.data["totals"]["closing_balance"], Decimal("60"))

    def test_rbac_permissions_ex_ap_rpt_01(self):
        regular_user = get_user_model().objects.create_user(
            username="ap10-regular", email="ap10reg@example.com", password="test"
        )
        membership = CompanyMembership.objects.create(
            company=self.company, user=regular_user.iam_account, membership_status="ACTIVE",
            created_by=self.user.iam_account
        )
        client = APIClient()
        client.force_authenticate(regular_user)

        # 1. No role / permission -> 403 Forbidden with EX-AP-RPT-01 message
        resp = client.get(self.base)
        self.assertEqual(resp.status_code, 403)
        self.assertIn("Bạn không có quyền xem báo cáo công nợ phải trả.", str(resp.data))

        # 2. Grant role with AP.REPORT.VIEW -> 200 OK
        role = Role.objects.create(company=self.company, code="AP10_VIEWER", name="AP10 Viewer")
        perm = Permission.objects.filter(code="AP.REPORT.VIEW").first()
        if perm:
            RolePermission.objects.create(role=role, permission=perm)
        UserRoleAssignment.objects.create(
            company_membership=membership, role=role, assigned_by=self.user.iam_account
        )

        resp_ok = client.get(self.base)
        self.assertEqual(resp_ok.status_code, 200)

    def test_exception_flows_validation_messages(self):
        # EX-AP-RPT-02: Missing report_type
        r = self.client.get(self.base + "&report_type=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("report_type")), "Vui lòng chọn Mẫu báo cáo.")

        # EX-AP-RPT-03: Missing start_date
        r = self.client.get(self.base.replace(f"start_date={self.start}", "start_date="))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("start_date")), "Vui lòng nhập Ngày bắt đầu kỳ báo cáo.")

        # EX-AP-RPT-04: Missing end_date
        r = self.client.get(self.base.replace(f"end_date={self.today}", "end_date="))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("end_date")), "Vui lòng nhập Ngày kết thúc kỳ báo cáo.")

        # EX-AP-RPT-05: Missing currency_id
        r = self.client.get(self.base + "&currency_id=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("currency_id")), "Vui lòng chọn Loại tiền.")

        # EX-AP-RPT-06: start_date > end_date
        r = self.client.get(self.base.replace(f"start_date={self.start}", f"start_date={self.today + timedelta(days=1)}"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("start_date")), "Ngày bắt đầu kỳ báo cáo không được lớn hơn Ngày kết thúc kỳ báo cáo.")

        # EX-AP-RPT-07: end_date > today
        future = self.today + timedelta(days=1)
        r = self.client.get(self.base.replace(f"end_date={self.today}", f"end_date={future}"))
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("end_date")), "Ngày kết thúc kỳ báo cáo không được lớn hơn ngày hiện tại.")

        # EX-AP-RPT-11: Invalid transaction_id in source endpoint
        url_source = self.base.replace("?", "source/?")
        r = self.client.get(url_source + "&transaction_id=nonexistent")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("transaction_id")), "Không tìm thấy chứng từ nguồn của giao dịch đã chọn.")

    def test_vietnamese_labels_and_alternative_flows(self):
        payable = self.payable("VN10", posting_day=self.start)
        self.payment("PAY_VN10", "30")

        # AF-AP-RPT-01: Mẫu báo cáo tiếng Việt "Tổng hợp theo nhà cung cấp"
        r_sum = self.client.get(self.base + "&report_type=Tổng hợp theo nhà cung cấp")
        self.assertEqual(r_sum.status_code, 200)
        self.assertEqual(r_sum.data["count"], 1)
        self.assertIn("total_opening", r_sum.data["totals"])
        self.assertIn("total_increase", r_sum.data["totals"])
        self.assertIn("total_decrease", r_sum.data["totals"])
        self.assertIn("total_closing", r_sum.data["totals"])

        # AF-AP-RPT-02: Mẫu báo cáo tiếng Việt "Chi tiết phát sinh"
        r_det = self.client.get(self.base + "&report_type=Chi tiết phát sinh")
        self.assertEqual(r_det.status_code, 200)
        self.assertEqual(r_det.data["count"], 2)
        types_vn = [row["transaction_type_display"] for row in r_det.data["rows"]]
        self.assertEqual(types_vn, ["Mua hàng", "Thanh toán"])

        # AF-AP-RPT-03: Hiển thị cả NCC hết công nợ với has_balance_only=false
        zero_supplier = Party.objects.create(company=self.company, code="ZERO10", name="Zero supplier")
        self.payable("ZERO_INV", supplier=zero_supplier, amount="40")
        self.payment("ZERO_PAY", "40", supplier=zero_supplier)
        r_all = self.client.get(self.base + "&has_balance_only=false")
        self.assertEqual(r_all.status_code, 200)
        self.assertEqual(r_all.data["count"], 2)

    def test_source_document_permission_check_br_ap_rpt_10(self):
        payable = self.payable("SRC10", posting_day=self.start)
        regular_user = get_user_model().objects.create_user(
            username="ap10-docuser", email="docuser@example.com", password="test"
        )
        membership = CompanyMembership.objects.create(
            company=self.company, user=regular_user.iam_account, membership_status="ACTIVE",
            created_by=self.user.iam_account
        )
        client = APIClient()
        client.force_authenticate(regular_user)

        # Grant AP.REPORT.VIEW only
        role = Role.objects.create(company=self.company, code="AP10_DOCVIEWER", name="AP10 Doc Viewer")
        perm_rpt = Permission.objects.filter(code="AP.REPORT.VIEW").first()
        if perm_rpt:
            RolePermission.objects.create(role=role, permission=perm_rpt)
        UserRoleAssignment.objects.create(
            company_membership=membership, role=role, assigned_by=self.user.iam_account
        )

        detail_resp = client.get(self.base + "&report_type=DETAIL")
        self.assertEqual(detail_resp.status_code, 200)
        tx_id = detail_resp.data["rows"][0]["transaction_id"]

        # Call source without AP.INVOICE.VIEW -> 403 Forbidden
        url_source = self.base.replace("?", "source/?") + f"&transaction_id={tx_id}"
        resp_denied = client.get(url_source)
        self.assertEqual(resp_denied.status_code, 403)
        self.assertIn("Bạn không có quyền xem chứng từ nguồn này.", str(resp_denied.data))

        # Grant AP.INVOICE.VIEW -> 200 OK
        perm_inv = Permission.objects.filter(code="AP.INVOICE.VIEW").first()
        if perm_inv:
            RolePermission.objects.create(role=role, permission=perm_inv)

        resp_ok = client.get(url_source)
        self.assertEqual(resp_ok.status_code, 200)
        self.assertEqual(resp_ok.data["source_id"], str(payable.invoice_id))

    def test_pagination_options_step_27(self):
        for size in (20, 50, 100):
            r = self.client.get(self.base + f"&page_size={size}")
            self.assertEqual(r.status_code, 200)
            self.assertEqual(r.data["page_size"], size)

        # Invalid page_size
        r_bad = self.client.get(self.base + "&page_size=25")
        self.assertEqual(r_bad.status_code, 400)
        self.assertIn("page_size chỉ nhận 20, 50 hoặc 100", str(r_bad.data))

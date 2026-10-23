"""Internal XLSX export from a frozen, company-scoped tax data snapshot."""
from decimal import Decimal
from io import BytesIO
import re

from django.db import transaction
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import PermissionDenied, ValidationError

from iam.permissions import has_effective_permission

from . import models, services, vat


SCHEMA_VERSION = "TAX12-REPORT-1"
TYPES = {
    "VAT_LEDGER": ("Sổ VAT", "TAX.SUMMARY.VIEW", [
        ("source_key", "Mã nguồn", "text"), ("direction", "Chiều", "text"),
        ("decision", "Quyết định", "text"), ("taxable_amount", "Tiền tính thuế", "number"),
        ("tax_amount", "Tiền thuế", "number"), ("deductible_tax_amount", "Thuế được khấu trừ", "number"),
        ("reason", "Lý do", "text"),
    ]),
    "TAX_OBLIGATION": ("Nghĩa vụ thuế", "TAX.OBLIGATION.VIEW", [
        ("tax_type", "Loại thuế", "text"), ("due_date", "Hạn nộp", "text"),
        ("assessed_amount", "Phải nộp", "number"), ("adjustment_amount", "Điều chỉnh", "number"),
        ("paid_amount", "Đã ghi nhận nội bộ", "number"), ("outstanding_amount", "Còn lại", "number"),
        ("basis_status", "Căn cứ", "text"), ("status", "Trạng thái", "text"),
    ]),
}


@transaction.atomic
def generate(*, request, company, period, export_type, purpose):
    if period.company_id != company.id:
        raise ValidationError({"period": "Kỳ thuế không thuộc doanh nghiệp."})
    if export_type not in TYPES:
        raise ValidationError({"export_type": "Loại xuất chưa được hỗ trợ trong MVP."})
    if purpose not in {"REPORT", "REVIEW", "ARCHIVE"}:
        raise ValidationError({"purpose": "MVP chỉ xuất báo cáo nội bộ; không tạo file để nộp."})
    title, permission, columns = TYPES[export_type]
    if not request.user.is_superuser and not has_effective_permission(services.actor(request), company, permission):
        raise PermissionDenied("Thiếu quyền xem dữ liệu nguồn.")
    warnings = []
    source_version = ""
    if export_type == "VAT_LEDGER":
        calculation = models.VatCalculation.objects.filter(company=company, period=period).order_by("-version").first()
        if not calculation:
            raise ValidationError("Chưa có kết quả tổng hợp VAT cho kỳ này.")
        if not vat.is_current(calculation):
            raise ValidationError("Kết quả VAT đã cũ; cần tính lại trước khi xuất.")
        entries = list(models.VatLedgerEntry.objects.filter(company=company, period=period).order_by("direction", "source_key"))
        if not entries:
            raise ValidationError("VAT Ledger chưa có dòng dữ liệu để xuất.")
        rows = [{key: str(getattr(entry, key)) for key, _, _ in columns} for entry in entries]
        source_version = f"VAT-{calculation.version}"
        warnings = list(calculation.warnings)
    else:
        obligations = list(models.TaxObligation.objects.filter(company=company, period=period).order_by("tax_type", "id"))
        if not obligations:
            raise ValidationError("Chưa có nghĩa vụ thuế trong kỳ để xuất.")
        rows = [{key: str(getattr(item, key)) for key, _, _ in columns} for item in obligations]
        warnings.append("Số đã ghi nhận là dữ liệu nội bộ; chưa thay thế xác nhận hạch toán của NSNN.")
    if len(rows) > 10000:
        raise ValidationError("Bản xuất vượt giới hạn 10.000 dòng của MVP.")
    snapshot = {"export_type": export_type, "purpose": purpose, "title": title,
        "company_name": company.name, "period": period.code, "source_version": source_version,
        "columns": columns, "rows": rows, "row_count": len(rows), "schema_version": SCHEMA_VERSION,
        "notice": "Báo cáo nội bộ; không phải tệp hồ sơ nộp CQT."}
    obj = models.TaxExport.objects.create(company=company, period=period, format_code="XLSX",
        format_version=SCHEMA_VERSION, data_snapshot=snapshot, warnings=warnings,
        exported_by=services.actor(request))
    services.audit(request, obj, "EXPORT", after={"export_type": export_type, "purpose": purpose,
        "period": period.code, "row_count": len(rows), "schema_version": SCHEMA_VERSION})
    return obj


def xlsx_bytes(export):
    snapshot = export.data_snapshot
    if snapshot.get("export_type") not in TYPES or export.format_code != "XLSX":
        raise ValidationError("Bản xuất này không có tệp XLSX thuộc MVP.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dữ liệu thuế"
    def safe_text(value):
        text = str(value)
        return "'" + text if text.startswith(("=", "+", "-", "@")) else text
    sheet.append([safe_text(snapshot["title"]), safe_text(snapshot["company_name"]), safe_text(snapshot["period"])])
    sheet.append([snapshot["notice"]])
    sheet.append(["Phiên bản nguồn", snapshot.get("source_version", ""), "Schema", snapshot["schema_version"]])
    sheet.append([])
    sheet.append([label for _, label, _ in snapshot["columns"]])
    for cell in sheet[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="008A50")
    for row in snapshot["rows"]:
        values = []
        for key, _, kind in snapshot["columns"]:
            value = row.get(key, "")
            if kind == "number":
                number = Decimal(value)
                values.append(str(number) if len(number.as_tuple().digits) > 15 else float(number))
            else:
                values.append(safe_text(value))
        sheet.append(values)
    sheet.freeze_panes = "A6"
    sheet.auto_filter.ref = f"A5:{sheet.cell(row=max(sheet.max_row, 5), column=len(snapshot['columns'])).coordinate}"
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(40, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def filename(export):
    snapshot = export.data_snapshot
    period = re.sub(r"[^A-Za-z0-9_-]", "-", snapshot["period"])
    return f"thue-{snapshot['export_type'].lower()}-{period}-{str(export.id)[:8]}.xlsx"

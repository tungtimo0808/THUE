from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from openpyxl import load_workbook

from django.contrib.auth import get_user_model
from django.utils import timezone
from rest_framework.test import APITestCase, APIClient

from accounts_payable.models import Payable, PayableStatus, SupplierInvoice
from accounts_receivable.models import (
    Receivable, ReceivableSchedule, ReceivableStatus,
    Receipt, ReceiptStatus, ReceiptAllocation,
    DebtNettingVoucher, DebtNettingVoucherLine,
    DebtAdjustmentVoucher, DebtAdjustmentVoucherLine,
)
from general_ledger.models import (
    Account, AccountNature, AccountStatus,
    ChartOfAccounts, NormalBalance,
)
from iam.models import CompanyMembership, UserAccount
from master_data.models import Currency, DocumentType, Party, PartyRole
from organization.models import Company, Branch


class DueTrackingTests(APITestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.prior_10d = self.today - timedelta(days=10)
        self.prior_20d = self.today - timedelta(days=20)
        self.future_5d = self.today + timedelta(days=5)
        self.future_15d = self.today + timedelta(days=15)

        # 1. Currencies
        self.currency_vnd = Currency.objects.get(code="VND")
        self.currency_usd, _ = Currency.objects.get_or_create(
            code="USD",
            defaults={"name": "US Dollar", "numeric_code": "840", "decimal_places": 2}
        )

        # 2. Company & Branch
        self.company = Company.objects.create(
            code="COMP-DUE",
            name="Công ty Theo Dõi Công Nợ",
            legal_name="Công ty TNHH Theo Dõi Công Nợ",
            tax_code="0309998881",
            accounting_currency=self.currency_vnd,
            legal_reporting_currency=self.currency_vnd,
        )
        self.branch = Branch.objects.create(
            company=self.company,
            code="BR-DUE-01",
            name="Chi nhánh Miền Nam",
        )

        # 3. Chart of Accounts & Accounts
        self.chart = ChartOfAccounts.objects.create(
            company=self.company,
            code="COA-DUE",
            name="Due Tracking COA",
            effective_from=date(2026, 1, 1),
            is_default=True,
        )
        self.account_131 = Account.objects.create(
            company=self.company,
            chart_of_accounts=self.chart,
            code="131",
            name="Phải thu của khách hàng",
            account_type=AccountNature.ASSET,
            normal_balance=NormalBalance.DEBIT,
            status=AccountStatus.ACTIVE,
        )
        self.account_1312 = Account.objects.create(
            company=self.company,
            chart_of_accounts=self.chart,
            code="1312",
            name="Phải thu khách hàng ngoại tệ",
            account_type=AccountNature.ASSET,
            normal_balance=NormalBalance.DEBIT,
            status=AccountStatus.ACTIVE,
        )
        self.account_511 = Account.objects.create(
            company=self.company,
            chart_of_accounts=self.chart,
            code="511",
            name="Doanh thu bán hàng",
            account_type=AccountNature.REVENUE,
            normal_balance=NormalBalance.CREDIT,
            status=AccountStatus.ACTIVE,
        )

        # 4. Users (Admin and Staff)
        User = get_user_model()
        self.admin_user = User.objects.create_superuser(
            username="due_admin", email="due_admin@example.com", password="password"
        )
        CompanyMembership.objects.create(
            company=self.company,
            user=self.admin_user.iam_account,
            created_by=self.admin_user.iam_account,
        )

        # Assigned Staff User
        self.staff_user = User.objects.create_user(
            username="staff_sales1", email="staff1@example.com", password="password"
        )
        self.staff_account = self.staff_user.iam_account
        self.staff_account.display_name = "Lê Văn Nhân Viên"
        self.staff_account.save(update_fields=("display_name",))
        CompanyMembership.objects.create(
            company=self.company,
            user=self.staff_account,
            created_by=self.admin_user.iam_account,
        )

        # 5. Customers
        # Customer 1 (Assigned to staff_account)
        self.customer1 = Party.objects.create(
            company=self.company,
            code="KH-ALPHA",
            name="Công ty Alpha",
            tax_code="0301112223",
            customer_group="VIP",
            owner=self.staff_account,
            default_currency=self.currency_vnd,
            status="ACTIVE",
        )
        PartyRole.objects.create(party=self.customer1, role_type="CUSTOMER", active_from=self.today)

        # Customer 2 (Unassigned owner)
        self.customer2 = Party.objects.create(
            company=self.company,
            code="KH-BETA",
            name="Công ty Beta",
            tax_code="0304445556",
            customer_group="RETAIL",
            owner=None,
            default_currency=self.currency_usd,
            status="ACTIVE",
        )
        PartyRole.objects.create(party=self.customer2, role_type="CUSTOMER", active_from=self.today)

        # 6. Test Client
        self.client = APIClient()
        self.client.force_authenticate(self.admin_user)
        self.base_url = f"/api/v1/accounts-receivable/due-tracking/?company_id={self.company.id}"

        # 7. Seed Receivables
        # Rec 1: Customer 1, VND, OVERDUE (due 10 days ago), single payment
        self.rec_overdue = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-OVERDUE-01",
            source_document_number="HD-2026-001",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("1000000"),
            balance=Decimal("1000000"),
            posting_date=self.prior_20d,
            document_date=self.prior_20d,
            due_date=self.prior_10d,
            status=ReceivableStatus.OPEN,
        )

        # Rec 2: Customer 1, VND, DUE_TODAY (due today), single payment
        self.rec_due_today = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-TODAY-02",
            source_document_number="HD-2026-002",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("500000"),
            balance=Decimal("500000"),
            posting_date=self.prior_10d,
            document_date=self.prior_10d,
            due_date=self.today,
            status=ReceivableStatus.OPEN,
        )

        # Rec 3: Customer 1, VND, UPCOMING (due in 5 days, within 7 days window)
        self.rec_upcoming = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-UPCOMING-03",
            source_document_number="HD-2026-003",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("800000"),
            balance=Decimal("800000"),
            posting_date=self.prior_10d,
            document_date=self.prior_10d,
            due_date=self.future_5d,
            status=ReceivableStatus.OPEN,
        )

        # Rec 4: Customer 1, VND, NOT_DUE (due in 15 days, beyond 7 days window)
        self.rec_not_due = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-NOTDUE-04",
            source_document_number="HD-2026-004",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("1200000"),
            balance=Decimal("1200000"),
            posting_date=self.prior_10d,
            document_date=self.prior_10d,
            due_date=self.future_15d,
            status=ReceivableStatus.OPEN,
        )

        # Rec 5: Customer 2, USD, OVERDUE (due 15 days ago)
        self.rec_usd_overdue = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-USD-01",
            source_document_number="INV-USD-101",
            customer=self.customer2,
            currency=self.currency_usd,
            account=self.account_1312,
            exchange_rate=Decimal("25000"),
            original_amount=Decimal("100"),
            balance=Decimal("100"),
            posting_date=self.prior_20d,
            document_date=self.prior_20d,
            due_date=self.today - timedelta(days=15),
            status=ReceivableStatus.OPEN,
        )

    # =========================================================================
    # Test Cases
    # =========================================================================

    def test_01_default_filter_shows_overdue_and_multi_currency_summaries(self):
        """
        Bước 1 / BR05 / BR09 / BR10:
        Truy cập mặc định hiển thị công nợ Quá hạn tại ngày hiện tại,
        sắp xếp ngày quá hạn giảm dần, kèm tổng hợp theo từng loại tiền và tổng quy đổi.
        """
        res = self.client.get(self.base_url)
        self.assertEqual(res.status_code, 200, res.data)

        # Default filter is OVERDUE
        self.assertEqual(res.data["filters"]["due_category"], "OVERDUE")
        self.assertEqual(res.data["as_of_date"], str(self.today))

        # Items returned should all be OVERDUE
        items = res.data["items"]
        self.assertEqual(len(items), 2)  # rec_usd_overdue (15d) and rec_overdue (10d)
        for it in items:
            self.assertEqual(it["due_category"], "OVERDUE")
            self.assertGreater(it["days_overdue"], 0)

        # BR10: Sort by days_overdue DESC
        self.assertEqual(items[0]["receivable_code"], "AR-USD-01")
        self.assertEqual(items[0]["days_overdue"], 15)
        self.assertEqual(items[1]["receivable_code"], "AR-OVERDUE-01")
        self.assertEqual(items[1]["days_overdue"], 10)

        # BR09 & Mục 2.3: Multi-currency summaries present full debt distribution
        summaries = {s["currency_code"]: s for s in res.data["summaries"]}
        self.assertIn("VND", summaries)
        self.assertIn("USD", summaries)

        vnd_sum = summaries["VND"]
        self.assertEqual(Decimal(vnd_sum["overdue_balance"]), Decimal("1000000"))
        self.assertEqual(Decimal(vnd_sum["due_today_balance"]), Decimal("500000"))
        self.assertEqual(Decimal(vnd_sum["upcoming_balance"]), Decimal("800000"))
        self.assertEqual(Decimal(vnd_sum["not_due_balance"]), Decimal("1200000"))
        self.assertEqual(Decimal(vnd_sum["total_remaining_balance"]), Decimal("3500000"))

        usd_sum = summaries["USD"]
        self.assertEqual(Decimal(usd_sum["overdue_balance"]), Decimal("100"))
        self.assertEqual(Decimal(usd_sum["total_remaining_balance"]), Decimal("100"))

        # Base summary (quy đổi theo tỷ giá ghi sổ)
        base = res.data["base_summary"]
        # VND: 3,500,000 * 1 = 3,500,000
        # USD: 100 * 25,000 = 2,500,000
        # Total base: 6,000,000
        self.assertEqual(Decimal(base["base_total_remaining_balance"]), Decimal("6000000"))

    def test_02_classify_due_today_not_considered_overdue(self):
        """
        BR05: Khoản đến hạn đúng ngày theo dõi (due_date == as_of_date)
        chưa được coi là quá hạn; thuộc nhóm DUE_TODAY.
        """
        url = f"{self.base_url}&due_category=DUE_TODAY"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-TODAY-02")
        self.assertEqual(items[0]["due_category"], "DUE_TODAY")
        self.assertEqual(items[0]["days_overdue"], 0)
        self.assertEqual(items[0]["days_due"], 0)

    def test_03_classify_upcoming_within_window(self):
        """
        Bước 2 / AF01 / BR06: Lọc Sắp đến hạn với số ngày cấu hình (upcoming_days).
        """
        url = f"{self.base_url}&due_category=UPCOMING&upcoming_days=7"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-UPCOMING-03")
        self.assertEqual(items[0]["due_category"], "UPCOMING")
        self.assertEqual(items[0]["days_due"], 5)
        self.assertEqual(items[0]["days_overdue"], 0)

    def test_04_classify_not_due_after_upcoming_window(self):
        """
        Mục 3: Khoản có hạn sau khoảng sắp đến hạn được phân loại vào NOT_DUE.
        """
        url = f"{self.base_url}&due_category=NOT_DUE&upcoming_days=7"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-NOTDUE-04")
        self.assertEqual(items[0]["due_category"], "NOT_DUE")
        self.assertEqual(items[0]["days_due"], 15)

    def test_05_classify_no_due_date_handled_separately(self):
        """
        AF04 / BR07: Khoản chưa có hạn thanh toán được xếp riêng vào NO_DUE_DATE;
        không tự coi là quá hạn hoặc suy đoán từ ngày chứng từ.
        """
        rec_no_due = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-NODUE-05",
            source_document_number="HD-2026-005",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("300000"),
            balance=Decimal("300000"),
            posting_date=self.prior_10d,
            document_date=self.prior_10d,
            due_date=None,  # No due date
            status=ReceivableStatus.OPEN,
        )

        url = f"{self.base_url}&due_category=NO_DUE_DATE"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-NODUE-05")
        self.assertEqual(items[0]["due_category"], "NO_DUE_DATE")
        self.assertIsNone(items[0]["due_date"])
        self.assertIsNone(items[0]["days_due"])
        self.assertIsNone(items[0]["days_overdue"])

    def test_06_installment_schedules_tracked_individually_br03(self):
        """
        BR03: Nếu khoản phải thu chia nhiều đợt thanh toán, theo dõi từng đợt riêng;
        không hiển thị thêm dòng tổng khoản phải thu gây cộng trùng.
        """
        rec_multi = Receivable.objects.create(
            company=self.company,
            branch=self.branch,
            code="AR-MULTI-06",
            source_document_number="HD-MULTI-06",
            customer=self.customer1,
            currency=self.currency_vnd,
            account=self.account_131,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("900000"),
            balance=Decimal("900000"),
            posting_date=self.prior_20d,
            due_date=self.prior_10d,
            status=ReceivableStatus.OPEN,
        )
        # Installment 1: OVERDUE (due 10 days ago), 300,000
        sch1 = ReceivableSchedule.objects.create(
            company=self.company,
            receivable=rec_multi,
            installment_no=1,
            amount=Decimal("300000"),
            balance=Decimal("300000"),
            due_date=self.prior_10d,
            status=ReceivableStatus.OPEN,
        )
        # Installment 2: DUE_TODAY (due today), 300,000
        sch2 = ReceivableSchedule.objects.create(
            company=self.company,
            receivable=rec_multi,
            installment_no=2,
            amount=Decimal("300000"),
            balance=Decimal("300000"),
            due_date=self.today,
            status=ReceivableStatus.OPEN,
        )
        # Installment 3: UPCOMING (due in 5 days), 300,000
        sch3 = ReceivableSchedule.objects.create(
            company=self.company,
            receivable=rec_multi,
            installment_no=3,
            amount=Decimal("300000"),
            balance=Decimal("300000"),
            due_date=self.future_5d,
            status=ReceivableStatus.OPEN,
        )

        # Query ALL to check distinct installment lines
        url = f"{self.base_url}&due_category=ALL&search=HD-MULTI-06"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        items = res.data["items"]
        # Exactly 3 lines (one for each installment), NO duplicate parent line (BR03)
        self.assertEqual(len(items), 3)

        categories = {it["installment_no"]: it["due_category"] for it in items}
        self.assertEqual(categories[1], "OVERDUE")
        self.assertEqual(categories[2], "DUE_TODAY")
        self.assertEqual(categories[3], "UPCOMING")

    def test_07_in_period_reductions_applied_accurately_br04(self):
        """
        BR04: Số dư tại ngày theo dõi tính theo các nghiệp vụ tăng, giảm có hiệu lực đến ngày đó.
        Tiền thu chưa phân bổ không tự làm giảm khoản phải thu.
        """
        # 1. Receipt and Allocation (decrease 300,000)
        receipt = Receipt.objects.create(
            company=self.company,
            branch=self.branch,
            customer=self.customer1,
            currency=self.currency_vnd,
            amount=Decimal("400000"),  # 300,000 allocated, 100,000 unallocated
            allocated_amount=Decimal("300000"),
            receipt_date=self.prior_10d,
            external_type="BANK",
            external_reference="REC-TEST-01",
            status=ReceiptStatus.CONFIRMED,
            created_by=self.admin_user.iam_account,
        )
        ReceiptAllocation.objects.create(
            receipt=receipt,
            receivable=self.rec_overdue,
            amount=Decimal("300000"),
            created_by=self.admin_user.iam_account,
            effective_date=self.prior_10d,
        )

        # 2. Netting Voucher (decrease 150,000)
        net_voucher = DebtNettingVoucher.objects.create(
            company=self.company,
            branch=self.branch,
            number="NET-TEST-01",
            document_date=self.prior_10d,
            posting_date=self.prior_10d,
            customer=self.customer1,
            supplier=self.customer1,
            currency=self.currency_vnd,
            total_amount=Decimal("150000"),
            receivable_account=self.account_131,
            payable_account=self.account_131,
            status=DebtNettingVoucher.Status.POSTED,
            created_by=self.admin_user.iam_account,
        )
        invoice = SupplierInvoice.objects.create(
            company=self.company,
            branch=self.branch,
            supplier=self.customer1,
            supplier_tax_code="0301112223",
            invoice_series="AA/26E",
            invoice_number="INV-DUE-01",
            invoice_date=self.prior_10d,
            currency=self.currency_vnd,
            exchange_rate=Decimal("1"),
            total_amount=Decimal("150000"),
            status="POSTED",
            created_by=self.admin_user.iam_account,
        )
        payable = Payable.objects.create(
            company=self.company,
            branch=self.branch,
            supplier=self.customer1,
            invoice=invoice,
            currency=self.currency_vnd,
            exchange_rate=Decimal("1"),
            original_amount=Decimal("150000"),
            balance=Decimal("150000"),
            due_date=self.prior_10d,
            status=PayableStatus.OPEN,
        )
        DebtNettingVoucherLine.objects.create(
            voucher=net_voucher,
            line_no=1,
            receivable=self.rec_overdue,
            payable=payable,
            amount=Decimal("150000"),
        )

        # 3. Adjustment Voucher (decrease 50,000)
        adj_voucher = DebtAdjustmentVoucher.objects.create(
            company=self.company,
            branch=self.branch,
            number="ADJ-TEST-01",
            document_date=self.prior_10d,
            posting_date=self.prior_10d,
            customer=self.customer1,
            adjustment_type="DECREASE",
            currency=self.currency_vnd,
            receivable_account=self.account_131,
            adjustment_reason="Chiết khấu thương mại",
            document_basis="Thỏa thuận",
            total_amount=Decimal("50000"),
            status=DebtAdjustmentVoucher.Status.POSTED,
            created_by=self.admin_user.iam_account,
        )
        DebtAdjustmentVoucherLine.objects.create(
            voucher=adj_voucher,
            line_no=1,
            receivable=self.rec_overdue,
            offset_account=self.account_511,
            amount=Decimal("50000"),
        )

        url = f"{self.base_url}&due_category=OVERDUE&search=AR-OVERDUE-01"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        item = res.data["items"][0]
        self.assertEqual(Decimal(item["original_amount"]), Decimal("1000000"))
        # Reduced = 300,000 (alloc) + 150,000 (net) + 50,000 (adj) = 500,000
        self.assertEqual(Decimal(item["reduced_amount"]), Decimal("500000"))
        # Remaining = 1,000,000 - 500,000 = 500,000 (The unallocated 100,000 does NOT reduce remaining)
        self.assertEqual(Decimal(item["remaining_balance"]), Decimal("500000"))

    def test_08_fully_settled_items_excluded_from_tracking(self):
        """
        Mục 3: Chỉ theo dõi các khoản có số tiền còn phải thu lớn hơn 0 tại ngày theo dõi.
        """
        # Fully pay rec_overdue
        receipt = Receipt.objects.create(
            company=self.company,
            branch=self.branch,
            customer=self.customer1,
            currency=self.currency_vnd,
            amount=Decimal("1000000"),
            allocated_amount=Decimal("1000000"),
            receipt_date=self.prior_10d,
            external_type="BANK",
            external_reference="REC-FULL",
            status=ReceiptStatus.CONFIRMED,
            created_by=self.admin_user.iam_account,
        )
        ReceiptAllocation.objects.create(
            receipt=receipt,
            receivable=self.rec_overdue,
            amount=Decimal("1000000"),
            created_by=self.admin_user.iam_account,
            effective_date=self.prior_10d,
        )

        url = f"{self.base_url}&due_category=OVERDUE&search=AR-OVERDUE-01"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        # Fully paid item is excluded!
        self.assertEqual(len(res.data["items"]), 0)

    def test_09_historical_as_of_date_query_af03_br08(self):
        """
        AF03 / BR08: Xem công nợ tại ngày trong quá khứ.
        Giao dịch phát sinh sau ngày theo dõi chưa được tính.
        """
        # Make a payment TODAY for rec_overdue
        receipt = Receipt.objects.create(
            company=self.company,
            branch=self.branch,
            customer=self.customer1,
            currency=self.currency_vnd,
            amount=Decimal("500000"),
            allocated_amount=Decimal("500000"),
            receipt_date=self.today,  # payment made today
            external_type="BANK",
            external_reference="REC-TODAY",
            status=ReceiptStatus.CONFIRMED,
            created_by=self.admin_user.iam_account,
        )
        ReceiptAllocation.objects.create(
            receipt=receipt,
            receivable=self.rec_overdue,
            amount=Decimal("500000"),
            created_by=self.admin_user.iam_account,
            effective_date=self.today,
        )

        # Query as of 5 days ago (prior to today's payment)
        as_of_past = self.today - timedelta(days=5)
        url = f"{self.base_url}&as_of_date={as_of_past}&due_category=OVERDUE&search=AR-OVERDUE-01"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)

        item = res.data["items"][0]
        # At 5 days ago, payment was not yet made, so remaining_balance was full 1,000,000!
        self.assertEqual(Decimal(item["remaining_balance"]), Decimal("1000000"))
        self.assertEqual(Decimal(item["reduced_amount"]), Decimal("0"))

    def test_10_assigned_staff_filtering_af02(self):
        """
        AF02: Lọc công nợ theo nhân viên phụ trách khách hàng hoặc Chưa phân công (UNASSIGNED).
        """
        # Filter by staff Lê Văn Nhân Viên
        url_staff = f"{self.base_url}&due_category=ALL&assigned_staff_id={self.staff_account.id}"
        res_staff = self.client.get(url_staff)
        self.assertEqual(res_staff.status_code, 200)
        for it in res_staff.data["items"]:
            self.assertEqual(it["assigned_staff_id"], str(self.staff_account.id))
            self.assertEqual(it["customer_code"], "KH-ALPHA")

        # Filter by UNASSIGNED
        url_unassigned = f"{self.base_url}&due_category=ALL&assigned_staff_id=UNASSIGNED"
        res_unassigned = self.client.get(url_unassigned)
        self.assertEqual(res_unassigned.status_code, 200)
        for it in res_unassigned.data["items"]:
            self.assertIsNone(it["assigned_staff_id"])
            self.assertEqual(it["customer_code"], "KH-BETA")

    def test_11_overdue_days_range_filtering(self):
        """
        Mục 2.1: Lọc khoản quá hạn theo khoảng số ngày (overdue_days_min, overdue_days_max).
        rec_usd_overdue is overdue 15 days, rec_overdue is overdue 10 days.
        """
        # Filter 12 to 20 days -> only rec_usd_overdue (15d) matches
        url = f"{self.base_url}&due_category=OVERDUE&overdue_days_min=12&overdue_days_max=20"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-USD-01")

    def test_12_search_filter_by_customer_or_document_code(self):
        """
        Mục 2.1: Tìm kiếm theo mã khách hàng, tên khách hàng hoặc số chứng từ.
        """
        # Search by document number
        url = f"{self.base_url}&due_category=ALL&search=HD-2026-003"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        items = res.data["items"]
        self.assertEqual(len(items), 1)
        self.assertEqual(items[0]["receivable_code"], "AR-UPCOMING-03")

    def test_13_single_item_detail_endpoint_steps_3_4(self):
        """
        Bước 3-4: Endpoint item-detail trả về lịch thanh toán, chi tiết các khoản đã thu,
        phân bổ, bù trừ, điều chỉnh và số dư.
        """
        detail_url = (
            f"/api/v1/accounts-receivable/due-tracking/item-detail/?"
            f"company_id={self.company.id}&receivable_id={self.rec_overdue.id}"
        )
        res = self.client.get(detail_url)
        self.assertEqual(res.status_code, 200)

        data = res.data
        self.assertEqual(data["code"], "AR-OVERDUE-01")
        self.assertEqual(data["customer"]["code"], "KH-ALPHA")
        self.assertEqual(Decimal(data["original_amount"]), Decimal("1000000"))
        self.assertIn("schedules", data)
        self.assertIn("allocations", data)
        self.assertIn("nettings", data)
        self.assertIn("adjustments", data)

    def test_14_prepare_collection_success_steps_5_af07_br12(self):
        """
        Bước 5 / AF07 / BR12: Chọn các khoản nợ cần thu và gọi prepare-collection.
        Kiểm tra cùng khách hàng, loại tiền, tài khoản và trả về payload cho UC-AR-03.
        """
        prep_url = f"/api/v1/accounts-receivable/due-tracking/prepare-collection/?company_id={self.company.id}"
        payload = {
            "selected_items": [
                {"receivable_id": str(self.rec_overdue.id)},
                {"receivable_id": str(self.rec_due_today.id)},
            ]
        }
        res = self.client.post(prep_url, payload, format="json")
        self.assertEqual(res.status_code, 200, res.data)

        data = res.data
        self.assertEqual(data["target_use_case"], "UC-AR-03")
        self.assertEqual(data["customer_code"], "KH-ALPHA")
        self.assertEqual(data["currency_code"], "VND")
        self.assertEqual(data["selected_count"], 2)
        # Suggested amount = 1,000,000 + 500,000 = 1,500,000
        self.assertEqual(Decimal(data["suggested_total_amount"]), Decimal("1500000"))
        self.assertEqual(len(data["allocations"]), 2)

    def test_15_prepare_collection_heterogeneous_customer_currency_account_rejected_ef05(self):
        """
        EF05: Các khoản chọn thu tiền khác khách hàng, loại tiền hoặc tài khoản
        bị chặn không cho lập chung một chứng từ thu tiền.
        """
        prep_url = f"/api/v1/accounts-receivable/due-tracking/prepare-collection/?company_id={self.company.id}"

        # 1. Khác khách hàng & loại tiền (Customer 1 VND vs Customer 2 USD)
        payload = {
            "selected_items": [
                {"receivable_id": str(self.rec_overdue.id)},
                {"receivable_id": str(self.rec_usd_overdue.id)},
            ]
        }
        res = self.client.post(prep_url, payload, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("khách hàng", str(res.data))

    def test_16_prepare_collection_zero_balance_rejected_ef06(self):
        """
        EF06: Khoản công nợ đã được thanh toán hết (số dư = 0) bị chặn khi chuyển thu tiền.
        """
        # Set balance to 0 on rec_overdue
        self.rec_overdue.balance = Decimal("0")
        self.rec_overdue.save(update_fields=("balance",))

        prep_url = f"/api/v1/accounts-receivable/due-tracking/prepare-collection/?company_id={self.company.id}"
        payload = {
            "selected_items": [
                {"receivable_id": str(self.rec_overdue.id)},
            ]
        }
        res = self.client.post(prep_url, payload, format="json")
        self.assertEqual(res.status_code, 400)
        self.assertIn("thanh toán hết", str(res.data))

    def test_17_export_excel_endpoint_af05_br14(self):
        """
        AF05 / BR14 / EF08: Xuất Excel danh sách công nợ đến hạn kèm đầy đủ
        ngày theo dõi, bộ lọc, tổng hợp đa loại tiền và chi tiết từng dòng.
        """
        export_url = f"/api/v1/accounts-receivable/due-tracking/export/?company_id={self.company.id}&due_category=ALL"
        res = self.client.get(export_url)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        self.assertIn("bao-cao-cong-no-den-han", res["Content-Disposition"])

        # Validate OpenPyXL content
        wb = load_workbook(filename=BytesIO(res.content))
        ws = wb.active
        self.assertEqual(ws.title, "Cong No Den Han")
        # Check title row
        self.assertIn("BÁO CÁO THEO DÕI CÔNG NỢ ĐẾN HẠN", ws.cell(1, 1).value)

    def test_18_refresh_endpoint_af06(self):
        """
        AF06: Làm mới số liệu theo bộ lọc hiện tại.
        """
        refresh_url = f"/api/v1/accounts-receivable/due-tracking/refresh/?company_id={self.company.id}"
        res = self.client.post(refresh_url, {"due_category": "ALL"}, format="json")
        self.assertEqual(res.status_code, 200)
        self.assertIn("items", res.data)
        self.assertIn("data_extracted_at", res.data)

    def test_19_empty_result_displays_zeros_af08(self):
        """
        AF08: Không có công nợ phù hợp bộ lọc -> danh sách rỗng, các chỉ tiêu = 0.
        """
        url = f"{self.base_url}&search=NON_EXISTENT_QUERY"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data["items"]), 0)
        self.assertEqual(res.data["total_filtered_items"], 0)
        self.assertEqual(Decimal(res.data["base_summary"]["base_total_remaining_balance"]), Decimal("0"))

    def test_20_ef01_permission_checks(self):
        """
        EF01: Nhân viên không có quyền bị chặn 403 Forbidden.
        """
        unauthorized_user = get_user_model().objects.create_user(
            username="unauth_due", password="password"
        )
        self.client.force_authenticate(unauthorized_user)

        res = self.client.get(self.base_url)
        self.assertEqual(res.status_code, 403)

    def test_21_ef02_future_as_of_date_rejected(self):
        """
        EF02: Ngày theo dõi trong tương lai bị chặn với 400 ValidationError.
        """
        future_date = self.today + timedelta(days=2)
        url = f"{self.base_url}&as_of_date={future_date}"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 400)
        self.assertIn("as_of_date", str(res.data))

    def test_22_ef03_invalid_upcoming_days_or_overdue_range_rejected(self):
        """
        EF03: upcoming_days <= 0 hoặc khoảng overdue min > max bị chặn với 400 ValidationError.
        """
        # upcoming_days <= 0
        url_bad_upcoming = f"{self.base_url}&upcoming_days=0"
        res1 = self.client.get(url_bad_upcoming)
        self.assertEqual(res1.status_code, 400)
        self.assertIn("upcoming_days", str(res1.data))

        # overdue_days_min > overdue_days_max
        url_bad_range = f"{self.base_url}&overdue_days_min=30&overdue_days_max=10"
        res2 = self.client.get(url_bad_range)
        self.assertEqual(res2.status_code, 400)
        self.assertIn("overdue_days", str(res2.data))

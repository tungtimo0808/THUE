"""Staged, company-scoped catalogue imports. File values never choose the company."""
import csv
import hashlib
import io
import itertools
import json
from dataclasses import dataclass
from datetime import date

from django.conf import settings
from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import IntegrityError, transaction
from django.forms.models import model_to_dict
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework import serializers, status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from audit_log.services import record_audit
from iam.permissions import assignment_permission_codes, effective_assignments, has_effective_permission
from organization.models import Company, Employee
from organization.serializers import EmployeeSerializer
from general_ledger.models import Account
from general_ledger.serializers import AccountSerializer
from . import models
from . import serializers as md_serializers
from .payment_methods import PaymentMethodSerializer, PaymentMethodViewSet
from .payment_terms import PaymentTermSerializer
from .tax_catalog import TaxRateConfigSerializer, TaxTypeSerializer


@dataclass(frozen=True)
class Target:
    model: type
    serializer: type
    resource: str
    key: str = "code"
    role: str = ""
    forbidden: tuple = ()


TARGETS = {
    "customers": Target(models.Party, md_serializers.PartySerializer, "MD.PARTY", role="CUSTOMER"),
    "suppliers": Target(models.Party, md_serializers.PartySerializer, "MD.PARTY", role="VENDOR"),
    "employees": Target(Employee, EmployeeSerializer, "MD.EMPLOYEE", key="employee_code"),
    "item-categories": Target(models.ItemCategory, md_serializers.ItemCategorySerializer, "MD.ITEM_GROUP"),
    "units-of-measure": Target(models.UnitOfMeasure, md_serializers.UnitOfMeasureSerializer, "MD.UOM"),
    "items": Target(models.Item, md_serializers.ItemSerializer, "MD.ITEM"),
    "warehouses": Target(models.Warehouse, md_serializers.WarehouseSerializer, "MD.WAREHOUSE"),
    "company-bank-accounts": Target(models.CompanyBankAccount, md_serializers.CompanyBankAccountSerializer, "MD.BANK_ACCOUNT"),
    "accounts": Target(Account, AccountSerializer, "MD.ACCOUNT"),
    "cost-centers": Target(models.CostCenter, md_serializers.CostCenterSerializer, "MD.ORGANIZATION"),
    "expense-items": Target(models.ExpenseItem, md_serializers.ExpenseItemSerializer, "MD.EXPENSE_ITEM"),
    "statistical-codes": Target(models.StatisticalCode, md_serializers.StatisticalCodeSerializer, "MD.STATISTICAL_CODE"),
    "tax-types": Target(models.TaxType, TaxTypeSerializer, "MD.TAX_TYPE"),
    "tax-rates": Target(models.TaxRate, TaxRateConfigSerializer, "MD.TAX_RATE"),
    "payment-terms": Target(models.PaymentTerm, PaymentTermSerializer, "MD.PAYMENT_TERM", forbidden=("default_purchase", "default_sale")),
    "payment-methods": Target(models.PaymentMethod, PaymentMethodSerializer, "MD.PAYMENT_METHOD", forbidden=("default_receipt", "default_payment", "integration_channel")),
}

SYSTEM_FIELDS = {"id", "pk", "company", "created_by", "updated_by", "created_at", "updated_at", "version",
                 "source", "is_system", "connection_id", "connection_provider", "connection_status", "auth_expires_at",
                 "last_sync_at", "last_error", "merged_into", "merged_at", "merged_by", "lookup_at", "lookup_source"}
SECURE_FIELDS = {"identity_number", "personal_tax_code", "address", "date_of_birth"}


def limits():
    return int(getattr(settings, "CATALOG_IMPORT_MAX_BYTES", 5 * 1024 * 1024)), int(getattr(settings, "CATALOG_IMPORT_MAX_ROWS", 5000))


def catalogue_grant(account, company, target, action_code):
    code = f"{target.resource}.{action_code}"
    if has_effective_permission(account, company, code):
        return True
    if not any(field.name == "branch" for field in target.model._meta.fields):
        return False
    return any(code in assignment_permission_codes(assignment) for assignment in effective_assignments(account, company))


def fields_for(target):
    instance = target.serializer()
    return {name: field for name, field in instance.fields.items()
            if not field.read_only and name not in SYSTEM_FIELDS and name not in target.forbidden}


def key_value(target, data):
    return str(data.get(target.key, "")).strip().upper()


def identity_value(target, data):
    code = key_value(target, data)
    return f"{code}|{data.get('effective_from', '')}" if target.model is models.TaxRate else code


def serialize_error(exc):
    if isinstance(exc, (ValidationError, DjangoValidationError)):
        detail = exc.detail if isinstance(exc, ValidationError) else (exc.message_dict if hasattr(exc, "message_dict") else exc.messages)
        return json.loads(json.dumps(detail, default=str))
    return {"detail": str(exc)}


def parse_upload(upload):
    max_bytes, max_rows = limits()
    if upload.size > max_bytes:
        raise ValidationError({"file": f"File exceeds {max_bytes} bytes."})
    suffix = upload.name.rsplit(".", 1)[-1].lower()
    if suffix not in ("csv", "xlsx"):
        raise ValidationError({"file": "CSV or XLSX file is required."})
    raw = upload.read(max_bytes + 1)
    if len(raw) > max_bytes:
        raise ValidationError({"file": "File is too large."})
    try:
        if suffix == "csv":
            text = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text, newline=""), strict=True)
            headers = reader.fieldnames or []
            rows = list(itertools.islice(reader, max_rows + 1))
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.active
                iterator = sheet.iter_rows(values_only=True)
                headers = [str(value).strip() if value is not None else "" for value in next(iterator, ())]
                rows = [dict(zip(headers, ["" if value is None else value.isoformat() if isinstance(value, date) else str(value)
                                           for value in values])) for values in itertools.islice(iterator, max_rows + 1)]
            finally:
                workbook.close()
    except (UnicodeDecodeError, csv.Error, ValueError, OSError, KeyError, TypeError, ImportError) as exc:
        raise ValidationError({"file": "Unreadable file."}) from exc
    if not headers or not all(headers) or len(set(headers)) != len(headers) or not rows:
        raise ValidationError({"file": "A unique header and at least one data row are required."})
    if len(rows) > max_rows:
        raise ValidationError({"file": f"File exceeds {max_rows} rows."})
    if any(None in row or any(isinstance(value, str) and value.lstrip().startswith("=") for value in row.values()) for row in rows):
        raise ValidationError({"file": "Malformed rows or formulas are not allowed."})
    return headers, rows, raw


def version_of(obj):
    if obj is None:
        return ""
    fields = model_to_dict(obj)
    return hashlib.sha256(json.dumps(fields, sort_keys=True, default=str).encode()).hexdigest()


def order_rows(rows, target):
    """Parents that appear in the same file are evaluated first; cycles fail by row."""
    by_code = {key_value(target, row.mapped_data): row for row in rows if key_value(target, row.mapped_data)}
    output, visiting, visited, cycles = [], set(), set(), set()

    def visit(row):
        if row.pk in visited:
            return
        if row.pk in visiting:
            cycles.add(row.pk)
            return
        visiting.add(row.pk)
        parent_code = str(row.mapped_data.get("parent", "")).strip().upper()
        parent = by_code.get(parent_code)
        if parent:
            visit(parent)
            if parent.pk in cycles:
                cycles.add(row.pk)
        visiting.remove(row.pk)
        visited.add(row.pk)
        output.append(row)

    for row in rows:
        visit(row)
    return output, cycles


def resolve_references(data, target, company):
    converted = dict(data)
    for name, value in data.items():
        try:
            field = target.model._meta.get_field(name)
        except Exception:
            continue
        if not field.is_relation or not field.many_to_one or not value:
            continue
        related = field.remote_field.model
        if related is Company:
            raise ValidationError({name: "Company is selected by the session."})
        identifier = "employee_code" if related is Employee else "code"
        if not any(f.name == identifier for f in related._meta.fields):
            raise ValidationError({name: "Reference by business code is unavailable."})
        queryset = related.objects.filter(**{f"{identifier}__iexact": str(value).strip()})
        if any(f.name == "company" for f in related._meta.fields):
            queryset = queryset.filter(company=company)
        elif related is models.TaxType:
            queryset = queryset.filter(company=company)
        obj = queryset.first()
        if not obj or (hasattr(obj, "status") and obj.status != "ACTIVE") or (hasattr(obj, "is_active") and not obj.is_active):
            raise ValidationError({name: f"Active reference code {value} was not found in this company."})
        converted[name] = str(obj.pk)
    return converted


def process_rows(batch, request=None, *, commit=False):
    target = TARGETS[batch.target]
    rows = list(batch.rows.all())
    ordered, cycles = order_rows(rows, target)
    seen = set()
    results = {}
    with transaction.atomic():
        for row in ordered:
            code = key_value(target, row.mapped_data)
            identity = identity_value(target, row.mapped_data)
            if row.pk in cycles:
                results[row.pk] = {"action": "", "result": "ERROR", "errors": {"parent": "Parent cycle in file."}, "record_version": ""}
                continue
            if not code or identity in seen:
                results[row.pk] = {"action": "", "result": "ERROR", "errors": {target.key: "Missing or repeated code in file."}, "record_version": ""}
                continue
            seen.add(identity)
            lookup = {f"{target.key}__iexact": code}
            if target.model is models.TaxRate:
                try:
                    lookup["effective_from"] = parse_date(str(row.mapped_data.get("effective_from", "")))
                except ValueError:
                    lookup["effective_from"] = None
            existing = target.model.objects.filter(company=batch.company, **lookup).first() if target.model is not models.TaxRate or lookup["effective_from"] else None
            previous_version = version_of(existing)
            if existing and batch.duplicate_mode in ("ERROR", "CREATE_ONLY"):
                results[row.pk] = {"action": "", "result": "ERROR", "errors": {target.key: "Code already exists."}, "record_version": previous_version}
                continue
            if existing and batch.duplicate_mode == "SKIP":
                results[row.pk] = {"action": "SKIP", "result": "SKIPPED", "errors": {}, "record_version": previous_version}
                continue
            action_name = "UPDATE" if existing else "CREATE"
            try:
                with transaction.atomic():
                    if existing and "status" in row.mapped_data and row.mapped_data["status"] != existing.status:
                        raise ValidationError({"status": "Use the catalogue status action instead of import update."})
                    if not existing and target.model is models.Party and row.mapped_data.get("status", "ACTIVE") not in ("ACTIVE", "DRAFT", "INACTIVE"):
                        raise ValidationError({"status": "Unsupported initial party status."})
                    if existing and target.model is Account and "chart_of_accounts" in row.mapped_data and row.mapped_data["chart_of_accounts"].strip().upper() != existing.chart_of_accounts.code.upper():
                        raise ValidationError({"chart_of_accounts": "Move accounts through the chart management workflow."})
                    if existing and batch.target == "payment-methods" and row.mapped_data.get("group", existing.group) != existing.group and PaymentMethodViewSet()._used(existing):
                        raise ValidationError({"group": "Group cannot change after use."})
                    if existing and batch.target == "tax-rates" and existing.tax_type_ref.source != "COMPANY":
                        raise PermissionDenied("System tax rate cannot be edited.")
                    data = resolve_references(row.mapped_data, target, batch.company)
                    if not batch.created_by.auth_user.is_superuser:
                        branch_id = data.get("branch") or getattr(existing, "branch_id", None)
                        scope = {"branch_id": branch_id,
                                 "warehouse_id": existing.pk if target.model is models.Warehouse and existing else None,
                                 "bank_account_id": existing.pk if target.model is models.CompanyBankAccount and existing else None}
                        for permission in ("IMPORT", "EDIT") if existing and batch.duplicate_mode == "UPDATE" else ("IMPORT",):
                            if not has_effective_permission(batch.created_by, batch.company, f"{target.resource}.{permission}", **scope):
                                raise PermissionDenied("Row is outside the permitted data scope.")
                    data["company"] = str(batch.company_id)
                    if batch.target == "tax-types":
                        data["source"] = "COMPANY"
                    if existing:
                        data.pop(target.key, None)
                        data.pop("company", None)
                    changed_fields = set(row.mapped_data) - {target.key}
                    before_values = {name: str(getattr(existing, name)) for name in changed_fields} if existing else None
                    serializer = target.serializer(existing, data=data, partial=bool(existing), context={"request": request})
                    serializer.is_valid(raise_exception=True)
                    save_options = {"version": existing.version + 1} if existing and any(field.name == "version" for field in existing._meta.fields) else {}
                    obj = serializer.save(**save_options)
                    if target.role:
                        models.PartyRole.objects.get_or_create(party=obj, role_type=target.role, defaults={"active_from": timezone.localdate()})
                    if commit:
                        after_values = {name: str(getattr(obj, name)) for name in changed_fields}
                        if batch.target == "company-bank-accounts" and "account_number" in changed_fields:
                            if before_values:
                                before_values["account_number"] = "***" + before_values["account_number"][-4:]
                            after_values["account_number"] = "***" + after_values["account_number"][-4:]
                        record_audit(company=batch.company, actor_user=batch.created_by, module_code="MDM", action_code="IMPORT",
                                     entity=obj, before=before_values,
                                     after={target.key: code, **after_values, "batch_id": str(batch.pk), "source_row": row.source_row})
                results[row.pk] = {"action": action_name, "result": "UPDATED" if existing else "CREATED", "errors": {},
                                   "record_version": previous_version, "record_id": str(obj.pk) if commit else ""}
            except (ValidationError, PermissionDenied, DjangoValidationError, IntegrityError, ValueError, TypeError) as exc:
                results[row.pk] = {"action": action_name, "result": "ERROR", "errors": serialize_error(exc), "record_version": previous_version}
        if not commit or (batch.error_mode == "ATOMIC" and any(item["result"] == "ERROR" for item in results.values())):
            transaction.set_rollback(True)
    return results


def result_hash(rows, results):
    payload = [{"row": row.source_row, "data": row.mapped_data, **results[row.pk]} for row in rows]
    return hashlib.sha256(json.dumps(payload, sort_keys=True, default=str).encode()).hexdigest()


def execute_batch(batch, request=None):
    """Worker-safe execution. Caller owns the transaction and company lock."""
    if hashlib.sha256(batch.source_file).hexdigest() != batch.source_hash:
        raise ValidationError({"file": "Stored source file failed integrity check."})
    if batch.company.status != "ACTIVE":
        raise ValidationError({"company": "Company is inactive."})
    if not batch.created_by.auth_user.is_superuser:
        target = TARGETS[batch.target]
        if not catalogue_grant(batch.created_by, batch.company, target, "IMPORT"):
            raise PermissionDenied("Import permission was revoked.")
        if batch.duplicate_mode == "UPDATE" and not catalogue_grant(batch.created_by, batch.company, target, "EDIT"):
            raise PermissionDenied("Edit permission was revoked.")
    rows = list(batch.rows.all())
    current = process_rows(batch, request)
    if result_hash(rows, current) != batch.preview_hash:
        raise ValidationError({"preview": "Catalogue changed since preview. Run preview again."}, code="conflict")
    batch.status = "IMPORTING"
    batch.save(update_fields=("status",))
    results = process_rows(batch, request, commit=True)
    atomic_failed = batch.error_mode == "ATOMIC" and any(item["result"] == "ERROR" for item in results.values())
    for row in rows:
        outcome = results[row.pk]
        row.action, row.result, row.errors, row.record_id = (outcome.get(key, "") for key in ("action", "result", "errors", "record_id"))
        if atomic_failed and row.result in ("CREATED", "UPDATED"):
            row.result = "ROLLED_BACK"
            row.record_id = ""
        row.save(update_fields=("action", "result", "errors", "record_id"))
    batch.created_rows = 0 if atomic_failed else sum(item["result"] == "CREATED" for item in results.values())
    batch.updated_rows = 0 if atomic_failed else sum(item["result"] == "UPDATED" for item in results.values())
    batch.skipped_rows = sum(item["result"] == "SKIPPED" for item in results.values())
    batch.error_rows = sum(item["result"] == "ERROR" for item in results.values())
    batch.status = "FAILED" if atomic_failed else "PARTIAL" if batch.error_rows or batch.skipped_rows else "COMPLETED"
    batch.error_message = "All catalogue changes were rolled back." if atomic_failed else ""
    batch.finished_at = timezone.now()
    batch.save()
    record_audit(company=batch.company, actor_user=batch.created_by, module_code="MDM", action_code="IMPORT_CONFIRM",
                 entity=batch, after={"target": batch.target, "status": batch.status, "created_rows": batch.created_rows,
                                      "updated_rows": batch.updated_rows, "error_rows": batch.error_rows})
    return batch


class CatalogImportViewSet(viewsets.GenericViewSet):
    permission_classes = (IsAuthenticated,)
    queryset = models.CatalogImportBatch.objects.all()

    def _company(self):
        company_id = self.request.query_params.get("company_id")
        if not company_id:
            raise ValidationError({"company_id": "Working company is required."})
        try:
            return Company.objects.get(pk=company_id)
        except (Company.DoesNotExist, ValueError):
            raise ValidationError({"company_id": "Unknown company."})

    def _authorize(self, company, target, *, update=False, view=False):
        if target not in TARGETS:
            raise ValidationError({"target": "Unsupported catalogue."})
        if self.request.user.is_superuser:
            return
        account = getattr(self.request.user, "iam_account", None)
        config = TARGETS[target]
        action_code = "VIEW" if view else "IMPORT"
        if not account or not catalogue_grant(account, company, config, action_code):
            raise PermissionDenied("Catalogue permission is required.")
        if update and not catalogue_grant(account, company, config, "EDIT"):
            raise PermissionDenied("Edit permission is required for update mode.")

    def get_queryset(self):
        return self.queryset.filter(company=self._company())

    def _batch(self, pk, *, write=False):
        batch = self.get_queryset().filter(pk=pk).first()
        if not batch:
            raise ValidationError({"batch": "Batch was not found."})
        account = getattr(self.request.user, "iam_account", None)
        if not self.request.user.is_superuser and batch.created_by_id != getattr(account, "pk", None):
            raise PermissionDenied("Only the batch owner can view it.")
        self._authorize(batch.company, batch.target, update=write and batch.duplicate_mode == "UPDATE")
        return batch

    def _audit_batch(self, batch, action):
        record_audit(company=batch.company, actor_user=self.request.user.iam_account, module_code="MDM", action_code=action,
                     entity=batch, after={"target": batch.target, "status": batch.status, "total_rows": batch.total_rows,
                                          "created_rows": batch.created_rows, "updated_rows": batch.updated_rows, "error_rows": batch.error_rows}, request=self.request)

    def _summary(self, batch):
        return {"id": str(batch.pk), "target": batch.target, "file_name": batch.file_name, "template_version": batch.template_version,
                "duplicate_mode": batch.duplicate_mode, "error_mode": batch.error_mode, "status": batch.status,
                "total_rows": batch.total_rows, "valid_rows": batch.valid_rows, "error_rows": batch.error_rows,
                "created_rows": batch.created_rows, "updated_rows": batch.updated_rows, "skipped_rows": batch.skipped_rows,
                "started_at": batch.started_at, "finished_at": batch.finished_at, "error_message": batch.error_message}

    @action(detail=False, methods=("get",))
    def targets(self, request):
        company = self._company()
        visible = []
        for name, target in TARGETS.items():
            try:
                self._authorize(company, name)
            except PermissionDenied:
                continue
            visible.append({"target": name, "key": target.key, "fields": sorted(fields_for(target))})
        return Response(visible)

    @action(detail=False, methods=("get",))
    def template(self, request):
        company = self._company()
        target_name = request.query_params.get("target")
        self._authorize(company, target_name)
        target = TARGETS[target_name]
        columns = [target.key] + sorted(name for name in fields_for(target) if name != target.key)
        if request.query_params.get("file_format") == "xlsx":
            from openpyxl import Workbook
            workbook = Workbook()
            workbook.active.title = "Import"
            workbook.active.append(columns)
            output = io.BytesIO()
            workbook.save(output)
            response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"] = f'attachment; filename="{target_name}-v1.xlsx"'
            return response
        output = io.StringIO()
        csv.writer(output).writerow(columns)
        response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{target_name}-v1.csv"'
        return response

    @action(detail=False, methods=("post",))
    def upload(self, request):
        company = self._company()
        if company.status != "ACTIVE":
            raise ValidationError({"company": "Company is inactive."})
        target_name = request.data.get("target")
        mode = request.data.get("duplicate_mode", "ERROR")
        error_mode = request.data.get("error_mode", "ATOMIC")
        self._authorize(company, target_name, update=mode == "UPDATE")
        if mode not in ("ERROR", "SKIP", "UPDATE", "CREATE_ONLY") or error_mode not in ("ATOMIC", "PARTIAL"):
            raise ValidationError({"mode": "Invalid duplicate or error mode."})
        upload = request.FILES.get("file")
        if not upload:
            raise ValidationError({"file": "File is required."})
        headers, rows, raw = parse_upload(upload)
        allowed = fields_for(TARGETS[target_name])
        if set(headers) & SYSTEM_FIELDS or set(headers) & SECURE_FIELDS:
            raise ValidationError({"file": "System or sensitive fields are not importable."})
        if target_name == "company-bank-accounts" and any(name.startswith("connection_") for name in headers):
            raise ValidationError({"file": "Connection data is not importable."})
        with transaction.atomic():
            batch = models.CatalogImportBatch.objects.create(company=company, target=target_name, file_name=upload.name[:255],
                source_file=raw, source_hash=hashlib.sha256(raw).hexdigest(), duplicate_mode=mode, error_mode=error_mode,
                total_rows=len(rows), created_by=request.user.iam_account)
            models.CatalogImportRow.objects.bulk_create([
                models.CatalogImportRow(batch=batch, source_row=index, source_data=row)
                for index, row in enumerate(rows, 2)
            ])
            self._audit_batch(batch, "IMPORT_UPLOAD")
        return Response({**self._summary(batch), "headers": headers, "suggested_mapping": {name: name for name in headers if name in allowed}}, status=status.HTTP_201_CREATED)

    def list(self, request):
        company = self._company()
        batches = []
        for batch in self.get_queryset().order_by("-started_at")[:100]:
            if not request.user.is_superuser and batch.created_by_id != getattr(getattr(request.user, "iam_account", None), "pk", None):
                continue
            try:
                self._authorize(company, batch.target)
            except PermissionDenied:
                continue
            batches.append(self._summary(batch))
        return Response(batches)

    def retrieve(self, request, pk=None):
        batch = self._batch(pk)
        return Response({**self._summary(batch), "rows": [
            {"source_row": row.source_row, "action": row.action, "result": row.result, "errors": row.errors,
             "record_id": row.record_id, "mapped_data": row.mapped_data}
            for row in batch.rows.all()]})

    @action(detail=True, methods=("post",))
    def preview(self, request, pk=None):
        batch = self._batch(pk, write=True)
        if hashlib.sha256(batch.source_file).hexdigest() != batch.source_hash:
            raise ValidationError({"file": "Stored source file failed integrity check."})
        if batch.status not in ("UPLOADED", "ERROR", "READY"):
            raise ValidationError({"status": "Batch cannot be previewed now."})
        mapping = request.data.get("mapping", batch.column_mapping or {
            name: name for name in batch.rows.first().source_data if name in fields_for(TARGETS[batch.target])
        })
        if not isinstance(mapping, dict):
            raise ValidationError({"mapping": "Object required."})
        allowed = fields_for(TARGETS[batch.target])
        if any(value not in allowed for value in mapping.values()) or len(set(mapping.values())) != len(mapping):
            raise ValidationError({"mapping": "Unknown or repeated destination field."})
        if batch.target in ("employees", "customers", "suppliers") and any(value in SECURE_FIELDS for value in mapping.values()):
            raise ValidationError({"mapping": "Sensitive fields require the dedicated catalogue API."})
        rows = list(batch.rows.all())
        headers = set(rows[0].source_data)
        if any(source not in headers for source in mapping):
            raise ValidationError({"mapping": "Unknown source column."})
        if TARGETS[batch.target].key not in mapping.values():
            raise ValidationError({"mapping": "Business code must be mapped."})
        with transaction.atomic():
            batch.status = "VALIDATING"
            batch.column_mapping = mapping
            batch.save(update_fields=("status", "column_mapping"))
            for row in rows:
                row.mapped_data = {destination: row.source_data.get(source, "") for source, destination in mapping.items() if row.source_data.get(source, "") != ""}
                row.save(update_fields=("mapped_data",))
            results = process_rows(batch, request)
            for row in rows:
                outcome = results[row.pk]
                row.action, row.result, row.errors, row.record_version = (outcome[key] for key in ("action", "result", "errors", "record_version"))
                row.save(update_fields=("action", "result", "errors", "record_version"))
            batch.error_rows = sum(item["result"] == "ERROR" for item in results.values())
            batch.skipped_rows = sum(item["result"] == "SKIPPED" for item in results.values())
            batch.valid_rows = batch.total_rows - batch.error_rows - batch.skipped_rows
            batch.preview_hash = result_hash(rows, results)
            batch.status = "ERROR" if batch.error_rows else "READY"
            batch.save()
            self._audit_batch(batch, "IMPORT_PREVIEW")
        return self.retrieve(request, pk)

    @action(detail=True, methods=("post",))
    def confirm(self, request, pk=None):
        batch = self._batch(pk, write=True)
        if batch.status not in ("READY", "ERROR") or batch.status == "ERROR" and batch.error_mode != "PARTIAL":
            raise ValidationError({"status": "Preview must be ready, or partial mode must be selected."})
        with transaction.atomic():
            batch = models.CatalogImportBatch.objects.select_for_update().get(pk=batch.pk)
            if batch.status not in ("READY", "ERROR"):
                raise ValidationError({"status": "Batch was already processed."})
            Company.objects.select_for_update().get(pk=batch.company_id)
            rows = list(batch.rows.all())
            current = process_rows(batch, request)
            if result_hash(rows, current) != batch.preview_hash:
                raise ValidationError({"preview": "Catalogue changed since preview. Run preview again."}, code="conflict")
            threshold = int(getattr(settings, "CATALOG_IMPORT_BACKGROUND_ROWS", 1000))
            if batch.total_rows >= threshold:
                batch.status = "QUEUED"
                batch.save(update_fields=("status",))
                self._audit_batch(batch, "IMPORT_QUEUE")
                return Response(self._summary(batch), status=status.HTTP_202_ACCEPTED)
            execute_batch(batch, request)
        return self.retrieve(request, pk)

    @action(detail=True, methods=("post",))
    def cancel(self, request, pk=None):
        batch = self._batch(pk, write=True)
        if batch.status not in ("UPLOADED", "READY", "ERROR"):
            raise ValidationError({"status": "Batch can no longer be cancelled."})
        batch.status = "CANCELLED"
        batch.finished_at = timezone.now()
        batch.save(update_fields=("status", "finished_at"))
        self._audit_batch(batch, "IMPORT_CANCEL")
        return Response(self._summary(batch))

    @action(detail=True, methods=("get",), url_path="results.csv")
    def results_csv(self, request, pk=None):
        batch = self._batch(pk)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(("source_row", "action", "result", "errors", "record_id"))
        for row in batch.rows.all():
            writer.writerow((row.source_row, row.action, row.result, json.dumps(row.errors, ensure_ascii=False), row.record_id))
        response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="catalog-import-{batch.pk}.csv"'
        return response

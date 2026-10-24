from decimal import Decimal
from datetime import timedelta
from io import BytesIO
import json
from pathlib import Path
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient
from audit_log.models import AuditLog
from accounts_payable.models import Payable,PayableDocument,Payment,PayableAdjustment
from general_ledger.models import Account,AccountClass,AccountingPeriod,ChartOfAccounts,JournalEntry,JournalLine
from master_data.models import Bank,CompanyBankAccount,Currency,Party
from organization.models import Branch,Company
from .models import AccountProfile,BankTransaction,Reconciliation,StatementImport,StatementLine,BankReconciliationPeriod,BankReconciliationItem,BankVoucherReview
class BankApiTests(TestCase):
 def setUp(self):
  self.today=timezone.localdate();self.currency=Currency.objects.get(code="VND");self.company=Company.objects.create(code="BANKCO",name="Bank Co",legal_name="Bank Co Ltd",tax_code="0100444433",accounting_currency=self.currency,legal_reporting_currency=self.currency);self.branch=Branch.objects.create(company=self.company,code="HN",name="Hà Nội");cl=AccountClass.objects.create(company=self.company,code="ASSET",name="Tài sản",normal_balance="DEBIT");chart=ChartOfAccounts.objects.create(company=self.company,code="TT99",name="Hệ thống",effective_from=self.today,is_default=True);self.gl=Account.objects.create(company=self.company,chart_of_accounts=chart,account_class=cl,code="1121",name="Tiền gửi",account_type="ASSET",normal_balance="DEBIT");self.bank=Bank.objects.create(code="TCBTEST",name="Ngân hàng test");self.ba=CompanyBankAccount.objects.create(company=self.company,branch=self.branch,bank=self.bank,account_number="19000001",account_name="BANK CO",currency=self.currency,gl_account=self.gl);self.ba2=CompanyBankAccount.objects.create(company=self.company,branch=self.branch,bank=self.bank,account_number="19000002",account_name="BANK CO 2",currency=self.currency,gl_account=self.gl);self.maker=get_user_model().objects.create_superuser(username="bank-maker",email="maker@bank.test",password="x");self.approver=get_user_model().objects.create_superuser(username="bank-approver",email="approver@bank.test",password="x");self.client=APIClient();self.client.force_authenticate(self.maker);self.q=f"?company_id={self.company.id}";p=self.client.post(f"/api/v1/bank/accounts/{self.q}",{"account":str(self.ba.id),"company":str(self.company.id),"integration_enabled":True,"connector_code":"MOCK"},format="json");self.assertEqual(p.status_code,201,p.data);self.profile=p.data["id"]
 def test_bank_payment_draft_validates_and_keeps_balance(self):
  debit=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,code="6421",name="Chi phí",account_type="EXPENSE",normal_balance="DEBIT")
  payload={"company":str(self.company.id),"account":self.profile,"number":"UNC-DRAFT-01","transaction_type":"PAYMENT","business_type":"OTHER_PAYMENT","transaction_date":str(self.today),"value_date":str(self.today),"counterparty_name":"Ngân hàng","currency":str(self.currency.id),"amount":"100","exchange_rate":"1","description":"Phí ngân hàng","payment_details":{"posting_date":str(self.today),"bank_fee":True,"lines":[{"description":"Phí ngân hàng","debit_account":str(debit.id),"credit_account":str(self.gl.id),"amount":"100"}]}}
  path=f"/api/v1/bank/transactions/{self.q}&branch_id={self.branch.id}"
  self.assertEqual(self.client.post(f"/api/v1/bank/transactions/{self.q}",payload,format="json").status_code,400)
  saved=self.client.post(path,payload,format="json");self.assertEqual(saved.status_code,201,saved.data)
  self.assertEqual(saved.data["status"],"DRAFT");self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("0"))
  self.assertEqual(self.client.post(path,payload,format="json").status_code,400)
  AccountingPeriod.objects.create(company=self.company,code="OPEN",fiscal_year=self.today.year,
   period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="OPEN")
  account=AccountProfile.objects.get(pk=self.profile);account.balance=Decimal("1000");account.save(update_fields=("balance",))
  tx_id=saved.data["id"]
  unscoped=self.client.post(f"/api/v1/bank/transactions/{tx_id}/submit/{self.q}",{},format="json")
  self.assertEqual(unscoped.status_code,400)
  submitted=self.client.post(f"/api/v1/bank/transactions/{tx_id}/submit/{self.q}&branch_id={self.branch.id}",{},format="json")
  self.assertEqual(submitted.status_code,200,submitted.data)
  approved=self.client.post(f"/api/v1/bank/transactions/{tx_id}/approve/{self.q}&branch_id={self.branch.id}",{},format="json")
  self.assertEqual(approved.status_code,200,approved.data)
  posted=self.client.post(f"/api/v1/bank/transactions/{tx_id}/post/{self.q}&branch_id={self.branch.id}",{},format="json")
  self.assertEqual(posted.status_code,200,posted.data)
  self.assertEqual(posted.data["gl_status"],"POSTED")
  self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("900"))
  payload["number"]="UNC-DRAFT-02";payload["payment_details"]["lines"][0]["credit_account"]=str(debit.id)
  self.assertEqual(self.client.post(path,payload,format="json").status_code,400)
 def test_bank_payment_draft_rejects_locked_period(self):
  debit=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,code="6421",name="Chi phí",account_type="EXPENSE",normal_balance="DEBIT")
  AccountingPeriod.objects.create(company=self.company,code="LOCKED",fiscal_year=self.today.year,period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="CLOSED")
  payload={"company":str(self.company.id),"account":self.profile,"number":"UNC-LOCKED","transaction_type":"PAYMENT","business_type":"OTHER_PAYMENT","transaction_date":str(self.today),"counterparty_name":"Ngân hàng","currency":str(self.currency.id),"amount":"100","exchange_rate":"1","description":"Phí ngân hàng","payment_details":{"posting_date":str(self.today),"bank_fee":True,"lines":[{"description":"Phí ngân hàng","debit_account":str(debit.id),"credit_account":str(self.gl.id),"amount":"100"}]}}
  response=self.client.post(f"/api/v1/bank/transactions/{self.q}&branch_id={self.branch.id}",payload,format="json")
  self.assertEqual(response.status_code,400);self.assertIn("đã khóa",str(response.data))
 def test_supplier_invoice_payment_posts_bank_debt_and_discount_together(self):
  AccountingPeriod.objects.create(company=self.company,code="OPEN",fiscal_year=self.today.year,
   period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="OPEN")
  payable_account=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,
   code="3311",name="Phải trả",account_type="LIABILITY",normal_balance="CREDIT",
   dimension_requirements={"SUPPLIER":"REQUIRED"})
  Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,
   code="5151",name="Chiết khấu thanh toán",account_type="REVENUE",normal_balance="CREDIT")
  supplier=Party.objects.create(company=self.company,code="NCC01",name="Nhà cung cấp",status="ACTIVE")
  document=PayableDocument.objects.create(company=self.company,branch=self.branch,number="CN-01",
   document_date=self.today,posting_date=self.today,due_date=self.today,supplier=supplier,
   payable_account=payable_account,currency=self.currency,exchange_rate=1,subtotal=500,
   tax_amount=0,total_amount=500,status="POSTED",created_by=self.maker.iam_account)
  debt=Payable.objects.create(company=self.company,branch=self.branch,supplier=supplier,document=document,
   currency=self.currency,exchange_rate=1,original_amount=500,balance=500)
  account=AccountProfile.objects.get(pk=self.profile);account.balance=Decimal("1000");account.save(update_fields=("balance",))
  payload={"company":str(self.company.id),"account":self.profile,"number":"UNC-INVOICE-01",
   "transaction_type":"PAYMENT","business_type":"SUPPLIER_INVOICE","transaction_date":str(self.today),
   "party":str(supplier.id),"counterparty_name":supplier.name,"counterparty_account":"123456",
   "currency":str(self.currency.id),"amount":"400","exchange_rate":"1","description":"Trả theo hóa đơn",
   "payment_details":{"posting_date":str(self.today),"branch":str(self.branch.id),
    "beneficiary_name":supplier.name,"beneficiary_account":"123456","beneficiary_bank":"Ngân hàng test",
    "lines":[{"description":"Trả công nợ","debit_account":str(payable_account.id),
              "credit_account":str(self.gl.id),"amount":"400"}],
    "allocations":[{"payable":str(debt.id),"amount":"400","discount":"50"}]}}
  created=self.client.post(f"/api/v1/bank/transactions/{self.q}&branch_id={self.branch.id}",payload,format="json")
  self.assertEqual(created.status_code,201,created.data)
  self.assertEqual(Payable.objects.get(pk=debt.pk).balance,Decimal("500"))
  tx_id=created.data["id"]
  self.client.post(f"/api/v1/bank/transactions/{tx_id}/submit/{self.q}&branch_id={self.branch.id}",{},format="json")
  self.client.post(f"/api/v1/bank/transactions/{tx_id}/approve/{self.q}&branch_id={self.branch.id}",{},format="json")
  posted=self.client.post(f"/api/v1/bank/transactions/{tx_id}/post/{self.q}&branch_id={self.branch.id}",{},format="json")
  self.assertEqual(posted.status_code,200,posted.data)
  self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("600"))
  self.assertEqual(Payable.objects.get(pk=debt.pk).balance,Decimal("50"))
  self.assertEqual(Payment.objects.get(external_reference=str(tx_id)).allocated_amount,Decimal("400"))
  self.assertEqual(PayableAdjustment.objects.get(payable=debt).amount,Decimal("-50"))
 def create_posted_receipt(self):
  x=self.client.post(f"/api/v1/bank/transactions/{self.q}",{"company":str(self.company.id),"account":self.profile,"number":"BC-001","transaction_type":"RECEIPT","business_type":"CUSTOMER_RECEIPT","transaction_date":str(self.today),"counterparty_name":"Khách hàng","currency":str(self.currency.id),"amount":"1000000","exchange_rate":"1","description":"Thu tiền"},format="json");self.assertEqual(x.status_code,201,x.data);tid=x.data["id"];self.client.post(f"/api/v1/bank/transactions/{tid}/submit/{self.q}",{},format="json");self.client.force_authenticate(self.approver);a=self.client.post(f"/api/v1/bank/transactions/{tid}/approve/{self.q}",{"comment":"Duyệt"},format="json");self.assertEqual(a.data["status"],"APPROVED");self.client.force_authenticate(self.maker);p=self.client.post(f"/api/v1/bank/transactions/{tid}/post/{self.q}",{"bank_reference":"REF001"},format="json");self.assertEqual(p.status_code,200,p.data);gl=self.client.post(f"/api/v1/bank/transactions/{tid}/post-gl/{self.q}",{"gl_reference":"GL-BANK-001"},format="json");self.assertEqual(gl.status_code,200,gl.data);return tid
 def test_transaction_statement_reconciliation_and_gl(self):
  tid=self.create_posted_receipt();self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("1000000"));statement=self.client.post(f"/api/v1/bank/statements/{self.q}",{"company":str(self.company.id),"account":self.profile,"source":"FILE","source_reference":"STMT-001","period_from":str(self.today),"period_to":str(self.today),"lines":[{"external_id":"LINE-001","transaction_date":str(self.today),"direction":"CREDIT","amount":"1000000","currency":str(self.currency.id),"reference":"REF001","description":"Thu tiền"}]},format="json");self.assertEqual(statement.status_code,201,statement.data);line=statement.data["lines"][0]["id"];self.assertEqual(self.client.post(f"/api/v1/bank/reconciliations/match/{self.q}",{"statement_line":line,"transaction":tid,"matched_amount":"1000001"},format="json").status_code,400);matched=self.client.post(f"/api/v1/bank/reconciliations/match/{self.q}",{"statement_line":line,"transaction":tid,"matched_amount":"1000000"},format="json");self.assertEqual(matched.status_code,200,matched.data);self.assertEqual(matched.data["status"],"MATCHED");self.assertTrue(AuditLog.objects.filter(company=self.company,module_code="BANK").exists())
 def test_transfer_and_idempotent_integration_request(self):
  self.create_posted_receipt();p2=self.client.post(f"/api/v1/bank/accounts/{self.q}",{"account":str(self.ba2.id),"company":str(self.company.id)},format="json");self.assertEqual(p2.status_code,201,p2.data)
  AccountingPeriod.objects.create(company=self.company,code="OPEN",fiscal_year=self.today.year,period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="OPEN")
  payload={"company":str(self.company.id),"number":"DC-001","source_account":self.profile,"destination_account":p2.data["id"],"transfer_date":str(self.today),"posting_date":str(self.today),"source_amount":"300000","destination_amount":"300000","exchange_rate":"1","description":"Điều chuyển","details":{"lines":[{"description":"Điều chuyển","debit_account":str(self.gl.id),"credit_account":str(self.gl.id),"amount":"300000"}]}}
  t=self.client.post(f"/api/v1/bank/transfers/{self.q}",payload,format="json");self.assertEqual(t.status_code,201,t.data)
  self.assertEqual(AccountProfile.objects.get(pk=p2.data["id"]).balance,Decimal("0"))
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{self.q}",payload,format="json").status_code,400)
  transfer_id=t.data["id"];c=self.client.post(f"/api/v1/bank/transfers/{transfer_id}/confirm/{self.q}",{},format="json");self.assertEqual(c.status_code,200,c.data);self.assertEqual(c.data["status"],"CONFIRMED");self.assertEqual(AccountProfile.objects.get(pk=p2.data["id"]).balance,Decimal("300000"));self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("700000"))
  entry=JournalEntry.objects.get(source_event_id=f"BANK_TRANSFER:{transfer_id}");self.assertEqual(entry.status,"POSTED")
  lines=list(JournalLine.objects.filter(entry=entry).order_by("line_no"));self.assertEqual([line.debit for line in lines],[Decimal("300000"),Decimal("0")]);self.assertEqual([line.credit for line in lines],[Decimal("0"),Decimal("300000")]);self.assertEqual([line.dimensions["bank_account_id"] for line in lines],[str(self.ba2.id),str(self.ba.id)])
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{transfer_id}/confirm/{self.q}",{},format="json").status_code,400)
  ledger=self.client.get(f"/api/v1/bank/accounts/{p2.data['id']}/ledger/{self.q}");self.assertEqual(ledger.status_code,200);self.assertEqual(ledger.data["total_receipts"],Decimal("300000"))
  source_ledger=self.client.get(f"/api/v1/bank/accounts/{self.profile}/ledger/{self.q}");self.assertEqual(source_ledger.status_code,200);self.assertEqual(source_ledger.data["total_payments"],Decimal("300000"))
  tx=self.client.post(f"/api/v1/bank/transactions/{self.q}",{"company":str(self.company.id),"account":self.profile,"number":"PC-API","transaction_type":"PAYMENT","business_type":"VENDOR_PAYMENT","transaction_date":str(self.today),"counterparty_name":"Nhà cung cấp","currency":str(self.currency.id),"amount":"100000","exchange_rate":"1","description":"Chi API"},format="json");sent=self.client.post(f"/api/v1/bank/transactions/{tx.data['id']}/send-bank/{self.q}",{"idempotency_key":"IDEMP-001","payload":{"amount":100000}},format="json");self.assertEqual(sent.status_code,200,sent.data);duplicate=self.client.post(f"/api/v1/bank/transactions/{tx.data['id']}/send-bank/{self.q}",{"idempotency_key":"IDEMP-001"},format="json");self.assertEqual(duplicate.status_code,400)
 def test_transfer_rejects_invalid_accounts_and_rolls_back_failed_post(self):
  p2=self.client.post(f"/api/v1/bank/accounts/{self.q}",{"account":str(self.ba2.id),"company":str(self.company.id)},format="json");self.assertEqual(p2.status_code,201,p2.data)
  payload={"company":str(self.company.id),"number":"CK-VALIDATE","source_account":self.profile,"destination_account":p2.data["id"],"transfer_date":str(self.today),"posting_date":str(self.today),"source_amount":"100","destination_amount":"100","exchange_rate":"1","description":"Chuyển nội bộ","details":{"lines":[{"description":"Chuyển nội bộ","debit_account":str(self.gl.id),"credit_account":str(self.gl.id),"amount":"100"}]}}
  path=f"/api/v1/bank/transfers/{self.q}"
  self.assertEqual(self.client.post(path,{**payload,"destination_account":self.profile},format="json").status_code,400)
  self.assertEqual(self.client.post(path,{**payload,"destination_amount":"90"},format="json").status_code,400)
  self.assertEqual(self.client.post(path,{**payload,"details":{"lines":[{**payload["details"]["lines"][0],"amount":"0"}]}},format="json").status_code,400)
  self.assertEqual(self.client.post(path,{**payload,"details":{"lines":[{**payload["details"]["lines"][0],"amount":"invalid"}]}},format="json").status_code,400)
  saved=self.client.post(path,payload,format="json");self.assertEqual(saved.status_code,201,saved.data)
  AccountingPeriod.objects.create(company=self.company,code="OPEN",fiscal_year=self.today.year,period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="OPEN")
  failed=self.client.post(f"/api/v1/bank/transfers/{saved.data['id']}/confirm/{self.q}",{},format="json")
  self.assertEqual(failed.status_code,400);self.assertEqual(AccountProfile.objects.get(pk=p2.data["id"]).balance,Decimal("0"));self.assertFalse(JournalEntry.objects.filter(source_entity_id=saved.data["id"]).exists())
 def test_transfer_edit_cancel_and_atomic_unpost(self):
  p2=self.client.post(f"/api/v1/bank/accounts/{self.q}",{"account":str(self.ba2.id),"company":str(self.company.id)},format="json");self.assertEqual(p2.status_code,201,p2.data)
  AccountingPeriod.objects.create(company=self.company,code="OPEN",fiscal_year=self.today.year,period_number=self.today.month,start_date=self.today.replace(day=1),end_date=self.today,status="OPEN")
  source=AccountProfile.objects.get(pk=self.profile);source.balance=Decimal("1000");source.save(update_fields=("balance",))
  payload={"company":str(self.company.id),"number":"CK-EDIT-01","source_account":self.profile,"destination_account":p2.data["id"],"transfer_date":str(self.today),"posting_date":str(self.today),"source_amount":"100","destination_amount":"100","exchange_rate":"1","description":"Transfer test","details":{"lines":[{"description":"Transfer line","debit_account":str(self.gl.id),"credit_account":str(self.gl.id),"amount":"100"}]}}
  path=f"/api/v1/bank/transfers/{self.q}"
  created=self.client.post(path,payload,format="json");self.assertEqual(created.status_code,201,created.data)
  transfer_id=created.data["id"];detail=f"/api/v1/bank/transfers/{transfer_id}/{self.q}"
  edited=self.client.patch(detail,{"version":created.data["version"],"description":"Updated transfer"},format="json");self.assertEqual(edited.status_code,200,edited.data)
  self.assertEqual(edited.data["version"],created.data["version"]+1)
  self.assertEqual(self.client.patch(detail,{"version":created.data["version"],"description":"Stale"},format="json").status_code,400)
  posted=self.client.post(f"/api/v1/bank/transfers/{transfer_id}/confirm/{self.q}",{},format="json");self.assertEqual(posted.status_code,200,posted.data)
  self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("900"));self.assertEqual(AccountProfile.objects.get(pk=p2.data["id"]).balance,Decimal("100"))
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{transfer_id}/cancel/{self.q}",{},format="json").status_code,400)
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{transfer_id}/reverse/{self.q}",{},format="json").status_code,400)
  reversed_response=self.client.post(f"/api/v1/bank/transfers/{transfer_id}/reverse/{self.q}",{"reason":"Correction","period":str(AccountingPeriod.objects.get(company=self.company).id),"posting_date":str(self.today)},format="json")
  self.assertEqual(reversed_response.status_code,200,reversed_response.data);self.assertEqual(reversed_response.data["status"],"REVERSED")
  self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("1000"));self.assertEqual(AccountProfile.objects.get(pk=p2.data["id"]).balance,Decimal("0"))
  self.assertEqual(JournalEntry.objects.get(source_event_id=f"BANK_TRANSFER:{transfer_id}").status,"REVERSED")
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{transfer_id}/reverse/{self.q}",{"reason":"Again"},format="json").status_code,400)
  payload["number"]="CK-CANCEL-01"
  draft=self.client.post(path,payload,format="json");self.assertEqual(draft.status_code,201,draft.data)
  cancelled=self.client.post(f"/api/v1/bank/transfers/{draft.data['id']}/cancel/{self.q}",{},format="json");self.assertEqual(cancelled.status_code,200,cancelled.data)
  self.assertEqual(cancelled.data["status"],"CANCELLED")
  self.assertEqual(self.client.post(f"/api/v1/bank/transfers/{draft.data['id']}/confirm/{self.q}",{},format="json").status_code,400)
 def test_statement_preview_and_atomic_file_import(self):
  from openpyxl import Workbook
  book=Workbook();book.active.append(["Ngày","Nội dung","Thu","Chi","Mã"]);book.active.append([str(self.today),"Khách chuyển",100,0,"BANK-1"]);book.active.append([str(self.today),"Phí ngân hàng",0,20,"BANK-2"])
  output=BytesIO();book.save(output)
  payload={"company":str(self.company.id),"account":self.profile,"currency":str(self.currency.id),
   "period_from":str(self.today),"period_to":str(self.today),"sheet_name":"Sheet","header_row":1,
   "first_data_row":2,"date_format":"DMY","number_format":"VI","column_mapping":{"transaction_date":"0","description":"1","receipt":"2","payment":"3","reference":"4"},"unique_reference":True,
   "rows":[{"source_row":2,"transaction_date":str(self.today),"description":"Khách chuyển","direction":"CREDIT","amount":"100","reference":"BANK-1","raw":{"Mã":"BANK-1"}},
    {"source_row":3,"transaction_date":str(self.today),"description":"Phí ngân hàng","direction":"DEBIT","amount":"20","reference":"BANK-2"},
    {"source_row":4,"transaction_date":"","description":"Lỗi ngày","direction":"CREDIT","amount":"10","reference":"BANK-3"},
    {"source_row":5,"transaction_date":str(self.today),"description":"Khách chuyển","direction":"CREDIT","amount":"100","reference":"BANK-1"}]}
  preview=self.client.post(f"/api/v1/bank/statements/preview-import/{self.q}",payload,format="json")
  self.assertEqual(preview.status_code,200,preview.data);self.assertEqual(preview.data["counts"],{"VALID":2,"ERROR":1,"DUPLICATE":1,"SUSPECT":0})
  payload["selected_rows"]=[2,3]
  def upload():return SimpleUploadedFile("statement.xlsx",output.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
  payload["rows"][0]["amount"]="999"
  tampered=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(tampered.status_code,400);self.assertEqual(StatementImport.objects.filter(company=self.company).count(),0)
  payload["rows"][0]["amount"]="100"
  imported=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(imported.status_code,201,imported.data);self.assertEqual(StatementLine.objects.filter(statement_id=imported.data["id"]).count(),2)
  self.assertEqual(set(StatementLine.objects.filter(statement_id=imported.data["id"]).values_list("external_id",flat=True)),{"BANK-1","BANK-2"})
  stored=StatementImport.objects.get(pk=imported.data["id"]);self.assertTrue(stored.source_file.name)
  self.assertTrue(imported.data["source_file"].endswith(".xlsx"))
  with self.assertRaises(ValueError):stored.source_file.url
  source=self.client.get(f"/api/v1/bank/statements/{stored.pk}/source-file/{self.q}")
  self.assertEqual(source.status_code,200);self.assertEqual(b"".join(source.streaming_content),output.getvalue())
  self.assertEqual(AccountProfile.objects.get(pk=self.profile).balance,Decimal("0"))
  self.assertEqual(JournalEntry.objects.filter(company=self.company).count(),0)
  duplicate=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(duplicate.status_code,400);self.assertEqual(StatementImport.objects.filter(company=self.company).count(),1)
  payload["selected_rows"]=[2];payload["unique_reference"]=False;payload["confirmed_suspects"]=[2]
  overlap=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(overlap.status_code,400);self.assertEqual(StatementImport.objects.filter(company=self.company).count(),1)
 def test_statement_suspect_requires_explicit_confirmation(self):
  from openpyxl import Workbook
  book=Workbook();book.active.append(["Ngày","Nội dung","Thu","Chi"]);book.active.append([str(self.today),"Giống nhau",50,0]);book.active.append([str(self.today),"Giống nhau",50,0]);output=BytesIO();book.save(output)
  payload={"company":str(self.company.id),"account":self.profile,"currency":str(self.currency.id),"period_from":str(self.today),"period_to":str(self.today),"sheet_name":"Sheet","column_mapping":{"transaction_date":"0","description":"1","receipt":"2","payment":"3"},"rows":[
   {"source_row":2,"transaction_date":str(self.today),"description":"Giống nhau","direction":"CREDIT","amount":"50"},
   {"source_row":3,"transaction_date":str(self.today),"description":"Giống nhau","direction":"CREDIT","amount":"50"}],"selected_rows":[2,3]}
  preview=self.client.post(f"/api/v1/bank/statements/preview-import/{self.q}",payload,format="json")
  self.assertEqual([row["status"] for row in preview.data["rows"]],["VALID","SUSPECT"])
  def upload():return SimpleUploadedFile("statement.xlsx",output.getvalue())
  response=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(response.status_code,400);self.assertEqual(StatementImport.objects.filter(company=self.company).count(),0)
  payload["confirmed_suspects"]=[3]
  response=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(response.status_code,201,response.data);self.assertEqual(len(response.data["lines"]),2)
  repeated=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(repeated.status_code,400);self.assertEqual(StatementImport.objects.filter(company=self.company).count(),1)
 def test_statement_xls_source_is_verified(self):
  content=(Path(__file__).resolve().parent/"fixtures"/"uc_bank_04_sample.xls").read_bytes()
  payload={"company":str(self.company.id),"account":self.profile,"currency":str(self.currency.id),
   "period_from":"2026-10-21","period_to":"2026-10-21","sheet_name":"Statement",
   "header_row":1,"first_data_row":2,"date_format":"DMY","number_format":"US",
   "column_mapping":{"transaction_date":"0","description":"1","receipt":"2","payment":"3"},
   "rows":[{"source_row":2,"transaction_date":"2026-10-21","description":"XLS receipt","direction":"CREDIT","amount":"125"}],"selected_rows":[2]}
  def upload():return SimpleUploadedFile("sample.xls",content,content_type="application/vnd.ms-excel")
  saved=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(saved.status_code,201,saved.data)
  self.assertEqual(StatementLine.objects.get(statement_id=saved.data["id"]).amount,Decimal("125"))
  payload["rows"][0]["description"]="Forged"
  rejected=self.client.post(f"/api/v1/bank/statements/import-file/{self.q}",{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(rejected.status_code,400)
 def test_statement_balance_warning_uses_selected_rows(self):
  from openpyxl import Workbook
  book=Workbook();book.active.append(["Date","Description","Receipt","Payment"])
  book.active.append([str(self.today),"First receipt",100,0]);book.active.append([str(self.today),"Second receipt",50,0])
  output=BytesIO();book.save(output)
  payload={"company":str(self.company.id),"account":self.profile,"currency":str(self.currency.id),
   "period_from":str(self.today),"period_to":str(self.today),"sheet_name":"Sheet","header_row":1,"first_data_row":2,
   "column_mapping":{"transaction_date":"0","description":"1","receipt":"2","payment":"3"},
   "opening_balance":"0","closing_balance":"150","selected_rows":[2],
   "rows":[{"source_row":2,"transaction_date":str(self.today),"description":"First receipt","direction":"CREDIT","amount":"100"},
           {"source_row":3,"transaction_date":str(self.today),"description":"Second receipt","direction":"CREDIT","amount":"50"}]}
  def upload():return SimpleUploadedFile("statement.xlsx",output.getvalue())
  path=f"/api/v1/bank/statements/import-file/{self.q}"
  blocked=self.client.post(path,{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(blocked.status_code,400);self.assertEqual(StatementImport.objects.count(),0)
  payload["confirm_balance_warning"]=True
  saved=self.client.post(path,{"payload":json.dumps(payload),"source_file":upload()},format="multipart")
  self.assertEqual(saved.status_code,201,saved.data)
  self.assertEqual(saved.data["raw_file"]["balance_difference"],"50")
 def test_statement_line_management_preserves_source_and_audits_removal(self):
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="FILE",source_reference="IMP-MANAGE",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  statement.raw_file={"file_name":"sao-ke.xlsx","sheet_name":"Statement"};statement.save(update_fields=("raw_file",))
  line=StatementLine.objects.create(statement=statement,external_id="000123",transaction_date=self.today,direction="CREDIT",amount=Decimal("125"),currency=self.currency,reference="0009",description="Thu thử",raw_data={"source_row":2,"counterparty_account":"987654321","ending_balance":"125"})
  path=f"/api/v1/bank/statement-lines/{line.id}/{self.q}"
  listed=self.client.get(f"/api/v1/bank/statement-lines/{self.q}&account={self.profile}&search=000123")
  self.assertEqual(listed.status_code,200,listed.data);self.assertEqual(len(listed.data),1)
  self.assertEqual(listed.data[0]["remaining_amount"],"125.0000")
  self.assertEqual(listed.data[0]["sheet_name"],"Statement")
  self.assertEqual(listed.data[0]["counterparty_account"],"987654321")
  self.assertEqual(listed.data[0]["ending_balance"],"125")
  self.assertEqual(listed.data[0]["created_by_name"],self.maker.iam_account.display_name)
  self.assertEqual(len(self.client.get(f"/api/v1/bank/statement-lines/{self.q}&search=987654321").data),1)
  self.assertEqual(self.client.patch(path,{"amount":"999","updated_at":""},format="json").status_code,400)
  self.assertEqual(self.client.patch(path,{"tracking_status":"NEEDS_REVIEW","updated_at":""},format="json").status_code,400)
  saved=self.client.patch(path,{"processing_note":"Cần đối chiếu","tracking_status":"NEEDS_REVIEW","review_reason":"Thiếu tên người chuyển","updated_at":""},format="json")
  self.assertEqual(saved.status_code,200,saved.data);self.assertEqual(saved.data["amount"],"125.0000")
  self.assertEqual(self.client.patch(path,{"processing_note":"Ghi đè","updated_at":""},format="json").status_code,400)
  self.assertEqual(self.client.post(f"/api/v1/bank/statement-lines/{line.id}/remove/{self.q}",{},format="json").status_code,400)
  removed=self.client.post(f"/api/v1/bank/statement-lines/{line.id}/remove/{self.q}",{"reason":"Nhập nhầm dòng","updated_at":saved.data["updated_at"]},format="json")
  self.assertEqual(removed.status_code,200,removed.data)
  line.refresh_from_db();self.assertEqual(line.external_id,"000123");self.assertIsNotNone(line.deleted_at)
  self.assertEqual(len(self.client.get(f"/api/v1/bank/statement-lines/{self.q}").data),0)
  self.assertTrue(AuditLog.objects.filter(company=self.company,action_code="STATEMENT_LINE_DELETE").exists())
 def test_statement_line_export_and_source_delete_rules(self):
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="CONNECTOR",source_reference="SYNC-001",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  line=StatementLine.objects.create(statement=statement,external_id="S-001",transaction_date=self.today,direction="DEBIT",amount=Decimal("50"),currency=self.currency)
  url=f"/api/v1/bank/statement-lines/{line.id}/remove/{self.q}"
  self.assertEqual(self.client.post(url,{"reason":"Sai"},format="json").status_code,400)
  exported=self.client.get(f"/api/v1/bank/statement-lines/export/{self.q}&account={self.profile}")
  self.assertEqual(exported.status_code,200);self.assertEqual(exported.content[:2],b"PK")
  statement.source="FILE";statement.reconciliation_confirmed_at=timezone.now();statement.save(update_fields=("source","reconciliation_confirmed_at"))
  self.assertEqual(self.client.patch(f"/api/v1/bank/statement-lines/{line.id}/{self.q}",{"updated_at":"","processing_note":"Đổi"},format="json").status_code,400)
  self.assertEqual(self.client.post(url,{"reason":"Sai"},format="json").status_code,400)
 def test_many_to_many_statement_matching_and_unmatch(self):
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="FILE",source_reference="IMP-MATCH",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  first=StatementLine.objects.create(statement=statement,external_id="L1",transaction_date=self.today,direction="CREDIT",amount=Decimal("150"),currency=self.currency)
  second=StatementLine.objects.create(statement=statement,external_id="L2",transaction_date=self.today,direction="CREDIT",amount=Decimal("75"),currency=self.currency)
  def voucher(number):return BankTransaction.objects.create(company=self.company,account_id=self.profile,number=number,transaction_type="RECEIPT",business_type="CUSTOMER_RECEIPT",transaction_date=self.today,counterparty_name="Khách A",currency=self.currency,amount=Decimal("100"),exchange_rate=1,base_amount=Decimal("100"),description="Thu tiền",status="POSTED",gl_status="POSTED",created_by=self.maker.iam_account)
  one,two=voucher("BC-M1"),voucher("BC-M2")
  candidates=self.client.get(f"/api/v1/bank/reconciliations/candidates/{self.q}&statement_line={first.pk}")
  self.assertEqual(candidates.status_code,200,candidates.data);self.assertEqual(len(candidates.data["candidates"]),2)
  batch={"statement_line":str(first.pk),"expected_line_remaining":"150","allocations":[{"transaction":str(one.pk),"matched_amount":"100","expected_transaction_remaining":"100"},{"transaction":str(two.pk),"matched_amount":"50","expected_transaction_remaining":"100"}]}
  matched=self.client.post(f"/api/v1/bank/reconciliations/match-batch/{self.q}",batch,format="json")
  self.assertEqual(matched.status_code,200,matched.data);self.assertEqual(matched.data["status"],"MATCHED");self.assertEqual(matched.data["remaining_amount"],"0.0000")
  first.refresh_from_db();self.assertEqual(first.match_status,"MATCHED")
  stale=self.client.post(f"/api/v1/bank/reconciliations/match-batch/{self.q}",{"statement_line":str(second.pk),"expected_line_remaining":"75","allocations":[{"transaction":str(two.pk),"matched_amount":"60","expected_transaction_remaining":"100"}]},format="json")
  self.assertEqual(stale.status_code,400);self.assertEqual(Reconciliation.objects.filter(cancelled_at__isnull=True).count(),2)
  partial=self.client.post(f"/api/v1/bank/reconciliations/match/{self.q}",{"statement_line":str(second.pk),"transaction":str(two.pk),"matched_amount":"50"},format="json")
  self.assertEqual(partial.status_code,200,partial.data);second.refresh_from_db();self.assertEqual(second.match_status,"PARTIALLY_MATCHED")
  detail=self.client.get(f"/api/v1/bank/statement-lines/{second.pk}/{self.q}")
  self.assertEqual(detail.data["matched_amount"],"50.0000");self.assertEqual(detail.data["remaining_amount"],"25.0000")
  self.assertEqual(len(detail.data["linked_vouchers"]),1)
  self.assertEqual(self.client.post(f"/api/v1/bank/reconciliations/match/{self.q}",{"statement_line":str(second.pk),"transaction":str(two.pk),"matched_amount":"1"},format="json").status_code,400)
  removed=self.client.post(f"/api/v1/bank/reconciliations/{partial.data['id']}/unmatch/{self.q}",{},format="json")
  self.assertEqual(removed.status_code,200,removed.data);second.refresh_from_db();self.assertEqual(second.match_status,"UNMATCHED")
  self.assertEqual(Reconciliation.objects.filter(cancelled_at__isnull=True).count(),2)
  self.assertTrue(AuditLog.objects.filter(company=self.company,action_code="UNMATCH").exists())
  deleted=self.client.post(f"/api/v1/bank/statement-lines/{second.pk}/remove/{self.q}",{"reason":"Giao dịch nhập nhầm sau khi hủy ghép","updated_at":detail.data["updated_at"]},format="json")
  self.assertEqual(deleted.status_code,200,deleted.data)
  statement.reconciliation_confirmed_at=timezone.now();statement.save(update_fields=("reconciliation_confirmed_at",))
  self.assertEqual(self.client.post(f"/api/v1/bank/reconciliations/{matched.data['links'][0]['id']}/unmatch/{self.q}",{},format="json").status_code,400)
 def test_statement_match_rejects_unposted_wrong_direction_and_currency(self):
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="FILE",source_reference="IMP-INVALID",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  line=StatementLine.objects.create(statement=statement,external_id="L-INVALID",transaction_date=self.today,direction="CREDIT",amount=Decimal("100"),currency=self.currency)
  voucher=BankTransaction.objects.create(company=self.company,account_id=self.profile,number="BC-INVALID",transaction_type="RECEIPT",business_type="CUSTOMER_RECEIPT",transaction_date=self.today,counterparty_name="Khách A",currency=self.currency,amount=Decimal("100"),exchange_rate=1,base_amount=Decimal("100"),description="Thu tiền",status="DRAFT",gl_status="PENDING",created_by=self.maker.iam_account)
  url=f"/api/v1/bank/reconciliations/match/{self.q}"
  payload={"statement_line":str(line.pk),"transaction":str(voucher.pk),"matched_amount":"100"}
  self.assertEqual(self.client.post(url,payload,format="json").status_code,400)
  voucher.status="POSTED";voucher.save(update_fields=("status",))
  self.assertEqual(self.client.post(url,payload,format="json").status_code,400)
  voucher.gl_status="POSTED";voucher.transaction_type="PAYMENT";voucher.save(update_fields=("gl_status","transaction_type"))
  self.assertEqual(self.client.post(url,payload,format="json").status_code,400)
  usd,_=Currency.objects.get_or_create(code="USD",defaults={"name":"US Dollar","numeric_code":"840","decimal_places":2})
  voucher.transaction_type="RECEIPT";voucher.currency=usd;voucher.save(update_fields=("transaction_type","currency"))
  self.assertEqual(self.client.post(url,payload,format="json").status_code,400)
  voucher.currency=self.currency;voucher.save(update_fields=("currency",))
  for invalid in ("0","-1","abc","101"):
   self.assertEqual(self.client.post(url,{**payload,"matched_amount":invalid},format="json").status_code,400)
  self.assertFalse(Reconciliation.objects.exists())
 def test_offline_period_preview_and_confirmation(self):
  yesterday=self.today-timedelta(days=1)
  def voucher(number,date,kind,value,posted=True):
   return BankTransaction.objects.create(company=self.company,account_id=self.profile,number=number,transaction_type=kind,business_type="CUSTOMER_RECEIPT" if kind=="RECEIPT" else "OTHER_PAYMENT",transaction_date=date,counterparty_name="Khách A",currency=self.currency,amount=Decimal(value),exchange_rate=1,base_amount=Decimal(value),description="Đối chiếu",status="POSTED",gl_status="POSTED" if posted else "PENDING",balance_applied=True,created_by=self.maker.iam_account)
  voucher("OLD",yesterday,"RECEIPT","40")
  receipt=voucher("R1",self.today,"RECEIPT","100")
  payment=voucher("P1",self.today,"PAYMENT","25")
  pending=voucher("PENDING",self.today,"RECEIPT","10",False)
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="FILE",source_reference="PERIOD-TEST",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  for index,(item,direction,value) in enumerate(((receipt,"CREDIT","100"),(payment,"DEBIT","25")),1):
   line=StatementLine.objects.create(statement=statement,external_id=f"LINE-{index}",transaction_date=self.today,direction=direction,amount=Decimal(value),currency=self.currency,match_status="MATCHED")
   Reconciliation.objects.create(company=self.company,statement_line=line,transaction=item,matched_amount=Decimal(value),status="MATCHED",created_by=self.maker.iam_account)
  profile=AccountProfile.objects.get(pk=self.profile);profile.balance=Decimal("125");profile.save(update_fields=("balance",))
  path=f"/api/v1/bank/reconciliation-periods/preview/{self.q}&account={self.profile}&date_from={self.today}&date_to={self.today}"
  preview=self.client.get(path)
  self.assertEqual(preview.status_code,200,preview.data)
  self.assertEqual(preview.data["opening_balance"],"40.0000")
  self.assertEqual({row["number"] for row in preview.data["transactions"]},{"R1","P1"})
  confirm=f"/api/v1/bank/reconciliation-periods/confirm/{self.q}"
  payload={"account":self.profile,"date_from":str(self.today),"date_to":str(self.today),"transactions":[str(receipt.pk),str(payment.pk)],"bank_closing_balance":"115"}
  self.assertEqual(self.client.post(confirm,{**payload,"bank_closing_balance":"114"},format="json").status_code,400)
  self.assertEqual(self.client.post(confirm,{**payload,"transactions":[str(pending.pk)]},format="json").status_code,400)
  result=self.client.post(confirm,payload,format="json")
  self.assertEqual(result.status_code,201,result.data)
  self.assertEqual(result.data["opening_balance"],Decimal("40"))
  self.assertEqual(result.data["closing_balance"],Decimal("115"))
  self.assertEqual(BankReconciliationPeriod.objects.count(),1)
  self.assertEqual(BankReconciliationItem.objects.count(),2)
  self.assertEqual(self.client.post(confirm,payload,format="json").status_code,400)
  self.assertTrue(all(row["confirmed"] for row in self.client.get(path).data["transactions"]))
  linked=Reconciliation.objects.filter(transaction=receipt).first()
  self.assertEqual(self.client.post(f"/api/v1/bank/reconciliations/{linked.pk}/unmatch/{self.q}",{},format="json").status_code,400)
  self.assertEqual(self.client.patch(f"/api/v1/bank/statement-lines/{linked.statement_line_id}/{self.q}",{"updated_at":"","processing_note":"Đổi sau xác nhận"},format="json").status_code,400)
  self.assertEqual(self.client.post(f"/api/v1/bank/transactions/{receipt.pk}/reverse/{self.q}",{"reason":"Đảo sau đối chiếu"},format="json").status_code,400)
  self.assertTrue(AuditLog.objects.filter(company=self.company,action_code="RECONCILIATION_CONFIRM").exists())
 def test_unmatched_list_waiting_and_resolution(self):
  statement=StatementImport.objects.create(company=self.company,account_id=self.profile,source="FILE",source_reference="UNMATCHED-TEST",period_from=self.today,period_to=self.today,created_by=self.maker.iam_account)
  line=StatementLine.objects.create(statement=statement,external_id="BANK-001",transaction_date=self.today,direction="CREDIT",amount=Decimal("100"),currency=self.currency,reference="REF-001",description="Thu tiền")
  voucher=BankTransaction.objects.create(company=self.company,account_id=self.profile,number="BC-UNMATCHED",transaction_type="RECEIPT",business_type="CUSTOMER_RECEIPT",transaction_date=self.today,counterparty_name="Khách A",currency=self.currency,amount=Decimal("100"),exchange_rate=1,base_amount=Decimal("100"),description="Thu tiền",bank_reference="REF-001",status="POSTED",gl_status="POSTED",balance_applied=True,created_by=self.maker.iam_account)
  path=f"/api/v1/bank/unmatched/items/{self.q}&account={self.profile}&date_from={self.today}&date_to={self.today}"
  listed=self.client.get(path)
  self.assertEqual(listed.status_code,200,listed.data)
  self.assertEqual({row["kind"] for row in listed.data["items"]},{"STATEMENT","VOUCHER"})
  self.assertEqual(self.client.post(f"/api/v1/bank/reconciliation-periods/confirm/{self.q}",{"account":self.profile,"date_from":str(self.today),"date_to":str(self.today),"transactions":[str(voucher.pk)],"bank_closing_balance":"100"},format="json").status_code,400)
  wait=f"/api/v1/bank/unmatched/mark-waiting/{self.q}"
  self.assertEqual(self.client.post(wait,{"kind":"STATEMENT","id":str(line.pk)},format="json").status_code,400)
  self.assertEqual(self.client.post(wait,{"kind":"STATEMENT","id":str(line.pk),"reason":"Chờ xác minh","note":"Liên hệ ngân hàng"},format="json").status_code,200)
  self.assertEqual(self.client.post(wait,{"kind":"VOUCHER","id":str(voucher.pk),"reason":"Chờ sao kê"},format="json").status_code,200)
  line.refresh_from_db();self.assertEqual(line.tracking_status,"WAITING")
  self.assertTrue(BankVoucherReview.objects.filter(transaction=voucher,reason="Chờ sao kê").exists())
  self.assertEqual(self.client.get(path).data["items"][0]["processing_status"],"WAITING")
  matched=self.client.post(f"/api/v1/bank/reconciliations/match/{self.q}",{"statement_line":str(line.pk),"transaction":str(voucher.pk),"matched_amount":"100"},format="json")
  self.assertEqual(matched.status_code,200,matched.data)
  line.refresh_from_db();self.assertEqual(line.tracking_status,"NORMAL")
  self.assertEqual(self.client.get(path).data["items"],[])
  self.assertTrue(AuditLog.objects.filter(company=self.company,action_code="UNMATCHED_WAIT").exists())
 def test_balance_report_uses_posted_gl_per_bank_account_and_exports(self):
  parent=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,account_class=self.gl.account_class,code="112",name="Tiền gửi",account_type="ASSET",normal_balance="DEBIT",is_postable=False)
  self.gl.parent=parent;self.gl.save(update_fields=("parent",))
  second=AccountProfile.objects.create(company=self.company,account=self.ba2)
  counterpart=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,account_class=self.gl.account_class,code="411",name="Vốn",account_type="EQUITY",normal_balance="CREDIT")
  yesterday=self.today-timedelta(days=1)
  period=AccountingPeriod.objects.create(company=self.company,code="REPORT",fiscal_year=self.today.year,period_number=self.today.month,start_date=yesterday,end_date=self.today+timedelta(days=1))
  def journal(number,date,bank_account,debit,credit,status="POSTED",entry_type="MANUAL"):
   entry=JournalEntry.objects.create(company=self.company,period=period,number=number,posting_date=date,description=number,status=status,entry_type=entry_type,created_by=self.maker.iam_account,total_debit=max(debit,credit),total_credit=max(debit,credit))
   JournalLine.objects.create(entry=entry,line_no=1,account=self.gl,debit=Decimal(debit),credit=Decimal(credit),currency=self.currency,foreign_amount=Decimal(max(debit,credit)),dimensions={"bank_account_id":str(bank_account)})
   JournalLine.objects.create(entry=entry,line_no=2,account=counterpart,debit=Decimal(credit),credit=Decimal(debit),currency=self.currency,foreign_amount=Decimal(max(debit,credit)))
  journal("OPEN-R",yesterday,self.ba.id,"100","0",entry_type="OPENING")
  journal("RECEIPT-R",self.today,self.ba.id,"50","0")
  journal("PAYMENT-R",self.today,self.ba2.id,"0","20")
  journal("TRANSFER-OUT",self.today,self.ba.id,"0","10")
  journal("TRANSFER-IN",self.today,self.ba2.id,"10","0")
  journal("DRAFT-R",self.today,self.ba.id,"999","0",status="DRAFT")
  journal("FUTURE-R",self.today+timedelta(days=1),self.ba.id,"777","0")
  ambiguous=JournalEntry.objects.create(company=self.company,period=period,number="AMBIGUOUS-R",posting_date=self.today,description="Không rõ tài khoản ngân hàng",status="POSTED",created_by=self.maker.iam_account,total_debit=Decimal("30"),total_credit=Decimal("30"))
  JournalLine.objects.create(entry=ambiguous,line_no=1,account=self.gl,debit=Decimal("30"),currency=self.currency,foreign_amount=Decimal("30"))
  JournalLine.objects.create(entry=ambiguous,line_no=2,account=counterpart,credit=Decimal("30"),currency=self.currency,foreign_amount=Decimal("30"))
  usd,_=Currency.objects.get_or_create(code="USD",defaults={"name":"US Dollar","numeric_code":"840","decimal_places":2})
  gl_usd=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,account_class=self.gl.account_class,parent=parent,code="1122",name="Tiền gửi USD",account_type="ASSET",normal_balance="DEBIT")
  ba_usd=CompanyBankAccount.objects.create(company=self.company,branch=self.branch,bank=self.bank,code="USDREPORT",account_number="19000003",account_name="BANK USD",currency=usd,gl_account=gl_usd)
  AccountProfile.objects.create(company=self.company,account=ba_usd,balance=Decimal("150000"))
  usd_entry=JournalEntry.objects.create(company=self.company,period=period,number="USD-R",posting_date=self.today,description="Thu USD",status="POSTED",created_by=self.maker.iam_account,total_debit=Decimal("2500000"),total_credit=Decimal("2500000"))
  JournalLine.objects.create(entry=usd_entry,line_no=1,account=gl_usd,debit=Decimal("2500000"),currency=usd,foreign_amount=Decimal("100"),exchange_rate=Decimal("25000"),dimensions={"bank_account_id":str(ba_usd.id)})
  JournalLine.objects.create(entry=usd_entry,line_no=2,account=counterpart,credit=Decimal("2500000"),currency=usd,foreign_amount=Decimal("100"),exchange_rate=Decimal("25000"))
  path=f"/api/v1/bank/balance-report/summary/{self.q}&account_code=112&date_from={self.today}&date_to={self.today}"
  result=self.client.get(path)
  self.assertEqual(result.status_code,200,result.data)
  rows={row["bank_account_no"]:row for row in result.data["rows"]}
  self.assertEqual(Decimal(rows["19000001"]["opening_balance"]),Decimal("100"))
  self.assertEqual(Decimal(rows["19000001"]["debit_amount"]),Decimal("50"))
  self.assertEqual(Decimal(rows["19000001"]["credit_amount"]),Decimal("10"))
  self.assertEqual(Decimal(rows["19000001"]["closing_balance"]),Decimal("140"))
  self.assertEqual(Decimal(rows["19000002"]["closing_balance"]),Decimal("-10"))
  self.assertEqual(Decimal(rows["19000003"]["closing_balance"]),Decimal("150100"))
  totals={row["currency"]:Decimal(row["closing_balance"]) for row in result.data["totals"]}
  self.assertEqual(totals,{"VND":Decimal("130"),"USD":Decimal("150100")})
  self.assertTrue(any("chưa xác định được tài khoản ngân hàng" in warning for warning in result.data["warnings"]))
  filtered=self.client.get(path+"&search=19000001")
  self.assertEqual(len(filtered.data["rows"]),1)
  self.assertEqual(Decimal(filtered.data["totals"][0]["closing_balance"]),Decimal("140"))
  book=self.client.get(path.replace("/summary/","/export/")+"&search=19000001")
  self.assertEqual(book.status_code,200);self.assertEqual(book.content[:2],b"PK")
  self.assertEqual(self.client.get(path.replace("account_code=112","account_code=111")).status_code,400)
  self.assertEqual(self.client.get(path.replace(f"company_id={self.company.id}","company_id=00000000-0000-4000-8000-000000000000")).status_code,400)
  saved=self.client.post(path.replace("/summary/","/save/"),{"name":"Số dư tháng","period_preset":"Tháng này"},format="json")
  self.assertEqual(saved.status_code,201,saved.data)
  presets=self.client.get(f"/api/v1/bank/balance-report/presets/{self.q}")
  self.assertEqual(len(presets.data),1)
  self.client.force_authenticate(self.approver)
  self.assertEqual(self.client.get(f"/api/v1/bank/balance-report/presets/{self.q}").data,[])
 def test_deposit_ledger_running_balance_filter_sort_scope_and_export(self):
  second=AccountProfile.objects.create(company=self.company,account=self.ba2)
  counterpart=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,code="411",name="Vốn",account_type="EQUITY",normal_balance="CREDIT")
  yesterday=self.today-timedelta(days=1)
  period=AccountingPeriod.objects.create(company=self.company,code="LEDGER",fiscal_year=self.today.year,
   period_number=self.today.month,start_date=yesterday,end_date=self.today)
  def entry(number,date,bank,debit,credit,status="POSTED",entry_type="MANUAL"):
   item=JournalEntry.objects.create(company=self.company,period=period,number=number,posting_date=date,
    description=number,status=status,entry_type=entry_type,created_by=self.maker.iam_account,
    total_debit=max(Decimal(debit),Decimal(credit)),total_credit=max(Decimal(debit),Decimal(credit)))
   JournalLine.objects.create(entry=item,line_no=1,account=self.gl,debit=debit,credit=credit,
    currency=self.currency,foreign_amount=max(Decimal(debit),Decimal(credit)),dimensions={"bank_account_id":str(bank)})
   JournalLine.objects.create(entry=item,line_no=2,account=counterpart,debit=credit,credit=debit,
    currency=self.currency,foreign_amount=max(Decimal(debit),Decimal(credit)))
  entry("OPEN",yesterday,self.ba.id,"100","0",entry_type="OPENING")
  entry("PAY",self.today,self.ba.id,"0","70")
  entry("RECEIVE",self.today,self.ba.id,"20","0")
  entry("SECOND",self.today,self.ba2.id,"5","0")
  entry("DRAFT",self.today,self.ba.id,"999","0",status="DRAFT")
  path=f"/api/v1/bank/deposit-ledger/summary/{self.q}&account_code=1121&date_from={self.today}&date_to={self.today}"
  result=self.client.get(path)
  self.assertEqual(result.status_code,200,result.data)
  groups={row["bank_account_no"]:row for row in result.data["groups"]}
  first=groups["19000001"]
  self.assertEqual(Decimal(first["opening_balance"]),Decimal("100"))
  self.assertEqual(Decimal(first["receipt_total"]),Decimal("20"))
  self.assertEqual(Decimal(first["payment_total"]),Decimal("70"))
  self.assertEqual(Decimal(first["closing_balance"]),Decimal("50"))
  self.assertEqual([Decimal(row["balance"]) for row in first["rows"]],[Decimal("30"),Decimal("50")])
  voucher=self.client.get(f"/api/v1/bank/deposit-ledger/voucher/{first['rows'][0]['journal_entry_id']}/{self.q}&account_code=1121&date_from={self.today}&date_to={self.today}")
  self.assertEqual(voucher.status_code,200,voucher.data)
  self.assertEqual([row["account_code"] for row in voucher.data["lines"]],["1121","411"])
  sorted_result=self.client.get(path+"&receipt_first=true").data["groups"][0]
  self.assertEqual([Decimal(row["balance"]) for row in sorted_result["rows"]],[Decimal("120"),Decimal("50")])
  searched=self.client.get(path+"&search=PAY").data["groups"][0]
  self.assertEqual(len(searched["rows"]),1)
  self.assertEqual(Decimal(searched["filtered_payment_total"]),Decimal("70"))
  self.assertEqual(Decimal(searched["closing_balance"]),Decimal("50"))
  scoped=self.client.get(path+f"&bank_account_id={self.ba2.id}").data["groups"]
  self.assertEqual([row["bank_account_no"] for row in scoped],["19000002"])
  self.assertEqual(self.client.get(path+"&bank_account_id=00000000-0000-4000-8000-000000000000").status_code,400)
  exported=self.client.get(path.replace("/summary/","/export/"))
  self.assertEqual(exported.status_code,200);self.assertEqual(exported.content[:2],b"PK")
  self.assertEqual(self.client.post(path.replace("/summary/","/email/"),{"to":"invalid","subject":"Sổ tháng"},format="json").status_code,400)
  self.assertEqual(self.client.post(path.replace("/summary/","/email/"),{"to":"recipient@example.com","subject":"Sổ tháng"},format="json").status_code,400)
  usd,_=Currency.objects.get_or_create(code="USD",defaults={"name":"US Dollar","numeric_code":"840","decimal_places":2})
  usd_gl=Account.objects.create(company=self.company,chart_of_accounts=self.gl.chart_of_accounts,code="1122",name="Tiền gửi USD",account_type="ASSET",normal_balance="DEBIT")
  usd_bank=CompanyBankAccount.objects.create(company=self.company,branch=self.branch,bank=self.bank,account_number="19000003",account_name="USD",currency=usd,gl_account=usd_gl)
  AccountProfile.objects.create(company=self.company,account=usd_bank,balance=Decimal("150"))
  usd_entry=JournalEntry.objects.create(company=self.company,period=period,number="USD-REC",posting_date=self.today,
   description="Thu USD",status="POSTED",created_by=self.maker.iam_account,total_debit=Decimal("2500000"),total_credit=Decimal("2500000"))
  JournalLine.objects.create(entry=usd_entry,line_no=1,account=usd_gl,debit=Decimal("2500000"),currency=usd,
   foreign_amount=Decimal("100"),dimensions={"bank_account_id":str(usd_bank.id)})
  JournalLine.objects.create(entry=usd_entry,line_no=2,account=counterpart,credit=Decimal("2500000"),currency=usd,foreign_amount=Decimal("100"))
  usd_result=self.client.get(path.replace("account_code=1121","account_code=1122"))
  self.assertEqual(usd_result.status_code,200,usd_result.data)
  self.assertEqual(Decimal(usd_result.data["groups"][0]["receipt_total"]),Decimal("100"))
  self.assertEqual(Decimal(usd_result.data["groups"][0]["closing_balance"]),Decimal("250"))
  saved=self.client.post(path.replace("/summary/","/save/"),{"name":"Sổ tháng","merge_same":True},format="json")
  self.assertEqual(saved.status_code,201,saved.data)
  self.assertEqual(len(self.client.get(f"/api/v1/bank/deposit-ledger/presets/{self.q}").data),1)
  self.client.force_authenticate(self.approver)
  self.assertEqual(self.client.get(f"/api/v1/bank/deposit-ledger/presets/{self.q}").data,[])

"""Posted general-ledger movements grouped by company bank account."""

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError

from general_ledger.models import Account, JournalLine
from master_data.models import CompanyBankAccount

from . import models

ZERO = Decimal("0")


def parameters(request):
    company = request.query_params.get("company_id")
    try:
        UUID(str(company))
    except (TypeError, ValueError):
        raise ValidationError("Doanh nghiệp không hợp lệ.")
    code = str(request.query_params.get("account_code") or "").strip()
    start = parse_date(str(request.query_params.get("date_from") or ""))
    end = parse_date(str(request.query_params.get("date_to") or ""))
    if not code:
        raise ValidationError("Vui lòng chọn tài khoản kế toán.")
    if not start or not end:
        raise ValidationError("Vui lòng nhập ngày hợp lệ theo định dạng DD/MM/YYYY.")
    if start > end:
        raise ValidationError("Từ ngày không được lớn hơn Đến ngày.")
    account = Account.objects.filter(company_id=company, code=code).first()
    if not account:
        raise ValidationError("Tài khoản kế toán không tồn tại hoặc không hợp lệ.")
    return company, code, start, end


def _baseline(profile):
    """Operational opening kept for accounts without an attributable GL opening entry."""
    value = profile.balance
    for voucher in models.BankTransaction.objects.filter(account=profile, status="POSTED", balance_applied=True):
        value -= voucher.amount if voucher.transaction_type in ("RECEIPT", "INTEREST", "OTHER_IN") else -voucher.amount
    for transfer in models.BankTransfer.objects.filter(destination_account=profile, status="CONFIRMED"):
        value -= transfer.destination_amount
    for transfer in models.BankTransfer.objects.filter(source_account=profile, status="CONFIRMED"):
        value += transfer.source_amount
    return value


def summary(request):
    company, code, start, end = parameters(request)
    branch = request.query_params.get("branch_id")
    search = str(request.query_params.get("search") or "").strip().casefold()
    allowed_accounts = list(Account.objects.filter(company_id=company, code__startswith=code))
    allowed_ids = {account.pk for account in allowed_accounts}
    profiles = list(models.AccountProfile.objects.filter(company_id=company,
        account__gl_account_id__in=allowed_ids).select_related("company", "account__bank", "account__currency", "account__gl_account"))
    if branch:
        profiles = [profile for profile in profiles if str(profile.account.branch_id) == branch]
    by_bank = {str(profile.account_id): profile for profile in profiles}
    by_gl = defaultdict(list)
    for profile in profiles:
        by_gl[profile.account.gl_account_id].append(profile)
    values = {profile.pk: {"opening": ZERO, "debit": ZERO, "credit": ZERO, "has_gl_opening": False} for profile in profiles}
    warnings = set()
    unmapped = CompanyBankAccount.objects.filter(company_id=company, gl_account__isnull=True, status="ACTIVE")
    if branch:
        unmapped = unmapped.filter(branch_id=branch)
    for bank_account in unmapped:
        warnings.add(f"Tài khoản ngân hàng {bank_account.account_number} chưa ánh xạ tài khoản kế toán.")
    journal_lines = JournalLine.objects.filter(entry__company_id=company, entry__status="POSTED",
        entry__posting_date__lte=end, account_id__in=allowed_ids).select_related("entry", "currency")
    for line in journal_lines:
        raw_bank_id = str((line.dimensions or {}).get("bank_account_id") or "")
        if raw_bank_id:
            profile = by_bank.get(raw_bank_id)
            if not profile or profile.account.gl_account_id != line.account_id:
                warnings.add(f"Bút toán {line.entry.number} có ánh xạ tài khoản ngân hàng không hợp lệ hoặc ngoài phạm vi báo cáo.")
                continue
        else:
            possible = by_gl.get(line.account_id, [])
            if len(possible) != 1:
                warnings.add(f"Tài khoản kế toán {line.account.code} có bút toán chưa xác định được tài khoản ngân hàng.")
                continue
            profile = possible[0]
            if profile.account.currency_id != line.currency_id:
                warnings.add(f"Bút toán {line.entry.number} thiếu loại tiền/tài khoản ngân hàng để quy về {profile.account.currency.code}.")
                continue
        if profile.account.currency_id == profile.company.accounting_currency_id:
            debit, credit = line.debit, line.credit
        elif line.currency_id == profile.account.currency_id and line.foreign_amount > ZERO:
            debit = line.foreign_amount if line.debit > ZERO else ZERO
            credit = line.foreign_amount if line.credit > ZERO else ZERO
        else:
            warnings.add(f"Bút toán {line.entry.number} thiếu số tiền nguyên tệ {profile.account.currency.code}.")
            continue
        row = values[profile.pk]
        if line.entry.entry_type == "OPENING":
            row["has_gl_opening"] = True
        if line.entry.posting_date < start:
            row["opening"] += debit - credit
        else:
            row["debit"] += debit
            row["credit"] += credit
    rows = []
    for profile in profiles:
        data = values[profile.pk]
        opening = data["opening"] + (ZERO if data["has_gl_opening"] else _baseline(profile))
        closing = opening + data["debit"] - data["credit"]
        rows.append({"id": str(profile.pk), "bank_account_id": str(profile.account_id),
            "bank_account_no": profile.account.account_number, "bank_name": profile.account.bank.name,
            "gl_account_code": profile.account.gl_account.code, "currency": profile.account.currency.code,
            "opening_balance": str(opening), "debit_amount": str(data["debit"]),
            "credit_amount": str(data["credit"]), "closing_balance": str(closing)})
    rows.sort(key=lambda row: (row["bank_account_no"], row["currency"]))
    if search:
        rows = [row for row in rows if search in row["bank_account_no"].casefold() or search in row["bank_name"].casefold()]
    totals = {}
    for row in rows:
        currency = row["currency"]
        total = totals.setdefault(currency, {"currency": currency, "opening_balance": ZERO,
            "debit_amount": ZERO, "credit_amount": ZERO, "closing_balance": ZERO})
        for field in ("opening_balance", "debit_amount", "credit_amount", "closing_balance"):
            total[field] += Decimal(row[field])
    return {"account_code": code, "date_from": str(start), "date_to": str(end),
        "rows": rows, "totals": [{key: str(value) if isinstance(value, Decimal) else value
            for key, value in total.items()} for total in sorted(totals.values(), key=lambda row: row["currency"])],
        "warnings": sorted(warnings)}


def excel(data):
    def safe(value):
        return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    book = Workbook()
    sheet = book.active
    sheet.title = "So du ngan hang"
    sheet.append(["BẢNG KÊ SỐ DƯ NGÂN HÀNG"])
    sheet.append([f"Từ ngày {data['date_from']} đến ngày {data['date_to']} · Tài khoản {data['account_code']}"])
    sheet.append(["Tài khoản ngân hàng", "Tên ngân hàng", "Loại tiền", "Số dư đầu kỳ", "Phát sinh Nợ", "Phát sinh Có", "Số dư cuối kỳ"])
    sheet["A1"].font = Font(bold=True, size=14)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
    for row in data["rows"]:
        sheet.append([safe(row["bank_account_no"]), safe(row["bank_name"]), row["currency"],
            *[float(Decimal(row[field])) for field in ("opening_balance", "debit_amount", "credit_amount", "closing_balance")]])
    for total in data["totals"]:
        sheet.append(["Tổng cộng", "", total["currency"],
            *[float(Decimal(total[field])) for field in ("opening_balance", "debit_amount", "credit_amount", "closing_balance")]])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
    for column, width in {"A": 24, "B": 46, "C": 14, "D": 21, "E": 21, "F": 21, "G": 21}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=4, min_col=4, max_col=7):
        for cell in row:
            cell.number_format = '#,##0.####;[Red](#,##0.####)'
    output = BytesIO()
    book.save(output)
    return output.getvalue()

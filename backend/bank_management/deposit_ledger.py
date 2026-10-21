"""Read-only bank deposit ledger built from posted journal lines."""

from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from uuid import UUID
from types import SimpleNamespace

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError

from general_ledger.models import Account, JournalLine

from . import balance_report, models

ZERO = Decimal("0")


def _flag(request, name, default):
    value = str(request.query_params.get(name, str(default).lower())).lower()
    if value not in ("true", "false", "1", "0"):
        raise ValidationError({name: "Giá trị tùy chọn không hợp lệ."})
    return value in ("true", "1")


def ledger(request):
    company, code, start, end = balance_report.parameters(request)
    account_id = str(request.query_params.get("bank_account_id") or "")
    if account_id:
        try:
            UUID(account_id)
        except ValueError:
            raise ValidationError({"bank_account_id": "Tài khoản ngân hàng không phù hợp với tài khoản kế toán đã chọn."})
    merge_same = _flag(request, "merge_same", True)
    merge_voucher = _flag(request, "merge_voucher", False)
    order_created = _flag(request, "order_created", True)
    receipt_first = _flag(request, "receipt_first", False)
    search = str(request.query_params.get("search") or "").strip().casefold()
    unfiltered_params = request.query_params.copy()
    unfiltered_params.pop("search", None)
    overview = balance_report.summary(SimpleNamespace(query_params=unfiltered_params))
    profiles = list(models.AccountProfile.objects.filter(company_id=company, account__gl_account__code__startswith=code)
        .select_related("account__gl_account", "account__currency", "account__bank", "company"))
    branch = request.query_params.get("branch_id")
    if branch:
        profiles = [profile for profile in profiles if str(profile.account.branch_id) == branch]
    if account_id and not any(str(profile.account_id) == account_id for profile in profiles):
        raise ValidationError({"bank_account_id": "Tài khoản ngân hàng không phù hợp với tài khoản kế toán đã chọn."})
    if account_id:
        profiles = [profile for profile in profiles if str(profile.account_id) == account_id]
    opening_by_profile = {row["id"]: row for row in overview["rows"]}
    by_bank = {str(profile.account_id): profile for profile in profiles}
    by_gl = defaultdict(list)
    for profile in profiles:
        by_gl[profile.account.gl_account_id].append(profile)
    account_ids = set(Account.objects.filter(company_id=company, code__startswith=code).values_list("id", flat=True))
    lines = list(JournalLine.objects.filter(entry__company_id=company, entry__status="POSTED",
        entry__posting_date__range=(start, end), account_id__in=account_ids)
        .select_related("entry", "account", "currency"))
    entry_ids = {line.entry_id for line in lines}
    counterparts = defaultdict(list)
    for line in JournalLine.objects.filter(entry_id__in=entry_ids).select_related("account"):
        counterparts[line.entry_id].append(line)
    events = defaultdict(list)
    for line in lines:
        raw_bank_id = str((line.dimensions or {}).get("bank_account_id") or "")
        if raw_bank_id:
            profile = by_bank.get(raw_bank_id)
            if not profile or profile.account.gl_account_id != line.account_id:
                continue
        else:
            choices = by_gl[line.account_id]
            if len(choices) != 1 or choices[0].account.currency_id != line.currency_id:
                continue
            profile = choices[0]
        if profile.account.currency_id == profile.company.accounting_currency_id:
            receipt, payment = line.debit, line.credit
        elif line.currency_id == profile.account.currency_id and line.foreign_amount > ZERO:
            receipt = line.foreign_amount if line.debit > ZERO else ZERO
            payment = line.foreign_amount if line.credit > ZERO else ZERO
        else:
            continue
        if not receipt and not payment:
            continue
        others = [other for other in counterparts[line.entry_id] if other.pk != line.pk and other.account_id != line.account_id]
        paired = [other for other in others if abs(other.line_no - line.line_no) == 1]
        counter_codes = sorted({other.account.code for other in (paired or others)})
        events[profile.pk].append({
            "journal_entry_id": str(line.entry_id), "journal_line_ids": [str(line.pk)],
            "source_entity_type": line.entry.source_entity_type or "",
            "source_entity_id": str(line.entry.source_entity_id) if line.entry.source_entity_id else "",
            "posting_date": str(line.entry.posting_date), "document_date": str(line.entry.document_date or line.entry.posting_date),
            "voucher_number": line.entry.number, "description": line.description or line.entry.description,
            "counter_account": ", ".join(counter_codes), "receipt": receipt, "payment": payment,
            "created_at": line.entry.created_at, "line_no": line.line_no,
        })
    groups = []
    for profile in profiles:
        overview_row = opening_by_profile.get(str(profile.pk))
        if not overview_row:
            continue
        opening = Decimal(overview_row["opening_balance"])
        raw = events[profile.pk]
        if opening == ZERO and not raw:
            continue
        if merge_same or merge_voucher:
            merged = {}
            for event in raw:
                direction = "receipt" if event["receipt"] else "payment"
                key = (event["journal_entry_id"], direction) if merge_voucher else (
                    event["journal_entry_id"], direction, event["description"], event["counter_account"])
                if key not in merged:
                    merged[key] = event.copy()
                else:
                    target = merged[key]
                    target["receipt"] += event["receipt"]
                    target["payment"] += event["payment"]
                    target["journal_line_ids"] += event["journal_line_ids"]
                    target["counter_account"] = ", ".join(sorted(set(filter(None,
                        (target["counter_account"] + ", " + event["counter_account"]).split(", ")))))
            raw = list(merged.values())
        raw.sort(key=lambda event: (event["posting_date"],
            0 if receipt_first and event["receipt"] else 1 if receipt_first else 0,
            str(event["created_at"]) if order_created else event["voucher_number"],
            event["line_no"], event["journal_line_ids"][0]))
        running = opening
        receipt_total = payment_total = ZERO
        rows = []
        for event in raw:
            running += event["receipt"] - event["payment"]
            receipt_total += event["receipt"]
            payment_total += event["payment"]
            rows.append({key: str(value) if isinstance(value, Decimal) else value for key, value in event.items()
                if key not in ("created_at", "line_no") } | {"balance": str(running)})
        visible = [row for row in rows if search in row["voucher_number"].casefold() or search in row["description"].casefold()] if search else rows
        if search and not visible:
            continue
        groups.append({"id": str(profile.pk), "bank_account_id": str(profile.account_id),
            "bank_account_no": profile.account.account_number, "bank_name": profile.account.bank.name,
            "currency": profile.account.currency.code, "gl_account_code": profile.account.gl_account.code,
            "opening_balance": str(opening), "receipt_total": str(receipt_total), "payment_total": str(payment_total),
            "closing_balance": str(running), "filtered_receipt_total": str(sum((Decimal(row["receipt"]) for row in visible), ZERO)),
            "filtered_payment_total": str(sum((Decimal(row["payment"]) for row in visible), ZERO)),
            "rows": visible, "row_count": len(rows)})
    groups.sort(key=lambda group: (group["bank_account_no"], group["currency"]))
    return {"account_code": code, "date_from": str(start), "date_to": str(end),
        "groups": groups, "warnings": overview["warnings"], "search": search}


def excel(data):
    def safe(value):
        return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
    book = Workbook()
    sheet = book.active
    sheet.title = "So tien gui"
    sheet.append(["SỔ TIỀN GỬI NGÂN HÀNG"])
    sheet.append([f"Từ ngày {data['date_from']} đến ngày {data['date_to']} · Tài khoản {data['account_code']}"])
    sheet.append(["Ngày hạch toán", "Ngày chứng từ", "Số chứng từ", "Diễn giải", "TK đối ứng", "Thu", "Chi", "Tồn"])
    for cell in sheet[3]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E2F0D9")
    for group in data["groups"]:
        sheet.append([safe(f"{group['bank_account_no']} - {group['bank_name']} ({group['currency']})")])
        sheet.append(["", "", "", "Số dư đầu kỳ", "", "", "", float(Decimal(group["opening_balance"]))])
        for row in group["rows"]:
            sheet.append([row["posting_date"], row["document_date"], safe(row["voucher_number"]),
                safe(row["description"]), safe(row["counter_account"]),
                float(Decimal(row["receipt"])), float(Decimal(row["payment"])), float(Decimal(row["balance"]))])
        sheet.append(["", "", "", "Tổng cộng toàn kỳ", "", float(Decimal(group["receipt_total"])),
            float(Decimal(group["payment_total"])), float(Decimal(group["closing_balance"]))])
        if data["search"]:
            sheet.append(["", "", "", "Tổng các dòng lọc", "", float(Decimal(group["filtered_receipt_total"])),
                float(Decimal(group["filtered_payment_total"]))])
    for column, width in {"A": 18, "B": 18, "C": 22, "D": 54, "E": 18, "F": 20, "G": 20, "H": 20}.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=4, min_col=6, max_col=8):
        for cell in row:
            cell.number_format = '#,##0.####;[Red](#,##0.####)'
    output = BytesIO()
    book.save(output)
    return output.getvalue()

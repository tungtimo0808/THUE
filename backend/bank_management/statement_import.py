"""Preview and atomically import bank statement rows without creating accounting vouchers."""

from decimal import Decimal, InvalidOperation
from datetime import date, datetime
from hashlib import sha256
from uuid import uuid4
from zipfile import BadZipFile

from django.db import transaction
from django.utils.dateparse import parse_date
from rest_framework.exceptions import ValidationError

from . import models


MAX_FILE_SIZE = 20 * 1024 * 1024


def _money(value):
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return None
    if not result.is_finite() or result <= 0 or result.as_tuple().exponent < -4 or result.adjusted() > 15:
        return None
    return result


def _optional_money(value):
    if value in (None, ""):
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        raise ValidationError({"balance": "Số dư sao kê không hợp lệ."})
    if not result.is_finite():
        raise ValidationError({"balance": "Số dư sao kê không hợp lệ."})
    return result


def _account(payload, company_id, branch_id=None):
    account = models.AccountProfile.objects.select_related(
        "account__currency", "account__bank", "account__branch"
    ).filter(pk=payload.get("account"), company_id=company_id, account__status="ACTIVE").first()
    if not account:
        raise ValidationError({"account": "Tài khoản ngân hàng không thuộc doanh nghiệp hoặc không hoạt động."})
    if branch_id and str(account.account.branch_id) != str(branch_id):
        raise ValidationError({"branch_id": "Chi nhánh không quản lý tài khoản ngân hàng đã chọn."})
    if str(payload.get("currency")) not in (str(account.account.currency_id), account.account.currency.code):
        raise ValidationError({"currency": "Loại tiền sao kê không khớp tài khoản ngân hàng."})
    return account


def _source_text(value):
    number_format = getattr(value, "number_format", "")
    value = getattr(value, "value", value)
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float) and value.is_integer():
        if number_format and set(number_format) == {"0"}:
            return str(int(value)).zfill(len(number_format))
        return str(int(value))
    if isinstance(value, int) and number_format and set(number_format) == {"0"}:
        return str(value).zfill(len(number_format))
    return str(value).strip()


def _source_date(value, date_format):
    value = getattr(value, "value", value)
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    raw = _source_text(value)
    parsed = parse_date(raw)
    if parsed:
        return parsed.isoformat()
    if isinstance(value, (int, float)) and value > 0:
        from openpyxl.utils.datetime import from_excel
        return from_excel(value).date().isoformat()
    import re
    parts = re.fullmatch(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", raw)
    if parts:
        first, second, year = map(int, parts.groups())
        if year < 100:
            year += 2000
        month, day = (first, second) if date_format == "MDY" else (second, first)
        try:
            return date(year, month, day).isoformat()
        except ValueError:
            pass
    return ""


def _source_amount(value, number_format):
    value = getattr(value, "value", value)
    raw = _source_text(value).replace(" ", "").replace("\xa0", "")
    if not raw:
        return Decimal("0")
    normalized = raw.replace(".", "").replace(",", ".") if number_format == "VI" and isinstance(value, str) else raw.replace(",", "")
    try:
        amount = Decimal(normalized)
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


def _verify_source_rows(payload, source_file, selected):
    """Re-read selected rows from the uploaded workbook; JSON rows are never authoritative."""
    try:
        header_row = int(payload.get("header_row", 1))
        first_row = int(payload.get("first_data_row", 2))
        if header_row < 1 or first_row <= header_row:
            raise ValueError()
        mapping = payload["column_mapping"]
        if not isinstance(mapping, dict):
            raise ValueError()
        indices = {key: int(value) for key, value in mapping.items() if value not in (None, "")}
        if any(index < 0 for index in indices.values()) or not {"transaction_date", "description"}.issubset(indices):
            raise ValueError()
        if not ({"receipt", "payment"}.issubset(indices) or "amount" in indices):
            raise ValueError()
    except (KeyError, TypeError, ValueError):
        raise ValidationError({"column_mapping": "Dòng tiêu đề, dòng dữ liệu hoặc ghép cột không hợp lệ."})
    try:
        if source_file.name.lower().endswith(".xlsx"):
            from openpyxl import load_workbook
            workbook = load_workbook(source_file, read_only=True, data_only=True)
            if payload.get("sheet_name") not in workbook.sheetnames:
                raise ValidationError({"sheet_name": "Trang tính không tồn tại trong tệp."})
            sheet = workbook[payload["sheet_name"]]
            headers = list(sheet.iter_rows(min_row=header_row, max_row=header_row))[0]
            source = {number: list(sheet.iter_rows(min_row=number, max_row=number))[0]
                      for number in selected if first_row <= number <= sheet.max_row}
            workbook.close()
        else:
            import xlrd
            workbook = xlrd.open_workbook(file_contents=source_file.read())
            if payload.get("sheet_name") not in workbook.sheet_names():
                raise ValidationError({"sheet_name": "Trang tính không tồn tại trong tệp."})
            sheet = workbook.sheet_by_name(payload["sheet_name"])
            headers = sheet.row_values(header_row - 1)
            source = {number: sheet.row_values(number - 1) for number in selected if first_row <= number <= sheet.nrows}
    except ValidationError:
        raise
    except Exception:
        raise ValidationError({"source_file": "Không đọc được trang tính trong tệp Excel."})
    finally:
        source_file.seek(0)
    submitted = {int(row["source_row"]): row for row in payload["rows"]}
    verified_source = {}
    for number in selected:
        if number not in source or number not in submitted:
            raise ValidationError({"selected_rows": f"Dòng {number} không tồn tại trong trang tính đã chọn."})
        cells, row = source[number], submitted[number]
        def cell(key):
            index = indices.get(key)
            return cells[index] if index is not None and index < len(cells) else None
        description = _source_text(cell("description"))
        reference = _source_text(cell("reference"))
        transaction_date = _source_date(cell("transaction_date"), payload.get("date_format"))
        value_date = _source_date(cell("value_date"), payload.get("date_format")) if cell("value_date") not in (None, "") else ""
        if description != str(row.get("description") or "").strip() or reference != str(row.get("reference") or "").strip() or transaction_date != row.get("transaction_date") or value_date != (row.get("value_date") or ""):
            raise ValidationError({"selected_rows": f"Dòng {number} khác nội dung hoặc ngày trong tệp nguồn; hãy kiểm tra lại."})
        for key in ("source_account", "source_currency", "counterparty", "counterparty_account", "ending_balance"):
            if key in indices and _source_text(cell(key)) != str(row.get(key) or "").strip():
                raise ValidationError({"selected_rows": f"Dòng {number}: cột {key} khác tệp nguồn."})
            if key not in indices and str(row.get(key) or "").strip():
                raise ValidationError({"selected_rows": f"Dòng {number}: cột {key} chưa được ghép từ tệp nguồn."})
        if {"receipt", "payment"}.issubset(indices):
            receipt = _source_amount(cell("receipt"), payload.get("number_format"))
            payment = _source_amount(cell("payment"), payload.get("number_format"))
            direction, amount = ("CREDIT", receipt) if receipt and receipt > 0 and payment == 0 else (("DEBIT", payment) if payment and payment > 0 and receipt == 0 else ("", None))
        else:
            signed = _source_amount(cell("amount"), payload.get("number_format"))
            sign_mode = payload.get("signed_mode")
            direction_cell = _source_text(cell("direction")).casefold()
            if "direction" in indices:
                direction = "CREDIT" if direction_cell in {"thu", "co", "có", "credit", "cr", "in"} else "DEBIT" if direction_cell in {"chi", "no", "nợ", "debit", "dr", "out"} else ""
            else:
                if sign_mode not in {"positive_receipt", "positive_payment"}:
                    raise ValidationError({"signed_mode": "Cần xác nhận quy ước dấu của cột số tiền."})
                direction = "" if signed is None else (("CREDIT" if signed > 0 else "DEBIT") if sign_mode == "positive_receipt" else ("DEBIT" if signed > 0 else "CREDIT"))
            amount = abs(signed) if signed is not None else None
        if not direction or amount is None or direction != row.get("direction") or amount != _money(row.get("amount")):
            raise ValidationError({"selected_rows": f"Dòng {number}: số tiền hoặc chiều Thu/Chi khác tệp nguồn."})
        verified_source[number] = {_source_text(header) or f"Cột {index + 1}": _source_text(value)
                                   for index, (header, value) in enumerate(zip(headers, cells))}
    return verified_source


def preview(payload, company_id, branch_id=None):
    account = _account(payload, company_id, branch_id)
    start, end = parse_date(str(payload.get("period_from") or "")), parse_date(str(payload.get("period_to") or ""))
    if not start or not end or end < start:
        raise ValidationError({"period_to": "Khoảng ngày sao kê không hợp lệ."})
    rows = payload.get("rows")
    if not isinstance(rows, list) or not rows or len(rows) > 20000:
        raise ValidationError({"rows": "Tệp phải có từ 1 đến 20.000 dòng giao dịch để kiểm tra."})
    if not str(payload.get("sheet_name") or "").strip() or not isinstance(payload.get("column_mapping"), dict):
        raise ValidationError({"column_mapping": "Thiếu trang tính hoặc cấu hình ghép cột."})
    unique_reference = bool(payload.get("unique_reference"))
    existing = models.StatementLine.objects.filter(statement__account=account, deleted_at__isnull=True).select_related("statement__account__account")
    existing_by_reference = {line.reference: line for line in existing if line.reference} if unique_reference else {}
    existing_fingerprints = {}
    for line in existing:
        key = (line.transaction_date, line.direction, line.amount, line.description.strip().casefold())
        existing_fingerprints.setdefault(key, line)
    seen_references, seen_fingerprints, source_rows = set(), {}, set()
    results = []
    receipts = payments = Decimal("0")
    for raw in rows:
        if not isinstance(raw, dict):
            raise ValidationError({"rows": "Dữ liệu dòng sao kê không hợp lệ."})
        try:
            source_row = int(raw.get("source_row"))
        except (ValueError, TypeError):
            source_row = 0
        date = parse_date(str(raw.get("transaction_date") or ""))
        value_date = parse_date(str(raw.get("value_date") or "")) if raw.get("value_date") else None
        direction = str(raw.get("direction") or "").upper()
        amount = _money(raw.get("amount"))
        description = str(raw.get("description") or "").strip()
        reference = str(raw.get("reference") or "").strip()
        reason = ""
        status = "VALID"
        matched = None
        if source_row <= 0 or source_row in source_rows:
            reason = "Số dòng nguồn không hợp lệ hoặc bị lặp."
        elif not date or not start <= date <= end:
            reason = "Ngày giao dịch không hợp lệ hoặc ngoài khoảng sao kê."
        elif raw.get("value_date") and not value_date:
            reason = "Ngày giá trị không hợp lệ."
        elif not description:
            reason = "Thiếu nội dung giao dịch."
        elif len(description) > 1000 or len(reference) > 150 or len(str(raw.get("counterparty") or "")) > 255:
            reason = "Nội dung, mã giao dịch hoặc tên đối tác vượt độ dài cho phép."
        elif direction not in ("CREDIT", "DEBIT"):
            reason = "Chiều giao dịch phải là Thu hoặc Chi."
        elif amount is None:
            reason = "Số tiền giao dịch phải lớn hơn 0."
        elif raw.get("source_account") and str(raw["source_account"]).strip() != account.account.account_number:
            reason = "Số tài khoản trong tệp không khớp tài khoản đã chọn."
        elif raw.get("source_currency") and str(raw["source_currency"]).strip().upper() != account.account.currency.code:
            reason = "Loại tiền trong tệp không khớp tài khoản đã chọn."
        if reason:
            status = "ERROR"
        elif unique_reference and reference and (reference in seen_references or reference in existing_by_reference):
            status, reason = "DUPLICATE", "Mã giao dịch duy nhất đã có trong tệp hoặc hệ thống."
            matched = existing_by_reference.get(reference)
        else:
            key = (date, direction, amount, description.casefold())
            matched = None if unique_reference and reference else seen_fingerprints.get(key) or existing_fingerprints.get(key)
            if matched:
                status, reason = "SUSPECT", "Ngày, chiều, số tiền và nội dung giống giao dịch khác; cần xác nhận."
            if unique_reference and reference:
                seen_references.add(reference)
            seen_fingerprints[key] = raw
        source_rows.add(source_row)
        if status in ("VALID", "SUSPECT"):
            if direction == "CREDIT":
                receipts += amount
            else:
                payments += amount
        results.append({"source_row": source_row, "status": status, "reason": reason,
            "transaction_date": str(date) if date else "", "direction": direction,
            "amount": str(amount) if amount is not None else "", "description": description,
            "reference": reference, "currency": account.account.currency.code,
            "ending_balance": str(raw.get("ending_balance") or ""),
            "counterparty": str(raw.get("counterparty") or ""),
            "counterparty_account": str(raw.get("counterparty_account") or ""),
            "existing": ({"number": matched.statement.source_reference,
            "transaction_date": str(matched.transaction_date), "direction": matched.direction,
            "amount": str(matched.amount), "description": matched.description,
            "reference": matched.reference,
            "account_number": matched.statement.account.account.account_number} if isinstance(matched, models.StatementLine) else None)})
    opening = _optional_money(payload.get("opening_balance"))
    closing = _optional_money(payload.get("closing_balance"))
    calculated = opening + receipts - payments if opening is not None else None
    return {"account": str(account.pk), "currency": account.account.currency.code,
        "rows": results, "total_receipts": str(receipts), "total_payments": str(payments),
        "opening_balance": str(opening) if opening is not None else None,
        "closing_balance": str(closing) if closing is not None else None,
        "calculated_closing": str(calculated) if calculated is not None else None,
        "balance_difference": str(closing - calculated) if closing is not None and calculated is not None else None,
        "counts": {status: sum(row["status"] == status for row in results) for status in ("VALID", "ERROR", "DUPLICATE", "SUSPECT")}}


@transaction.atomic
def confirm(payload, source_file, company, actor_user, request):
    if not source_file or not source_file.name.lower().endswith((".xlsx", ".xls")) or source_file.size > MAX_FILE_SIZE or source_file.size == 0:
        raise ValidationError({"source_file": "Chỉ nhận tệp .xlsx hoặc .xls, tối đa 20 MB."})
    signature = source_file.read(8)
    source_file.seek(0)
    if source_file.name.lower().endswith(".xlsx") and not signature.startswith(b"PK") or source_file.name.lower().endswith(".xls") and signature != bytes.fromhex("D0CF11E0A1B11AE1"):
        raise ValidationError({"source_file": "Tệp Excel không đọc được hoặc không đúng định dạng."})
    if source_file.name.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(source_file, read_only=True, data_only=True)
            if payload.get("sheet_name") not in workbook.sheetnames:
                raise ValidationError({"sheet_name": "Trang tính không tồn tại trong tệp."})
            workbook.close()
        except (BadZipFile, OSError, ValueError, KeyError):
            raise ValidationError({"source_file": "Tệp Excel bị hỏng hoặc được bảo vệ bằng mật khẩu."})
        finally:
            source_file.seek(0)
    account = _account(payload, company.pk, request.query_params.get("branch_id"))
    models.AccountProfile.objects.select_for_update().get(pk=account.pk)
    result = preview(payload, company.pk, request.query_params.get("branch_id"))
    if not isinstance(payload.get("selected_rows"), list) or not isinstance(payload.get("confirmed_suspects", []), list):
        raise ValidationError({"selected_rows": "Danh sách dòng import không hợp lệ."})
    try:
        selected = {int(value) for value in payload.get("selected_rows", [])}
        confirmed = {int(value) for value in payload.get("confirmed_suspects", [])}
    except (ValueError,TypeError):
        raise ValidationError({"selected_rows": "Danh sách dòng import không hợp lệ."})
    row_by_number = {row["source_row"]: row for row in result["rows"]}
    if not selected or len(selected) > len(row_by_number) or not selected.issubset(row_by_number):
        raise ValidationError({"selected_rows": "Chưa chọn dòng giao dịch hợp lệ để nhập."})
    for number in selected:
        status = row_by_number[number]["status"]
        if status not in ("VALID", "SUSPECT") or status == "SUSPECT" and number not in confirmed:
            raise ValidationError({"selected_rows": f"Dòng {number} có lỗi, trùng hoặc chưa xác nhận nghi trùng; hãy kiểm tra lại."})
    verified_source = _verify_source_rows(payload, source_file, selected)
    selected_receipts = sum((Decimal(row_by_number[number]["amount"]) for number in selected if row_by_number[number]["direction"] == "CREDIT"), Decimal("0"))
    selected_payments = sum((Decimal(row_by_number[number]["amount"]) for number in selected if row_by_number[number]["direction"] == "DEBIT"), Decimal("0"))
    selected_difference = (Decimal(result["closing_balance"]) - Decimal(result["opening_balance"]) - selected_receipts + selected_payments
                           if result["opening_balance"] is not None and result["closing_balance"] is not None else None)
    if selected_difference not in (None, Decimal("0")) and not payload.get("confirm_balance_warning"):
        raise ValidationError({"balance": "Số dư cuối sao kê không khớp; cần xác nhận cảnh báo."})
    file_hash = sha256()
    for chunk in source_file.chunks():
        file_hash.update(chunk)
    source_file.seek(0)
    file_digest = file_hash.hexdigest()
    import_key = sha256(f"{file_digest}:{payload['sheet_name']}:{','.join(map(str, sorted(selected)))}".encode()).hexdigest()
    if models.StatementImport.objects.filter(account=account, raw_file__import_key=import_key).exists():
        raise ValidationError({"source_file": "Các dòng đã chọn của tệp này đã được import trước đó."})
    previous = models.StatementLine.objects.filter(statement__account=account,
        statement__raw_file__file_hash=file_digest,
        statement__raw_file__sheet_name=payload["sheet_name"],
        deleted_at__isnull=True).values_list("raw_data", flat=True)
    repeated_rows = selected & {int(data.get("source_row")) for data in previous if isinstance(data, dict) and str(data.get("source_row", "")).isdigit()}
    if repeated_rows:
        raise ValidationError({"selected_rows": f"Dòng {min(repeated_rows)} của tệp này đã được import; hãy bỏ chọn dòng trùng."})
    from .services import audit
    statement = models.StatementImport.objects.create(company=company, account=account, source="FILE",
        source_reference=f"IMP-{uuid4().hex[:12].upper()}", period_from=payload["period_from"],period_to=payload["period_to"],
        source_file=source_file, created_by=actor_user,
        raw_file={"file_name": source_file.name, "file_size": source_file.size,
            "sheet_name": payload["sheet_name"], "header_row": payload.get("header_row"),
            "first_data_row": payload.get("first_data_row"), "date_format": payload.get("date_format"),
            "number_format": payload.get("number_format"), "column_mapping": payload["column_mapping"],
            "opening_balance": result["opening_balance"], "closing_balance": result["closing_balance"],
            "balance_difference": str(selected_difference) if selected_difference is not None else None,
            "confirmed_suspects": sorted(confirmed & selected),
            "active_total_receipts": str(selected_receipts),
            "active_total_payments": str(selected_payments),
            "active_line_count": len(selected),
            "file_hash": file_digest, "import_key": import_key})
    used_external_ids = {str(raw.get("reference") or "").strip() for raw in payload["rows"] if int(raw["source_row"]) in selected and str(raw.get("reference") or "").strip()}
    for raw in payload["rows"]:
        source_row = int(raw["source_row"])
        if source_row not in selected:
            continue
        reference = str(raw.get("reference") or "").strip()
        external_id = reference if reference and not statement.lines.filter(external_id=reference).exists() else f"line-{source_row}"
        while external_id in used_external_ids and external_id != reference:
            external_id += "-source"
        models.StatementLine.objects.create(statement=statement,external_id=external_id,
            transaction_date=raw["transaction_date"],value_date=raw.get("value_date") or None,
            direction=raw["direction"],amount=raw["amount"],currency=account.account.currency,
            reference=str(raw.get("reference") or ""),description=raw["description"],
            counterparty=str(raw.get("counterparty") or ""),raw_data={"source_row": source_row,
                "source": verified_source[source_row], "counterparty_account": str(raw.get("counterparty_account") or ""),
                "ending_balance": raw.get("ending_balance"), "confirmed_suspect": source_row in confirmed})
    audit(request,statement,"IMPORT",after={"lines":len(selected),"source_reference":statement.source_reference})
    return statement

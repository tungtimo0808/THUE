import calendar
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from common_services.models import FiscalPeriod, FiscalPeriodStatus, FiscalYear, FiscalYearStatus
from general_ledger.models import (
    Account, AccountNature, AccountStatus, AccountingPeriod,
    ChartOfAccounts, JournalEntry, NormalBalance,
)
from iam.models import CompanyMembership, Permission, Role, RolePermission, UserRoleAssignment
from master_data.models import Currency, Party, PartyRole, PaymentTerm, TaxRate
from organization.models import Branch, Company
from accounts_payable.models import (
    DocumentStatus, Netting, Payable, PayableAdjustment,
    PayableDocument, PayableDocumentLine, PayableStatus,
    Payment, PaymentAllocation, SupplierInvoice
)


class SupplierDebtTrackingTests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.currency = Currency.objects.get(code="VND")
        self.company = Company.objects.create(
            code="CO_AP_TRACK",
            name="Công ty TNHH Theo Dõi Công Nợ",
            legal_name="Công ty TNHH Theo Dõi Công Nợ",
            tax_code="0109998888",
            accounting_currency=self.currency,
            legal_reporting_currency=self.currency,
        )
        self.branch = Branch.objects.create(company=self.company, code="B_HCM", name="Chi nhánh TP.HCM")

        # Users
        self.admin_user = get_user_model().objects.create_superuser(
            username="admin_user", email="admin@example.com", password="pass"
        )
        self.tracking_user = get_user_model().objects.create_user(
            username="tracking_user", email="tracking@example.com", password="pass"
        )
        self.tracking_account = self.tracking_user.iam_account
        self.unauth_user = get_user_model().objects.create_user(
            username="unauth_user", email="unauth@example.com", password="pass"
        )
        self.unauth_account = self.unauth_user.iam_account

        # Membership
        self.membership = CompanyMembership.objects.create(
            company=self.company, user=self.tracking_account,
            created_by=self.admin_user.iam_account
        )

        # Permissions: grant AP.TRACKING.VIEW, EXPORT, PAYABLE, DOCUMENT, PAYMENT, ADJUSTMENT, NETTING
        self.role = Role.objects.create(company=self.company, code="AP_TRACKER", name="AP Tracking Officer")
        for pcode in (
            "AP.TRACKING.VIEW", "AP.TRACKING.EXPORT", "AP.PAYABLE.VIEW",
            "AP.DOCUMENT.VIEW", "AP.PAYMENT.VIEW", "AP.ADJUSTMENT.VIEW", "AP.NETTING.VIEW"
        ):
            perm = Permission.objects.filter(code=pcode).first()
            if perm:
                RolePermission.objects.create(role=self.role, permission=perm)
        UserRoleAssignment.objects.create(
            company_membership=self.membership, role=self.role, assigned_by=self.admin_user.iam_account
        )

        # Fiscal Year & Periods
        fy = FiscalYear.objects.create(
            company=self.company,
            code=f"FY_{self.today.year}",
            name=f"Năm {self.today.year}",
            start_date=date(self.today.year, 1, 1),
            end_date=date(self.today.year, 12, 31),
            status=FiscalYearStatus.ACTIVE,
            created_by=self.admin_user.iam_account,
        )
        for m in range(1, 13):
            last_day = calendar.monthrange(self.today.year, m)[1]
            FiscalPeriod.objects.create(
                company=self.company, fiscal_year=fy, code=f"FP_{self.today.year}_{m:02d}",
                name=f"Kỳ {m}", period_number=m, start_date=date(self.today.year, m, 1),
                end_date=date(self.today.year, m, last_day), status=FiscalPeriodStatus.OPEN
            )
            AccountingPeriod.objects.create(
                company=self.company, code=f"GL_{self.today.year}_{m:02d}",
                fiscal_year=self.today.year, period_number=m, start_date=date(self.today.year, m, 1),
                end_date=date(self.today.year, m, last_day), status="OPEN"
            )

        # COA & Accounts
        self.chart = ChartOfAccounts.objects.create(
            company=self.company, code="COA_TRACK", name="COA AP", is_default=True, effective_from=self.today
        )
        self.account_331 = Account.objects.create(
            company=self.company, chart_of_accounts=self.chart, code="331",
            name="Phải trả người bán", account_type=AccountNature.LIABILITY,
            normal_balance=NormalBalance.CREDIT, status=AccountStatus.ACTIVE, requires_party=True
        )
        self.account_642 = Account.objects.create(
            company=self.company, chart_of_accounts=self.chart, code="642",
            name="Chi phí quản lý", account_type=AccountNature.EXPENSE,
            normal_balance=NormalBalance.DEBIT, status=AccountStatus.ACTIVE
        )

        # Suppliers
        self.supplier_a = Party.objects.create(
            company=self.company, code="NCC_A", name="Nhà cung cấp A Alpha",
            tax_code="0101111222", default_currency=self.currency, status="ACTIVE"
        )
        PartyRole.objects.create(party=self.supplier_a, role_type="VENDOR", active_from=self.today)

        self.supplier_b = Party.objects.create(
            company=self.company, code="NCC_B", name="Nhà cung cấp B Beta",
            tax_code="0102222333", default_currency=self.currency, status="ACTIVE"
        )
        PartyRole.objects.create(party=self.supplier_b, role_type="VENDOR", active_from=self.today)

        self.supplier_c = Party.objects.create(
            company=self.company, code="NCC_C", name="Nhà cung cấp C Zero Balance",
            tax_code="0103333444", default_currency=self.currency, status="ACTIVE"
        )
        PartyRole.objects.create(party=self.supplier_c, role_type="VENDOR", active_from=self.today)

        self.client = APIClient()
        self.client.force_authenticate(self.tracking_user)
        self.query = f"?company_id={self.company.id}"

    def create_payable_with_document(self, supplier, amount, doc_date, due_date, status="POSTED"):
        doc = PayableDocument.objects.create(
            company=self.company,
            branch=self.branch,
            number=f"DOC_{supplier.code}_{doc_date.strftime('%Y%m%d')}_{amount}",
            document_date=doc_date,
            posting_date=doc_date,
            due_date=due_date,
            supplier=supplier,
            payable_account=self.account_331,
            currency=self.currency,
            subtotal=amount,
            tax_amount=Decimal("0"),
            total_amount=amount,
            status=DocumentStatus.POSTED if status == "POSTED" else DocumentStatus.DRAFT,
            created_by=self.tracking_account,
        )
        PayableDocumentLine.objects.create(
            document=doc,
            line_no=1,
            description="Chi phí mua hàng hóa",
            debit_account=self.account_642,
            amount=amount,
            total_amount=amount,
        )
        if status == "POSTED":
            payable = Payable.objects.create(
                company=self.company,
                branch=self.branch,
                supplier=supplier,
                document=doc,
                currency=self.currency,
                exchange_rate=Decimal("1"),
                original_amount=amount,
                paid_amount=Decimal("0"),
                netted_amount=Decimal("0"),
                adjustment_amount=Decimal("0"),
                balance=amount,
                due_date=due_date,
                status=PayableStatus.OPEN,
            )
            return payable
        return None

    def test_main_flow_summary_and_detail_steps1_to_37(self):
        """
        Kiểm thử toàn diện Luồng chính Bước 1 đến 37:
        - Báo cáo tổng hợp công nợ theo nhà cung cấp (Opening, Increase, Decrease, Closing, Due, Overdue)
        - Báo cáo chi tiết từng khoản nợ sắp xếp theo hạn thanh toán tăng dần
        """
        # Giả lập kỳ theo dõi: Tháng hiện tại
        from_date = self.today.replace(day=1)
        to_date = self.today

        # 1. NCC A:
        # - Khoản 1: Phát sinh trước kỳ (doc_date = 20 ngày trước from_date, due_date = 5 ngày trước from_date) -> 50.000.000
        #   Thanh toán trước kỳ 20.000.000 -> Opening payable = 30.000.000
        pre_date = from_date - timedelta(days=20)
        p1 = self.create_payable_with_document(self.supplier_a, Decimal("50000000"), pre_date, from_date - timedelta(days=5))
        pm1 = Payment.objects.create(
            company=self.company, supplier=self.supplier_a, currency=self.currency,
            payment_date=from_date - timedelta(days=10), amount=Decimal("20000000"),
            allocated_amount=Decimal("20000000"), external_type="BANK_TRANSFER",
            external_reference="UNC_PRE_01", status="CONFIRMED", created_by=self.tracking_account
        )
        PaymentAllocation.objects.create(
            payment=pm1, payable=p1, amount=Decimal("20000000"), created_by=self.tracking_account
        )
        p1.paid_amount = Decimal("20000000")
        p1.balance = Decimal("30000000")
        p1.save()

        # - Khoản 2: Phát sinh trong kỳ (doc_date = from_date + 2 ngày, due_date = to_date - 2 ngày) -> 40.000.000
        #   Quá hạn 2 ngày tại to_date!
        p2 = self.create_payable_with_document(self.supplier_a, Decimal("40000000"), from_date + timedelta(days=2), to_date - timedelta(days=2))

        # - Khoản 3: Phát sinh trong kỳ (doc_date = from_date + 3 ngày, due_date = to_date + 10 ngày) -> 30.000.000
        #   Chưa đến hạn tại to_date!
        #   Trong kỳ thanh toán 10.000.000 cho khoản 3
        p3 = self.create_payable_with_document(self.supplier_a, Decimal("30000000"), from_date + timedelta(days=3), to_date + timedelta(days=10))
        pm2 = Payment.objects.create(
            company=self.company, supplier=self.supplier_a, currency=self.currency,
            payment_date=from_date + timedelta(days=4), amount=Decimal("10000000"),
            allocated_amount=Decimal("10000000"), external_type="BANK_TRANSFER",
            external_reference="UNC_CUR_01", status="CONFIRMED", created_by=self.tracking_account
        )
        PaymentAllocation.objects.create(
            payment=pm2, payable=p3, amount=Decimal("10000000"), created_by=self.tracking_account
        )
        p3.paid_amount = Decimal("10000000")
        p3.balance = Decimal("20000000")
        p3.save()

        # 2. Truy vấn API Summary (Bước 1-24)
        url = f"/api/v1/accounts-payable/tracking/summary/{self.query}&from_date={from_date}&to_date={to_date}"
        res = self.client.get(url)
        self.assertEqual(res.status_code, 200, res.data)

        results = res.data["results"]
        supp_a_row = next((r for r in results if r["supplier_id"] == str(self.supplier_a.id)), None)
        self.assertIsNotNone(supp_a_row)

        # Kiểm tra công thức tài chính cho NCC A:
        # Opening: 30.000.000
        # Increase: 40.000.000 + 30.000.000 = 70.000.000
        # Decrease: 10.000.000
        # Closing: 30.000.000 + 70.000.000 - 10.000.000 = 90.000.000
        self.assertEqual(Decimal(supp_a_row["opening_payable_balance"]), Decimal("30000000"))
        self.assertEqual(Decimal(supp_a_row["period_increase"]), Decimal("70000000"))
        self.assertEqual(Decimal(supp_a_row["period_decrease"]), Decimal("10000000"))
        self.assertEqual(Decimal(supp_a_row["closing_payable_balance"]), Decimal("90000000"))

        # Quá hạn tại to_date: Khoản 1 (30M) + Khoản 2 (40M) = 70.000.000
        self.assertEqual(Decimal(supp_a_row["overdue_amount"]), Decimal("70000000"))
        # Số ngày quá hạn max: Khoản 1 quá hạn (to_date - (from_date - 5)) ngày
        expected_max_od = (to_date - (from_date - timedelta(days=5))).days
        self.assertEqual(supp_a_row["max_overdue_days"], expected_max_od)

        # 3. Truy vấn API Detail cho NCC A (Bước 25-37)
        detail_url = f"/api/v1/accounts-payable/tracking/detail/{self.query}&supplier_id={self.supplier_a.id}&from_date={from_date}&to_date={to_date}"
        det_res = self.client.get(detail_url)
        self.assertEqual(det_res.status_code, 200, det_res.data)

        det_lines = det_res.data["results"]
        self.assertEqual(len(det_lines), 3)

        # BR-AP-TRK-43: Kiểm tra thứ tự hạn thanh toán tăng dần
        due_dates = [x["due_date"] for x in det_lines]
        self.assertEqual(due_dates, sorted(due_dates))

        # Kiểm tra khoản 3 (chưa đến hạn, thanh toán một phần):
        line3 = next(x for x in det_lines if x["payable_id"] == str(p3.id))
        self.assertEqual(Decimal(line3["remaining_amount"]), Decimal("20000000"))
        self.assertEqual(line3["due_status"], "NOT_DUE")
        self.assertEqual(line3["status"], "PARTIALLY_PAID")
        self.assertEqual(line3["aging_bucket"], "NOT_DUE")
        self.assertEqual(line3["days_remaining"], 10)
        self.assertEqual(line3["overdue_days"], 0)

        # Kiểm tra khoản 2 (quá hạn 2 ngày):
        line2 = next(x for x in det_lines if x["payable_id"] == str(p2.id))
        self.assertEqual(Decimal(line2["remaining_amount"]), Decimal("40000000"))
        self.assertEqual(line2["due_status"], "OVERDUE")
        self.assertEqual(line2["status"], "OVERDUE")
        self.assertEqual(line2["aging_bucket"], "1_30")
        self.assertEqual(line2["overdue_days"], 2)
        self.assertTrue(line2["is_overdue"])

    def test_payable_timeline_and_preview_steps38_to_43(self):
        """
        Kiểm thử Bước 38-43:
        - Hiển thị timeline lịch sử phát sinh (ghi nhận, thanh toán, điều chỉnh, bù trừ)
        - Mở màn hình xem trước chứng từ chỉ đọc
        """
        p = self.create_payable_with_document(self.supplier_a, Decimal("100000000"), self.today - timedelta(days=15), self.today)
        # 1. Thanh toán 30M
        pay = Payment.objects.create(
            company=self.company, supplier=self.supplier_a, currency=self.currency,
            payment_date=self.today - timedelta(days=10), amount=Decimal("30000000"),
            allocated_amount=Decimal("30000000"), external_type="BANK_TRANSFER",
            external_reference="UNC_0099", status="CONFIRMED", created_by=self.tracking_account
        )
        PaymentAllocation.objects.create(payment=pay, payable=p, amount=Decimal("30000000"), created_by=self.tracking_account)

        # 2. Điều chỉnh giảm 10M
        PayableAdjustment.objects.create(
            company=self.company, payable=p, adjustment_date=self.today - timedelta(days=8),
            amount=Decimal("-10000000"), reason="Giảm giá hàng lỗi", posting_reference="ADJ_01",
            status="POSTED", created_by=self.tracking_account
        )

        # 3. Bù trừ công nợ 5M
        Netting.objects.create(
            company=self.company, number="NET_01", payable=p, counterparty_reference="BB_01",
            netting_date=self.today - timedelta(days=5), amount=Decimal("5000000"),
            reason="Bù trừ với khoản phải thu", status="POSTED", created_by=self.tracking_account
        )

        # Gọi API timeline
        time_res = self.client.get(f"/api/v1/accounts-payable/tracking/timeline/{self.query}&payable_id={p.id}")
        self.assertEqual(time_res.status_code, 200, time_res.data)

        events = time_res.data["events"]
        self.assertEqual(len(events), 4)
        self.assertEqual([e["event_type"] for e in events], ["RECOGNITION", "PAYMENT", "ADJUSTMENT", "NETTING"])
        self.assertTrue(all(e["can_view_document"] is True for e in events))

        # Preview chứng từ nguồn ở chế độ READ_ONLY
        prev_res = self.client.get(
            f"/api/v1/accounts-payable/tracking/document-preview/{self.query}&document_type=AP.DOCUMENT&document_id={p.document.id}"
        )
        self.assertEqual(prev_res.status_code, 200, prev_res.data)
        self.assertEqual(prev_res.data["mode"], "READ_ONLY")
        self.assertEqual(prev_res.data["number"], p.document.number)

    def test_permission_checks_ef01_ef02_br01_br04(self):
        """
        EF-01, EF-02, BR-01, BR-04:
        - Chặn user không có quyền theo dõi công nợ (403).
        - Admin không mặc nhiên có quyền nếu chưa được phân quyền AP (BR-04).
        - User không thuộc công ty bị chặn (EF-02).
        """
        # 1. User có membership trong công ty nhưng không có quyền AP.TRACKING.VIEW -> EF-01
        CompanyMembership.objects.create(
            company=self.company, user=self.unauth_account, created_by=self.admin_user.iam_account
        )
        self.client.force_authenticate(self.unauth_user)
        res = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}")
        self.assertEqual(res.status_code, 403)
        self.assertIn("Bạn không có quyền theo dõi công nợ nhà cung cấp", str(res.data))

        # 2. User không có membership trong công ty -> EF-02
        outsider = get_user_model().objects.create_user(
            username="outsider_user", email="outsider@example.com", password="pass"
        )
        self.client.force_authenticate(outsider)
        out_res = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}")
        self.assertEqual(out_res.status_code, 403)
        self.assertIn("Không xác định được phạm vi dữ liệu của tài khoản", str(out_res.data))

        # 3. Admin không mặc nhiên có quyền (BR-04)
        self.client.force_authenticate(self.admin_user)
        admin_res = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}")
        self.assertEqual(admin_res.status_code, 403)

    def test_document_preview_permission_denied_af10_br03(self):
        """
        AF-10, BR-03: User có quyền xem công nợ nhưng KHÔNG có quyền xem loại chứng từ cụ thể (ví dụ AP.PAYMENT).
        """
        p = self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today, self.today)
        pay = Payment.objects.create(
            company=self.company, supplier=self.supplier_a, currency=self.currency,
            payment_date=self.today, amount=Decimal("5000000"), allocated_amount=Decimal("5000000"),
            external_type="BANK_TRANSFER", external_reference="UNC_REVOKED", status="CONFIRMED",
            created_by=self.tracking_account
        )

        # Tước quyền AP.PAYMENT.VIEW của role
        p_perm = Permission.objects.filter(code="AP.PAYMENT.VIEW").first()
        if p_perm:
            RolePermission.objects.filter(role=self.role, permission=p_perm).delete()

        res = self.client.get(
            f"/api/v1/accounts-payable/tracking/document-preview/{self.query}&document_type=AP.PAYMENT&document_id={pay.id}"
        )
        self.assertEqual(res.status_code, 403)
        self.assertIn("Bạn không có quyền xem chứng từ này", str(res.data))

    def test_date_validation_ef03_ef04(self):
        """
        EF-03: Sai định dạng ngày.
        EF-04: Từ ngày lớn hơn Đến ngày.
        """
        # EF-03: Sai định dạng ngày
        res_bad_fmt = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}&from_date=2026/10/01")
        self.assertEqual(res_bad_fmt.status_code, 400)
        self.assertIn("Ngày theo dõi không đúng định dạng", str(res_bad_fmt.data))

        # EF-04: from_date > to_date
        res_bad_order = self.client.get(
            f"/api/v1/accounts-payable/tracking/summary/{self.query}&from_date=2026-10-15&to_date=2026-10-10"
        )
        self.assertEqual(res_bad_order.status_code, 400)
        self.assertIn("Từ ngày không được lớn hơn Đến ngày", str(res_bad_order.data))

    def test_deleted_or_inaccessible_supplier_ef06(self):
        """EF-06: Nhà cung cấp không tồn tại."""
        res = self.client.get(f"/api/v1/accounts-payable/tracking/detail/{self.query}&supplier_id={self.company.id}")
        self.assertEqual(res.status_code, 400)
        self.assertIn("Nhà cung cấp không tồn tại", str(res.data))

    def test_cancelled_payable_in_timeline_ef08(self):
        """EF-08: Khoản công nợ đã bị hủy không thể xem timeline."""
        p = self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today, self.today)
        p.status = PayableStatus.CANCELLED
        p.save()

        res = self.client.get(f"/api/v1/accounts-payable/tracking/timeline/{self.query}&payable_id={p.id}")
        self.assertEqual(res.status_code, 400)
        self.assertIn("Khoản công nợ không còn hiệu lực", str(res.data))

    def test_advance_payment_separate_tracking_br18_br21_br22_af03(self):
        """
        BR-18, BR-21, BR-22, AF-03:
        Khoản trả trước được theo dõi riêng biệt, không tự động cấn trừ nếu chưa có chứng từ bù trừ.
        """
        # Tạo khoản nợ 50M
        self.create_payable_with_document(self.supplier_b, Decimal("50000000"), self.today, self.today + timedelta(days=10))

        # Thanh toán 80M cho NCC B (trong đó 50M đã phân bổ, 30M chưa phân bổ là tiền trả trước)
        Payment.objects.create(
            company=self.company, supplier=self.supplier_b, currency=self.currency,
            payment_date=self.today, amount=Decimal("80000000"), allocated_amount=Decimal("50000000"),
            external_type="BANK_TRANSFER", external_reference="UNC_ADVANCE_01",
            status="CONFIRMED", created_by=self.tracking_account
        )

        res = self.client.get(
            f"/api/v1/accounts-payable/tracking/summary/{self.query}&from_date={self.today}&to_date={self.today}&supplier_id={self.supplier_b.id}"
        )
        self.assertEqual(res.status_code, 200)
        row = res.data["results"][0]

        # Trả trước cuối kỳ = 30.000.000 (80M - 50M)
        self.assertEqual(Decimal(row["closing_advance_balance"]), Decimal("30000000"))

    def test_show_zero_balance_filter_af13_br24(self):
        """
        AF-13, BR-24:
        - Ẩn số dư bằng 0: loại nhà cung cấp không có số dư nợ và không có tiền trả trước.
        - Hiển thị nếu NCC có số dư nợ = 0 nhưng còn khoản trả trước (BR-24).
        """
        # NCC C không phát sinh gì -> mặc định bị ẩn
        res_default = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}")
        self.assertEqual(res_default.status_code, 200)
        self.assertFalse(any(r["supplier_id"] == str(self.supplier_c.id) for r in res_default.data["results"]))

        # Bật show_zero_balance=true -> hiển thị NCC C
        res_zero = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}&show_zero_balance=true")
        self.assertEqual(res_zero.status_code, 200)
        self.assertTrue(any(r["supplier_id"] == str(self.supplier_c.id) for r in res_zero.data["results"]))

    def test_aging_buckets_br37_to_42(self):
        """
        BR-37..42: Phân loại chính xác 5 nhóm tuổi nợ (NOT_DUE, 1_30, 31_60, 61_90, OVER_90).
        """
        # 1. Chưa đến hạn: hạn sau 5 ngày
        self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today, self.today + timedelta(days=5))
        # 2. Quá hạn 10 ngày (1_30)
        self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today - timedelta(days=20), self.today - timedelta(days=10))
        # 3. Quá hạn 45 ngày (31_60)
        self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today - timedelta(days=60), self.today - timedelta(days=45))
        # 4. Quá hạn 75 ngày (61_90)
        self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today - timedelta(days=90), self.today - timedelta(days=75))
        # 5. Quá hạn 120 ngày (OVER_90)
        self.create_payable_with_document(self.supplier_a, Decimal("10000000"), self.today - timedelta(days=150), self.today - timedelta(days=120))

        res = self.client.get(
            f"/api/v1/accounts-payable/tracking/detail/{self.query}&supplier_id={self.supplier_a.id}&from_date={self.today - timedelta(days=200)}&to_date={self.today}"
        )
        self.assertEqual(res.status_code, 200)
        buckets = {line["aging_bucket"] for line in res.data["results"]}
        self.assertEqual(buckets, {"NOT_DUE", "1_30", "31_60", "61_90", "OVER_90"})

    def test_payable_status_prioritizes_overdue_br36_af08(self):
        """
        BR-36, AF-08: Khoản công nợ thanh toán một phần nhưng đã quá hạn thì ưu tiên hiển thị trạng thái Quá hạn.
        """
        p = self.create_payable_with_document(self.supplier_a, Decimal("50000000"), self.today - timedelta(days=30), self.today - timedelta(days=10))
        # Thanh toán 20M
        pay = Payment.objects.create(
            company=self.company, supplier=self.supplier_a, currency=self.currency,
            payment_date=self.today - timedelta(days=20), amount=Decimal("20000000"),
            allocated_amount=Decimal("20000000"), external_type="BANK_TRANSFER",
            external_reference="UNC_PARTIAL", status="CONFIRMED", created_by=self.tracking_account
        )
        PaymentAllocation.objects.create(payment=pay, payable=p, amount=Decimal("20000000"), created_by=self.tracking_account)
        p.paid_amount = Decimal("20000000")
        p.balance = Decimal("30000000")
        p.save()

        res = self.client.get(
            f"/api/v1/accounts-payable/tracking/detail/{self.query}&supplier_id={self.supplier_a.id}&from_date={self.today - timedelta(days=40)}&to_date={self.today}"
        )
        line = next(x for x in res.data["results"] if x["payable_id"] == str(p.id))
        self.assertEqual(line["status"], "OVERDUE")
        self.assertEqual(line["due_status"], "OVERDUE")

    def test_unposted_documents_excluded_br06_to_br10(self):
        """
        BR-06..10: Chứng từ Nháp, Chưa duyệt, Đã hủy không được tính vào công nợ.
        """
        # Tạo chứng từ DRAFT
        self.create_payable_with_document(self.supplier_b, Decimal("99000000"), self.today, self.today, status="DRAFT")

        res = self.client.get(f"/api/v1/accounts-payable/tracking/summary/{self.query}&supplier_id={self.supplier_b.id}")
        # NCC B không có khoản nào ghi sổ -> Không xuất hiện trong danh sách khi show_zero_balance=false
        self.assertEqual(len(res.data["results"]), 0)

    def test_excel_export_summary_and_detail_af11_af12_br05_br46(self):
        """
        AF-11, AF-12, BR-05, BR-46:
        - Xuất Excel Báo cáo Tổng hợp (export-summary)
        - Xuất Excel Báo cáo Chi tiết (export-detail)
        - Kiểm tra định dạng workbook và nội dung.
        """
        # Tạo dữ liệu mẫu
        self.create_payable_with_document(self.supplier_a, Decimal("60000000"), self.today - timedelta(days=10), self.today - timedelta(days=2))

        # 1. Xuất tổng hợp
        exp_sum = self.client.get(f"/api/v1/accounts-payable/tracking/export-summary/{self.query}")
        self.assertEqual(exp_sum.status_code, 200)
        self.assertEqual(exp_sum["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb_sum = load_workbook(BytesIO(exp_sum.content))
        self.assertIn("Tổng hợp công nợ NCC", wb_sum.sheetnames)
        ws_sum = wb_sum["Tổng hợp công nợ NCC"]
        self.assertIn("BÁO CÁO THEO DÕI CÔNG NỢ NHÀ CUNG CẤP", ws_sum.cell(2, 1).value)

        # 2. Xuất chi tiết
        exp_det = self.client.get(f"/api/v1/accounts-payable/tracking/export-detail/{self.query}&supplier_id={self.supplier_a.id}")
        self.assertEqual(exp_det.status_code, 200)
        self.assertEqual(exp_det["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb_det = load_workbook(BytesIO(exp_det.content))
        self.assertIn("Chi tiết công nợ NCC", wb_det.sheetnames)
        ws_det = wb_det["Chi tiết công nợ NCC"]
        self.assertIn("CHI TIẾT CÔNG NỢ NHÀ CUNG CẤP", ws_det.cell(2, 1).value)

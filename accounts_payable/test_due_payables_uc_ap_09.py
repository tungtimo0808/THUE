from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from iam.models import CompanyMembership, Permission, Role, RolePermission, UserRoleAssignment
from master_data.models import Currency, Party
from organization.models import Company
from .models import (Payable, PayableAdjustment, Payment, PaymentAllocation,
                     PaymentAllocationVoucher, PaymentAllocationVoucherLine,
                     PaymentRequest, PaymentRequestLine, SupplierInvoice)


class DuePayablesUcAp09Tests(TestCase):
    def setUp(self):
        self.today = timezone.localdate()
        self.currency = Currency.objects.get(code="VND")
        self.company = Company.objects.create(code="AP09", name="AP09", legal_name="AP09",
                                              tax_code="0100000009", accounting_currency=self.currency,
                                              legal_reporting_currency=self.currency)
        self.user = get_user_model().objects.create_superuser(username="ap09-user", email="ap09@example.com", password="test")
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self.supplier = Party.objects.create(company=self.company, code="NCC09", name="Supplier 09")
        self.base = f"/api/v1/accounts-payable/due-payables/?company_id={self.company.id}"

    def payable(self, number, due_offset, amount="100", supplier=None, posted_offset=-20):
        invoice = SupplierInvoice.objects.create(
            company=self.company, supplier=supplier or self.supplier, supplier_tax_code="0100000009",
            invoice_series="AA", invoice_number=number, invoice_date=self.today + timedelta(days=posted_offset),
            posting_date=self.today + timedelta(days=posted_offset), currency=self.currency,
            created_by=self.user.iam_account, status="POSTED")
        return Payable.objects.create(company=self.company, supplier=supplier or self.supplier,
                                      invoice=invoice, currency=self.currency, exchange_rate=1,
                                      original_amount=amount, balance=amount,
                                      due_date=self.today + timedelta(days=due_offset) if due_offset is not None else None)

    def test_groups_order_totals_and_exports(self):
        overdue = self.payable("OLD", -2)
        self.payable("TODAY", 0)
        self.payable("SOON", 3)
        self.payable("LATER", 10)
        self.payable("UNKNOWN", None)
        Payment.objects.create(company=self.company, supplier=self.supplier, currency=self.currency,
                               payment_date=self.today - timedelta(days=1), amount="30", allocated_amount="30",
                               status="POSTED", created_by=self.user.iam_account, number="PAY09")
        payment = Payment.objects.get(number="PAY09")
        PaymentAllocation.objects.create(payment=payment, payable=overdue, amount="30", created_by=self.user.iam_account)
        PayableAdjustment.objects.create(company=self.company, payable=overdue, adjustment_date=self.today,
                                         amount="-10", reason="Giảm giá", status="POSTED", created_by=self.user.iam_account)
        response = self.client.get(self.base)
        self.assertEqual(response.status_code, 200, response.data)
        rows = response.data["rows"]
        self.assertEqual([r["due_group"] for r in rows], ["OVERDUE", "DUE_TODAY", "DUE_SOON", "NOT_DUE", "NO_DUE_DATE"])
        self.assertEqual(rows[0]["remaining_amount"], Decimal("60"))
        self.assertEqual(rows[0]["paid_amount"], Decimal("30"))
        self.assertEqual(rows[0]["days"], 2)
        self.assertEqual(response.data["totals"]["remaining_amount"], Decimal("460"))
        self.assertEqual(response.data["totals"]["overdue"], Decimal("60"))
        detail = self.client.get(f"/api/v1/accounts-payable/due-payables/{overdue.id}/?company_id={self.company.id}")
        self.assertEqual(detail.status_code, 200, detail.data)
        self.assertEqual([e["type"] for e in detail.data["history"]], ["ORIGINAL", "PAYMENT", "ADJUSTMENT_DECREASE"])
        xlsx = self.client.get(self.base.replace("?", "export-excel/?"))
        self.assertEqual(xlsx.status_code, 200)
        sheet = load_workbook(BytesIO(xlsx.content), read_only=True).active
        self.assertEqual([sheet.cell(i, 2).value for i in range(3, 8)], ["OLD", "TODAY", "SOON", "LATER", "UNKNOWN"])
        pdf = self.client.get(self.base.replace("?", "export-pdf/?"))
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_historical_balance_and_filters(self):
        payable = self.payable("HIST", -2)
        payment = Payment.objects.create(company=self.company, supplier=self.supplier, currency=self.currency,
                                         payment_date=self.today, amount="100", allocated_amount="100",
                                         status="POSTED", created_by=self.user.iam_account, number="PAYHIST")
        PaymentAllocation.objects.create(payment=payment, payable=payable, amount="100", created_by=self.user.iam_account)
        payable.balance = 0
        payable.status = "PAID"
        payable.save()
        yesterday = self.client.get(self.base + f"&as_of_date={self.today-timedelta(days=1)}")
        self.assertEqual(yesterday.data["count"], 1)
        self.assertEqual(yesterday.data["rows"][0]["remaining_amount"], Decimal("100"))
        self.assertEqual(self.client.get(self.base).data["count"], 0)
        self.assertEqual(self.client.get(self.base + "&due_group=DUE_SOON").data["count"], 0)
        self.assertEqual(self.client.get(self.base + "&warning_days=0").status_code, 400)
        self.assertEqual(self.client.get(self.base + f"&as_of_date={self.today+timedelta(days=1)}").status_code, 400)
        self.assertEqual(self.client.get(self.base + "&currency_id=").status_code, 400)

    def test_prepare_request_uses_current_unreserved_balance(self):
        payable = self.payable("REQ", 3)
        request = PaymentRequest.objects.create(company=self.company, supplier=self.supplier, currency=self.currency,
                                                number="REQ09", requested_date=self.today, planned_payment_date=self.today,
                                                total_amount="40", status="APPROVED", created_by=self.user.iam_account)
        PaymentRequestLine.objects.create(request=request, line_no=1, payable=payable, amount="40")
        url = self.base.replace("?", "prepare-payment-request/?")
        result = self.client.post(url, {"payable_ids": [str(payable.id)]}, format="json")
        self.assertEqual(result.status_code, 200, result.data)
        self.assertEqual(result.data["total_amount"], Decimal("60"))
        self.assertEqual(self.client.post(url, {"payable_ids": []}, format="json").status_code, 400)
        payable.balance = 40
        payable.save()
        self.assertEqual(self.client.post(url, {"payable_ids": [str(payable.id)]}, format="json").status_code, 400)

    def test_reversal_filter_and_supplier_selection(self):
        first = self.payable("FIRST", -3)
        second_supplier = Party.objects.create(company=self.company, code="NCC09B", name="Supplier B")
        second = self.payable("SECOND", None, supplier=second_supplier)
        original = PayableAdjustment.objects.create(company=self.company, payable=first,
                                                    adjustment_date=self.today - timedelta(days=2), amount="-20",
                                                    reason="Trả hàng", status="POSTED", created_by=self.user.iam_account)
        PayableAdjustment.objects.create(company=self.company, payable=first, adjustment_date=self.today,
                                         amount="20", reason="Đảo trả hàng", status="POSTED", reverses=original,
                                         created_by=self.user.iam_account)
        yesterday = self.client.get(self.base + f"&as_of_date={self.today-timedelta(days=1)}&due_group=OVERDUE&supplier_ids={self.supplier.id}")
        self.assertEqual(yesterday.status_code, 200, yesterday.data)
        self.assertEqual(yesterday.data["rows"][0]["remaining_amount"], Decimal("80"))
        today = self.client.get(self.base + "&due_group=OVERDUE&payment_status=OPEN")
        self.assertEqual(today.data["rows"][0]["remaining_amount"], Decimal("100"))
        self.assertEqual(today.data["count"], 1)
        self.assertEqual(self.client.get(self.base + "&due_group=NO_DUE_DATE").data["rows"][0]["payable_id"], str(second.id))
        self.assertEqual(self.client.get(self.base + "&payment_status=PARTIALLY_PAID").data["count"], 0)
        url = self.base.replace("?", "prepare-payment-request/?")
        mixed = self.client.post(url, {"payable_ids": [str(first.id), str(second.id)]}, format="json")
        self.assertEqual(mixed.status_code, 400)
        self.assertIn("cùng một nhà cung cấp", str(mixed.data))

    def test_unposted_and_empty_result(self):
        payable = self.payable("DRAFT", 0)
        payable.invoice.status = "DRAFT"
        payable.invoice.save()
        response = self.client.get(self.base)
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(response.data["count"], 0)
        self.assertEqual(response.data["totals"]["remaining_amount"], 0)
        self.assertEqual(response.data["message"], "Không có khoản công nợ phù hợp với điều kiện tìm kiếm.")

    def test_allocated_advance_enters_balance_on_allocation_date(self):
        payable = self.payable("ADV", 5)
        payment = Payment.objects.create(company=self.company, supplier=self.supplier, currency=self.currency,
                                         payment_date=self.today - timedelta(days=5), amount="40", allocated_amount="40",
                                         status="POSTED", created_by=self.user.iam_account, number="ADV09")
        voucher = PaymentAllocationVoucher.objects.create(company=self.company, number="PB09", payment=payment,
                                                          allocation_date=self.today, description="Phân bổ tạm ứng",
                                                          total_amount="40", created_by=self.user.iam_account)
        PaymentAllocationVoucherLine.objects.create(voucher=voucher, payable=payable, line_no=1, amount="40")
        yesterday = self.client.get(self.base + f"&as_of_date={self.today-timedelta(days=1)}")
        self.assertEqual(yesterday.data["rows"][0]["remaining_amount"], Decimal("100"))
        today = self.client.get(self.base + "&payment_status=PARTIALLY_PAID")
        self.assertEqual(today.data["rows"][0]["remaining_amount"], Decimal("60"))
        self.assertEqual(today.data["rows"][0]["payable_id"], str(payable.id))

    def test_rbac_permissions_ex_ap_due_01(self):
        regular_user = get_user_model().objects.create_user(
            username="ap09-regular", email="ap09reg@example.com", password="test"
        )
        membership = CompanyMembership.objects.create(
            company=self.company, user=regular_user.iam_account, membership_status="ACTIVE",
            created_by=self.user.iam_account
        )
        client = APIClient()
        client.force_authenticate(regular_user)

        # 1. Without permission AP.PAYABLE.VIEW -> 403 Forbidden with EX-AP-DUE-01 message
        resp = client.get(self.base)
        self.assertEqual(resp.status_code, 403)
        self.assertIn("Bạn không có quyền xem công nợ phải trả.", str(resp.data))

        # 2. Grant role with AP.PAYABLE.VIEW -> 200 OK
        role = Role.objects.create(company=self.company, code="AP09_VIEWER", name="AP09 Viewer")
        perm = Permission.objects.filter(code="AP.PAYABLE.VIEW").first()
        if perm:
            RolePermission.objects.create(role=role, permission=perm)
        UserRoleAssignment.objects.create(
            company_membership=membership, role=role, assigned_by=self.user.iam_account
        )

        resp_ok = client.get(self.base)
        self.assertEqual(resp_ok.status_code, 200)

    def test_exception_flows_validation_messages(self):
        # EX-AP-DUE-02: Missing as_of_date
        r = self.client.get(self.base + "&as_of_date=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("as_of_date")), "Vui lòng nhập Ngày xác định công nợ.")

        # EX-AP-DUE-03: Missing currency_id / currency
        r = self.client.get(self.base + "&currency_id=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("currency_id")), "Vui lòng chọn Loại tiền.")

        # EX-AP-DUE-04: Missing due_group
        r = self.client.get(self.base + "&due_group=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("due_group")), "Vui lòng chọn Nhóm thời hạn.")

        # EX-AP-DUE-05: Missing warning_days
        r = self.client.get(self.base + "&warning_days=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("warning_days")), "Vui lòng nhập Số ngày cảnh báo trước hạn.")

        # EX-AP-DUE-06: Missing payment_status
        r = self.client.get(self.base + "&payment_status=")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("payment_status")), "Vui lòng chọn Trạng thái thanh toán.")

        # EX-AP-DUE-07: as_of_date > today
        future = self.today + timedelta(days=1)
        r = self.client.get(self.base + f"&as_of_date={future}")
        self.assertEqual(r.status_code, 400)
        self.assertEqual(str(r.data.get("as_of_date")), "Ngày xác định công nợ không được lớn hơn ngày hiện tại.")

        # EX-AP-DUE-08: Invalid warning_days (0, 366, invalid)
        for bad_val in (0, 366, "invalid"):
            r = self.client.get(self.base + f"&warning_days={bad_val}")
            self.assertEqual(r.status_code, 400)
            self.assertEqual(str(r.data.get("warning_days")), "Số ngày cảnh báo trước hạn phải là số nguyên từ 1 đến 365.")

        # EX-AP-DUE-09: prepare_payment_request without payable_ids
        url = self.base.replace("?", "prepare-payment-request/?")
        r = self.client.post(url, {"payable_ids": []}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Vui lòng chọn ít nhất một khoản công nợ.", str(r.data))

        # EX-AP-DUE-10: prepare_payment_request with multiple suppliers
        other_supplier = Party.objects.create(company=self.company, code="NCC_DIFF", name="Supplier Diff")
        p1 = self.payable("PAY1", 1)
        p2 = self.payable("PAY2", 1, supplier=other_supplier)
        r = self.client.post(url, {"payable_ids": [str(p1.id), str(p2.id)]}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Các khoản công nợ được chọn phải thuộc cùng một nhà cung cấp.", str(r.data))

        # EX-AP-DUE-11: prepare_payment_request when no remaining available amount
        p1.balance = Decimal("0")
        p1.save()
        r = self.client.post(url, {"payable_ids": [str(p1.id)]}, format="json")
        self.assertEqual(r.status_code, 400)
        self.assertIn("không còn số tiền có thể lập đề nghị thanh toán.", str(r.data))

    def test_vietnamese_labels_and_alternative_flows(self):
        p_overdue = self.payable("VN_OVERDUE", -3)
        p_today = self.payable("VN_TODAY", 0)
        p_soon = self.payable("VN_SOON", 2)
        p_notdue = self.payable("VN_NOTDUE", 15)
        p_nodue = self.payable("VN_NODUE", None)

        pmt = Payment.objects.create(company=self.company, supplier=self.supplier, currency=self.currency,
                                     payment_date=self.today, amount="20", allocated_amount="20",
                                     status="POSTED", created_by=self.user.iam_account, number="PAY_VN")
        PaymentAllocation.objects.create(payment=pmt, payable=p_soon, amount="20", created_by=self.user.iam_account)

        # AF-AP-DUE-01: Tất cả
        r = self.client.get(self.base + "&due_group=Tất cả")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 5)
        self.assertEqual(r.data["totals"]["total_remaining"], Decimal("480"))

        # AF-AP-DUE-02: Quá hạn
        r = self.client.get(self.base + "&due_group=Quá hạn")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_overdue.id))
        self.assertEqual(r.data["rows"][0]["days"], 3)

        # Đến hạn hôm nay
        r = self.client.get(self.base + "&due_group=Đến hạn hôm nay")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_today.id))
        self.assertEqual(r.data["rows"][0]["days"], 0)

        # AF-AP-DUE-03: Sắp đến hạn
        r = self.client.get(self.base + "&due_group=Sắp đến hạn")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_soon.id))
        self.assertEqual(r.data["rows"][0]["days"], 2)

        # Chưa đến hạn
        r = self.client.get(self.base + "&due_group=Chưa đến hạn")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_notdue.id))

        # AF-AP-DUE-04: Chưa xác định hạn
        r = self.client.get(self.base + "&due_group=Chưa xác định hạn")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_nodue.id))
        self.assertIsNone(r.data["rows"][0]["days"])

        # payment_status=Chưa thanh toán
        r = self.client.get(self.base + "&payment_status=Chưa thanh toán")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 4)

        # payment_status=Đã thanh toán một phần
        r = self.client.get(self.base + "&payment_status=Đã thanh toán một phần")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.data["count"], 1)
        self.assertEqual(r.data["rows"][0]["payable_id"], str(p_soon.id))

        # Totals aliases check
        totals = r.data["totals"]
        self.assertIn("total_remaining", totals)
        self.assertIn("total_overdue", totals)
        self.assertIn("total_due_today", totals)
        self.assertIn("total_due_soon", totals)
        self.assertIn("total_not_due", totals)
        self.assertIn("total_undetermined", totals)

    def test_sorting_order_and_business_rules(self):
        p_overdue_older = self.payable("SORT_O1", -5, posted_offset=-30)
        p_overdue_newer = self.payable("SORT_O2", -2, posted_offset=-10)
        p_overdue_same_due_doc1 = self.payable("SORT_O3A", -3, posted_offset=-25)
        p_overdue_same_due_doc2 = self.payable("SORT_O3B", -3, posted_offset=-15)

        p_today = self.payable("SORT_TODAY", 0)
        p_soon1 = self.payable("SORT_S1", 1)
        p_soon2 = self.payable("SORT_S2", 4)
        p_notdue = self.payable("SORT_ND", 10)
        p_nodue = self.payable("SORT_NONE", None)

        resp = self.client.get(self.base)
        self.assertEqual(resp.status_code, 200)
        rows = resp.data["rows"]
        num_list = [r["document_number"] for r in rows]

        expected_prefix = ["SORT_O1", "SORT_O3A", "SORT_O3B", "SORT_O2", "SORT_TODAY", "SORT_S1", "SORT_S2", "SORT_ND", "SORT_NONE"]
        self.assertEqual(num_list, expected_prefix)

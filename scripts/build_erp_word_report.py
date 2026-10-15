from __future__ import annotations

import re
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets" / "one-year-simulation"
SOURCE = DOCS / "ERP_ONE_YEAR_SIMULATION.md"
OUTPUT = DOCS / "ERP_ONE_YEAR_SIMULATION.docx"

GREEN = "606C38"
GREEN_DARK = "283618"
CREAM = "FEFAE0"
ORANGE = "DDA15E"
BLUE = "0B7894"
GRAY = "667085"


def money_label(value: float) -> str:
    return f"{value:,.0f}".replace(",", ".")


def make_charts() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})

    labels = ["Doanh thu", "Giá vốn", "Chi phí vận hành", "LNTT", "LNST"]
    values = [6000, 3000, 2100, 900, 720]
    colors = ["#606C38", "#DDA15E", "#BC6C25", "#0B7894", "#283618"]
    fig, ax = plt.subplots(figsize=(10, 5.3))
    bars = ax.bar(labels, values, color=colors, width=0.62)
    ax.set_title("Kết quả kinh doanh năm 2026", fontsize=16, fontweight="bold")
    ax.set_ylabel("Triệu VND")
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis="y", alpha=0.22)
    ax.bar_label(bars, labels=[money_label(v) for v in values], padding=4, fontweight="bold")
    fig.tight_layout()
    fig.savefig(ASSETS / "chart-profit-loss.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    months = list(range(1, 13))
    closing_cash = [3690, 3740, 3790, 3840, 3890, 3940, 3990, 4040, 4090, 4140, 4190, 4100]
    fig, ax = plt.subplots(figsize=(10, 5.2))
    ax.plot(months, closing_cash, color="#0B7894", marker="o", linewidth=2.8)
    ax.fill_between(months, closing_cash, min(closing_cash) - 100, color="#0B7894", alpha=0.12)
    ax.set_title("Tiền cuối kỳ theo tháng", fontsize=16, fontweight="bold")
    ax.set_xlabel("Tháng")
    ax.set_ylabel("Triệu VND")
    ax.set_xticks(months)
    ax.grid(alpha=0.2)
    ax.spines[["top", "right"]].set_visible(False)
    for x, y in zip(months, closing_cash):
        ax.annotate(money_label(y), (x, y), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=8)
    fig.tight_layout()
    fig.savefig(ASSETS / "chart-monthly-cash.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    asset_names = ["Tiền", "Phải thu", "Tồn kho", "TSCĐ ròng"]
    asset_values = [4100, 500, 600, 960]
    source_names = ["Phải trả", "Vốn góp", "LN chưa phân phối"]
    source_values = [440, 5000, 720]
    fig, axes = plt.subplots(1, 2, figsize=(10, 5.5))
    for ax, names, vals, title, palette in (
        (axes[0], asset_names, asset_values, "Tài sản: 6.160", ["#0B7894", "#606C38", "#DDA15E", "#BC6C25"]),
        (axes[1], source_names, source_values, "Nguồn vốn: 6.160", ["#DDA15E", "#283618", "#606C38"]),
    ):
        wedges, _, autotexts = ax.pie(
            vals,
            labels=names,
            autopct=lambda p: f"{p:.1f}%",
            startangle=90,
            colors=palette,
            textprops={"fontsize": 9},
        )
        for text in autotexts:
            text.set_fontweight("bold")
        ax.set_title(title, fontweight="bold")
    fig.suptitle("Cơ cấu bảng cân đối cuối năm", fontsize=16, fontweight="bold")
    fig.tight_layout()
    fig.savefig(ASSETS / "chart-balance-sheet.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    steps = [
        ("Khởi tạo", "Common, Organization, IAM"),
        ("Dữ liệu nền", "Master Data và kỳ kế toán"),
        ("Mua hàng", "Purchases, Inventory, AP"),
        ("Bán hàng", "Sales, Inventory, AR"),
        ("Tiền", "Cash và Bank"),
        ("Tài sản", "Fixed Assets và khấu hao"),
        ("Thuế", "Invoice, filing, obligation"),
        ("Sổ cái", "GL, đối chiếu, kết chuyển"),
        ("Kết quả", "Reporting, Audit, đóng kỳ"),
    ]
    fig, ax = plt.subplots(figsize=(13, 4.8))
    ax.set_xlim(0, 13)
    ax.set_ylim(0, 5)
    ax.axis("off")
    coords = [(0.3 + i * 1.42, 2.25) for i in range(len(steps))]
    for i, ((title, subtitle), (x, y)) in enumerate(zip(steps, coords)):
        fill = "#CDF4D3" if i in {0, 8} else "#C2E5FF" if i in {1, 7} else "#FFECBD"
        edge = "#66D575" if i in {0, 8} else "#3DADFF" if i in {1, 7} else "#FFC943"
        box = FancyBboxPatch((x, y), 1.15, 1.15, boxstyle="round,pad=0.06,rounding_size=0.08", facecolor=fill, edgecolor=edge, linewidth=1.8)
        ax.add_patch(box)
        ax.text(x + 0.575, y + 0.72, title, ha="center", va="center", fontsize=10, fontweight="bold")
        ax.text(x + 0.575, y + 0.32, subtitle, ha="center", va="center", fontsize=7.5, wrap=True)
        if i < len(steps) - 1:
            ax.annotate("", xy=(coords[i + 1][0] - 0.06, y + 0.575), xytext=(x + 1.2, y + 0.575), arrowprops={"arrowstyle": "->", "color": "#667085", "lw": 1.7})
    ax.set_title("Quy trình vận hành ERP xuyên suốt một năm", fontsize=17, fontweight="bold", pad=20)
    fig.tight_layout()
    fig.savefig(ASSETS / "diagram-erp-workflow.png", dpi=180, bbox_inches="tight")
    plt.close(fig)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_toc(document: Document) -> None:
    p = document.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "Nhấn Ctrl+A, F9 trong Word để cập nhật mục lục."
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def add_picture(document: Document, filename: str, caption: str, width: float = 6.65) -> None:
    path = ASSETS / filename
    if not path.exists():
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    caption_p = document.add_paragraph(caption)
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption_p.runs:
        caption_p.runs[0].italic = True
        caption_p.runs[0].font.size = Pt(9)
        caption_p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)


def style_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for level, size, color in ((1, 18, GREEN_DARK), (2, 15, GREEN), (3, 12, BLUE)):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)

    for sec in document.sections:
        header = sec.header.paragraphs[0]
        header.text = "SAO VIỆT ERP | MÔ PHỎNG VẬN HÀNH NĂM 2026"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRAY)
        add_page_number(sec.footer.paragraphs[0])


def add_cover(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("BÁO CÁO MÔ PHỎNG\nVẬN HÀNH ERP TRONG MỘT NĂM")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor.from_string(GREEN_DARK)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    run = p.add_run("Công ty TNHH Thương mại Sao Việt\nKỳ báo cáo: 01/01/2026 – 31/12/2026")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(GREEN)

    add_picture(document, "diagram-erp-workflow.png", "Tổng quan chuỗi vận hành từ thiết lập đến báo cáo", width=6.8)

    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    run = p.add_run("451 endpoint path · 788 HTTP operation · 56/56 backend test pass")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    document.add_page_break()

    document.add_heading("Mục lục", level=1)
    add_toc(document)
    document.add_page_break()


def parse_table(lines: list[str], start: int, document: Document) -> int:
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
        i += 1
    if len(rows) < 2:
        return start
    data = [rows[0]] + rows[2:]
    columns = max(len(row) for row in data)
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(data):
        cells = table.add_row().cells
        for col in range(columns):
            value = row[col] if col < len(row) else ""
            value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
            set_cell_text(cells[col], value, bold=row_index == 0)
            if row_index == 0:
                set_cell_shading(cells[col], GREEN)
                for run in cells[col].paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cells[col], "F5F7F2")
    document.add_paragraph()
    return i


def add_markdown_text(document: Document) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_language = ""
    code_lines: list[str] = []
    skipped_title = False

    image_after_heading = {
        "3.1. Báo cáo kết quả kinh doanh": [("chart-profit-loss.png", "Biểu đồ 1 — Doanh thu, chi phí và lợi nhuận năm 2026")],
        "3.5. Bảng cân đối cuối năm": [("chart-balance-sheet.png", "Biểu đồ 2 — Cơ cấu tài sản và nguồn vốn cuối năm")],
        "4. Sơ đồ vận hành toàn hệ thống": [("diagram-erp-workflow.png", "Sơ đồ 1 — Luồng nghiệp vụ và dữ liệu toàn hệ thống")],
        "5. Dòng tiền và lợi nhuận theo 12 tháng": [("chart-monthly-cash.png", "Biểu đồ 3 — Biến động tiền cuối kỳ trong 12 tháng")],
        "7.1. Common Services — tài khoản và không gian làm việc": [("ui-dashboard.png", "Ảnh 1 — Dashboard và không gian làm việc của ứng dụng hiện tại")],
        "8.1. Purchases": [("ui-purchase-order.png", "Ảnh 3 — Màn hình chi tiết đơn mua hàng")],
        "10.1. Cash Management": [("ui-cash-entry-form.png", "Ảnh 2 — Form nhập phiếu thu 500 triệu trên frontend hiện tại")],
        "14. General Ledger — hội tụ toàn bộ số liệu": [("ui-ledger-process.png", "Ảnh 4 — Màn hình quy trình Tổng hợp/GL")],
        "15. Reporting và Audit — kết quả cuối cùng": [("ui-report-center.png", "Ảnh 5 — Trung tâm báo cáo")],
    }

    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if not in_code:
                in_code = True
                code_language = stripped[3:].strip()
                code_lines = []
            else:
                if code_language != "mermaid":
                    p = document.add_paragraph()
                    set_cell_shading_like_paragraph(p, "F4F5F2")
                    run = p.add_run("\n".join(code_lines))
                    run.font.name = "Consolas"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                    run.font.size = Pt(8.5)
                in_code = False
                code_language = ""
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            i = parse_table(lines, i, document)
            continue
        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 3)
            text = heading.group(2).strip()
            if level == 1 and not skipped_title:
                skipped_title = True
                i += 1
                continue
            document.add_heading(text, level=level)
            for filename, caption in image_after_heading.get(text, []):
                add_picture(document, filename, caption)
            i += 1
            continue
        if stripped == "---":
            p = document.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D9D9D9")
            borders.append(bottom)
            pPr.append(borders)
            i += 1
            continue
        if re.match(r"^[-*]\s+", stripped):
            text = re.sub(r"^[-*]\s+", "", stripped)
            p = document.add_paragraph(style="List Bullet")
            add_inline_markdown(p, text)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", stripped):
            text = re.sub(r"^\d+\.\s+", "", stripped)
            p = document.add_paragraph(style="List Number")
            add_inline_markdown(p, text)
            i += 1
            continue
        p = document.add_paragraph()
        add_inline_markdown(p, stripped.lstrip("> "))
        if stripped.startswith(">"):
            p.style = document.styles["Intense Quote"]
        i += 1


def set_cell_shading_like_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_inline_markdown(paragraph, text: str) -> None:
    text = re.sub(r"\[([^]]+)]\(([^)]+)\)", r"\1", text)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(BLUE)
        else:
            paragraph.add_run(part)


def build() -> None:
    make_charts()
    document = Document()
    style_document(document)
    add_cover(document)
    add_markdown_text(document)
    document.core_properties.title = "Mô phỏng vận hành ERP trong một năm"
    document.core_properties.subject = "Quy trình, API, số liệu tài chính, biểu đồ và form nhập liệu"
    document.core_properties.author = "Dự án ERP Kế toán – Thuế"
    document.core_properties.keywords = "ERP, API, kế toán, thuế, GL, mô phỏng, 2026"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

"""Create executable Postman artifacts and a request-level workbook sheet.

This module is called by export_common_services_test_cases.py.  It deliberately
keeps one API request per execution row so testers never have to infer method,
URL, parameters, headers, or JSON body from a high-level scenario.
"""

from __future__ import annotations

import json
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


POSTMAN_HEADERS = [
    "Request ID", "Test Case ID", "UC ID", "Tên test case", "Method", "Request URL",
    "Tài khoản/Token", "Path variables", "Query parameters", "Headers", "Request Body",
    "Chuẩn bị dữ liệu / Cách lấy biến", "Expected HTTP", "Expected response & side effects",
    "Postman Tests", "Status", "Actual Result", "Tester", "Ngày chạy", "Defect ID",
    "Evidence/Link", "Kiểm tra nhập liệu",
]

DEFAULT_HEADERS = "Authorization: Bearer {{access_token}}\nAccept: application/json\nContent-Type: application/json"
PUBLIC_HEADERS = "Accept: application/json\nContent-Type: application/json"


TEMPLATES = {
    "register": ("POST", "{{base_url}}/api/v1/common/auth/register/", PUBLIC_HEADERS,
                 "", '{\n  "email": "{{register_email}}",\n  "phone": "{{register_phone}}",\n  "display_name": "Người dùng QA",\n  "password": "{{register_password}}",\n  "password_confirm": "{{register_password_confirm}}",\n  "terms_accepted": true,\n  "terms_version": "2026-01",\n  "privacy_version": "2026-01"\n}'),
    "verify": ("POST", "{{base_url}}/api/v1/common/auth/verify-registration/", PUBLIC_HEADERS,
               "", '{\n  "challenge": "{{registration_challenge}}",\n  "code": "{{otp_code}}"\n}'),
    "resend": ("POST", "{{base_url}}/api/v1/common/auth/resend-verification/", PUBLIC_HEADERS,
               "", '{\n  "challenge": "{{registration_challenge}}"\n}'),
    "change_registration_email": ("POST", "{{base_url}}/api/v1/common/auth/change-registration-email/", PUBLIC_HEADERS,
                                  "", '{\n  "challenge": "{{registration_challenge}}",\n  "email": "{{changed_email}}"\n}'),
    "company_search": ("GET", "{{base_url}}/api/v1/common/companies/search/?q={{company_search_q}}", DEFAULT_HEADERS,
                       "q={{company_search_q}}", ""),
    "join_create": ("POST", "{{base_url}}/api/v1/common/join-requests/", DEFAULT_HEADERS,
                    "", '{\n  "company": "{{c1_id}}",\n  "message": "Xin tham gia doanh nghiệp",\n  "invitation_code": "{{invitation_code}}"\n}'),
    "join_list": ("GET", "{{base_url}}/api/v1/common/join-requests/", DEFAULT_HEADERS, "", ""),
    "join_cancel": ("POST", "{{base_url}}/api/v1/common/join-requests/{{join_request_id}}/cancel/", DEFAULT_HEADERS,
                    "", "{}"),
    "join_admin_list": ("GET", "{{base_url}}/api/v1/common/company-join-requests/?company_id={{c1_id}}&status=PENDING", DEFAULT_HEADERS,
                        "company_id={{c1_id}}\nstatus=PENDING", ""),
    "join_decide": ("POST", "{{base_url}}/api/v1/common/company-join-requests/{{join_request_id}}/decide/?company_id={{c1_id}}", DEFAULT_HEADERS,
                    "company_id={{c1_id}}", '{\n  "company": "{{c1_id}}",\n  "decision": "APPROVED",\n  "reason": "Đồng ý"\n}'),
    "login": ("POST", "{{base_url}}/api/v1/common/auth/login/", PUBLIC_HEADERS,
              "", '{\n  "login": "{{login}}",\n  "password": "{{password}}"\n}'),
    "refresh": ("POST", "{{base_url}}/api/token/refresh/", PUBLIC_HEADERS,
                "", '{\n  "refresh": "{{refresh_token}}"\n}'),
    "logout": ("POST", "{{base_url}}/api/v1/common/auth/logout/", DEFAULT_HEADERS,
               "", '{\n  "refresh": "{{refresh_token}}"\n}'),
    "forgot": ("POST", "{{base_url}}/api/v1/common/auth/forgot-password/", PUBLIC_HEADERS,
               "", '{\n  "identifier": "{{login}}"\n}'),
    "reset": ("POST", "{{base_url}}/api/v1/common/auth/reset-password/", PUBLIC_HEADERS,
              "", '{\n  "challenge": "{{password_reset_challenge}}",\n  "code": "{{otp_code}}",\n  "new_password": "{{new_password}}",\n  "new_password_confirm": "{{new_password_confirm}}"\n}'),
    "change_password": ("POST", "{{base_url}}/api/v1/common/auth/change-password/", DEFAULT_HEADERS,
                        "", '{\n  "current_password": "{{password}}",\n  "new_password": "{{new_password}}",\n  "new_password_confirm": "{{new_password_confirm}}"\n}'),
    "profile_get": ("GET", "{{base_url}}/api/v1/common/profile/", DEFAULT_HEADERS, "", ""),
    "profile_patch": ("PATCH", "{{base_url}}/api/v1/common/profile/", DEFAULT_HEADERS,
                      "", '{\n  "display_name": "Nguyễn Văn QA",\n  "phone": "0901234567",\n  "first_name": "Văn",\n  "last_name": "Nguyễn",\n  "version": "{{profile_version}}"\n}'),
    "confirm_email": ("POST", "{{base_url}}/api/v1/common/profile/confirm-email/", DEFAULT_HEADERS,
                      "", '{\n  "challenge": "{{email_change_challenge}}",\n  "code": "{{otp_code}}"\n}'),
    "scopes": ("GET", "{{base_url}}/api/v1/common/scopes/", DEFAULT_HEADERS, "", ""),
    "switch_scope": ("POST", "{{base_url}}/api/v1/common/companies/switch/", DEFAULT_HEADERS,
                     "", '{\n  "company": "{{c1_id}}",\n  "branch": "{{b1_id}}",\n  "all_branches": false,\n  "confirm_discard": true\n}'),
    "notifications": ("GET", "{{base_url}}/api/v1/common/notifications/?company_id={{c1_id}}&page=1&page_size=20", DEFAULT_HEADERS,
                      "company_id={{c1_id}}\npage=1\npage_size=20", ""),
    "notification_detail": ("GET", "{{base_url}}/api/v1/common/notifications/{{notification_id}}/", DEFAULT_HEADERS, "", ""),
    "notification_read": ("POST", "{{base_url}}/api/v1/common/notifications/{{notification_id}}/read/", DEFAULT_HEADERS, "", "{}"),
    "notification_unread": ("POST", "{{base_url}}/api/v1/common/notifications/{{notification_id}}/unread/", DEFAULT_HEADERS, "", "{}"),
    "notification_read_all": ("POST", "{{base_url}}/api/v1/common/notifications/read-all/?company_id={{c1_id}}", DEFAULT_HEADERS,
                              "company_id={{c1_id}}", "{}"),
    "notification_count": ("GET", "{{base_url}}/api/v1/common/notifications/unread-count/?company_id={{c1_id}}", DEFAULT_HEADERS,
                           "company_id={{c1_id}}", ""),
    "notification_archive": ("POST", "{{base_url}}/api/v1/common/notifications/{{notification_id}}/archive/", DEFAULT_HEADERS, "", "{}"),
    "notification_open": ("POST", "{{base_url}}/api/v1/common/notifications/{{notification_id}}/open/", DEFAULT_HEADERS, "", "{}"),
    "search": ("GET", "{{base_url}}/api/v1/common/search/?company={{c1_id}}&branch={{b1_id}}&q={{search_q}}&page=1&page_size=20", DEFAULT_HEADERS,
               "company={{c1_id}}\nbranch={{b1_id}}\nq={{search_q}}\npage=1\npage_size=20", ""),
    "search_open": ("POST", "{{base_url}}/api/v1/common/search/open/", DEFAULT_HEADERS,
                    "", '{\n  "company": "{{c1_id}}",\n  "branch": "{{b1_id}}",\n  "entity_type": "{{search_entity_type}}",\n  "entity_id": "{{search_entity_id}}"\n}'),
    "work_list": ("GET", "{{base_url}}/api/v1/common/work-items/?company_id={{c1_id}}&branch_id={{b1_id}}", DEFAULT_HEADERS,
                  "company_id={{c1_id}}\nbranch_id={{b1_id}}", ""),
    "work_detail": ("GET", "{{base_url}}/api/v1/common/work-items/{{work_item_id}}/?company_id={{c1_id}}&branch_id={{b1_id}}", DEFAULT_HEADERS,
                    "company_id={{c1_id}}\nbranch_id={{b1_id}}", ""),
    "work_count": ("GET", "{{base_url}}/api/v1/common/work-items/count/?company_id={{c1_id}}&branch_id={{b1_id}}", DEFAULT_HEADERS,
                   "company_id={{c1_id}}\nbranch_id={{b1_id}}", ""),
    "work_open": ("POST", "{{base_url}}/api/v1/common/work-items/{{work_item_id}}/open/?company_id={{c1_id}}&branch_id={{b1_id}}", DEFAULT_HEADERS,
                  "company_id={{c1_id}}\nbranch_id={{b1_id}}", "{}"),
    "work_history": ("GET", "{{base_url}}/api/v1/common/work-items/{{work_item_id}}/history/?company_id={{c1_id}}&branch_id={{b1_id}}", DEFAULT_HEADERS,
                     "company_id={{c1_id}}\nbranch_id={{b1_id}}", ""),
    "ai_conversation": ("POST", "{{base_url}}/api/v1/common/ai-conversations/", DEFAULT_HEADERS,
                        "", '{\n  "company": "{{c1_id}}",\n  "branch": "{{b1_id}}",\n  "title": "Phân tích tháng"\n}'),
    "ai_query": ("POST", "{{base_url}}/api/v1/common/ai/query/", DEFAULT_HEADERS,
                 "", '{\n  "company": "{{c1_id}}",\n  "conversation": "{{ai_conversation_id}}",\n  "question": "{{ai_question}}",\n  "request_type": "ANALYSIS",\n  "request_id": "{{$guid}}",\n  "context": {"screen": "dashboard"}\n}'),
    "ai_confirm": ("POST", "{{base_url}}/api/v1/common/ai/{{ai_interaction_id}}/confirm/?company_id={{c1_id}}", DEFAULT_HEADERS,
                   "company_id={{c1_id}}", "{}"),
    "ai_feedback": ("POST", "{{base_url}}/api/v1/common/ai/{{ai_interaction_id}}/feedback/?company_id={{c1_id}}", DEFAULT_HEADERS,
                    "company_id={{c1_id}}", '{\n  "helpful": true,\n  "comment": "Hữu ích"\n}'),
    "ai_source": ("POST", "{{base_url}}/api/v1/common/ai/{{ai_interaction_id}}/source/{{source_index}}/?company_id={{c1_id}}", DEFAULT_HEADERS,
                  "company_id={{c1_id}}", "{}"),
}


def expected_codes(text: str, default: int) -> list[int]:
    codes = [int(value) for value in re.findall(r"(?<!\d)(20[0-9]|40[0-9]|429|503)(?!\d)", text)]
    return sorted(set(codes)) or [default]


def choose_templates(uc: str, case_id: str, scenario: str) -> list[str]:
    number = int(case_id[-2:])
    if uc == "UC-COM-01":
        return ["verify"] if number in {2, 3, 4, 5, 13} else ["resend"] if number in {6, 7} else ["change_registration_email"] if number in {11, 12} else ["register"]
    if uc == "UC-COM-02":
        return ["company_search"] if number <= 2 else ["join_cancel"] if number in {8, 9} else ["join_decide"] if number >= 10 else ["join_create"]
    if uc == "UC-COM-03":
        return ["refresh"] if number == 12 else ["login"]
    if uc == "UC-COM-04":
        return ["logout"]
    if uc == "UC-COM-05":
        return ["forgot"] if number in {1, 2, 3, 10} else ["forgot", "reset"] if number == 11 else ["reset"]
    if uc == "UC-COM-06":
        return ["change_password"]
    if uc == "UC-COM-07":
        return ["profile_get"] if number == 1 else ["confirm_email"] if number in {11, 12} else ["profile_patch"]
    if uc == "UC-COM-08":
        return ["scopes"] if number in {1, 11} else ["switch_scope"]
    if uc == "UC-COM-09":
        mapping = {3: "notification_detail", 4: "notification_read", 5: "notification_read_all", 6: "notification_count", 7: "notification_archive", 8: "notification_open", 9: "notification_open", 10: "notification_open"}
        return [mapping.get(number, "notifications")]
    if uc == "UC-COM-10":
        return ["search_open"] if number in {8, 9} else ["search"]
    if uc == "UC-COM-11":
        if number == 3:
            return ["work_list", "work_detail", "work_open", "work_history"]
        if number == 7:
            return ["work_count"]
        if number in {8, 9}:
            return ["work_open"]
        if number == 10:
            return ["work_history"]
        return ["work_list"]
    if uc == "UC-COM-12":
        if number in {4, 5}:
            return ["ai_conversation"]
        if number == 12:
            return ["ai_confirm"]
        if number == 13:
            return ["ai_feedback"]
        if number == 14:
            return ["ai_source"]
        return ["ai_query"]
    raise ValueError(uc)


def actor_for(uc: str, case_id: str, scenario: str) -> str:
    value = scenario.lower()
    if uc in {"UC-COM-01", "UC-COM-03", "UC-COM-05"}:
        return "Không cần token (public endpoint)"
    if "user khác" in value or "c2" in value or "ngoài scope" in value:
        return "{{token_u4}} – U4 chỉ có C1/B1"
    if uc == "UC-COM-02" and int(case_id[-2:]) >= 10:
        return "{{token_u7}} – admin C1"
    return "{{access_token}} – actor theo TEST_DATA"


def setup_text(uc: str, case_id: str, scenario: str, template: str) -> str:
    base = [
        "1. Chạy request đăng nhập phù hợp và lưu access token vào Environment.",
        "2. Điền UUID từ dữ liệu seed/staging; không tự đặt UUID giả nếu case yêu cầu object tồn tại.",
        "3. Đối chiếu actor, membership, Permission và Data Scope trong sheet TEST_DATA.",
    ]
    if uc in {"UC-COM-01", "UC-COM-05"}:
        base[0] = "1. Lấy challenge từ response trước đó và OTP từ email sandbox; không ghi OTP vào defect/evidence."
    if case_id == "COM11-03":
        if template == "work_list":
            base += ["4. Đổi branch_id trên URL thành {{b2_id}} trong khi U4 chỉ được B1; expected 403."]
        else:
            base += ["4. Đặt {{work_item_id}}={{c2_work_item_id}}, giữ company_id={{c1_id}}; expected 404 và không lộ dữ liệu C2."]
    if "đồng thời" in scenario.lower() or "hai admin" in scenario.lower():
        base += ["4. Dùng Postman Runner/Newman hoặc hai cửa sổ gửi request gần như đồng thời."]
    if "5 lần" in scenario.lower():
        base += ["4. Gửi request sai đủ 5 lần; lần tiếp theo dùng dữ liệu đúng để kiểm tra challenge/session đã bị vô hiệu."]
    return "\n".join(base)


def request_specs(cases: dict[str, list[dict[str, str]]]) -> list[dict[str, object]]:
    specs: list[dict[str, object]] = []
    for uc, rows in cases.items():
        for case in rows:
            templates = choose_templates(uc, case["id"], case["scenario"])
            for index, template_name in enumerate(templates):
                method, url, headers, query, body = TEMPLATES[template_name]
                suffix = chr(65 + index) if len(templates) > 1 else ""
                request_id = f"{case['id']}{suffix}"
                default = 201 if method == "POST" and template_name in {"register", "join_create", "ai_conversation", "ai_query"} else 204 if template_name in {"logout", "notification_archive"} else 200
                codes = expected_codes(case["expected"], default)
                # Exact cross-tenant expectations for the motivating example.
                if case["id"] == "COM11-03":
                    codes = [403] if template_name == "work_list" else [404]
                    if template_name == "work_list":
                        url = url.replace("{{b1_id}}", "{{b2_id}}")
                        query = query.replace("{{b1_id}}", "{{b2_id}}")
                    else:
                        url = url.replace("{{work_item_id}}", "{{c2_work_item_id}}")
                tests = (
                    f"pm.test('HTTP thuộc {codes}', () => pm.expect({json.dumps(codes)}).to.include(pm.response.code));\n"
                    "pm.test('Response không lộ secret', () => { const t=pm.response.text().toLowerCase(); "
                    "['password','refresh_token','access_token','otp_code'].forEach(k => pm.expect(t).not.to.include(k)); });"
                )
                specs.append({
                    "request_id": request_id, "case_id": case["id"], "uc": uc,
                    "title": case["scenario"], "method": method, "url": url,
                    "actor": actor_for(uc, case["id"], case["scenario"]),
                    "path_vars": "Các biến {{...}} trong path phải lấy từ response/seed; xem cột chuẩn bị dữ liệu.",
                    "query": query or "Không có", "headers": headers, "body": body or "Không có body",
                    "setup": setup_text(uc, case["id"], case["scenario"], template_name),
                    "codes": codes, "expected": case["expected"], "tests": tests,
                })
    return specs


def add_postman_sheet(workbook_path: Path, specs: list[dict[str, object]]) -> None:
    wb = load_workbook(workbook_path)
    if "POSTMAN_CASES" in wb.sheetnames:
        del wb["POSTMAN_CASES"]
    ws = wb.create_sheet("POSTMAN_CASES", 4)
    last_col = len(POSTMAN_HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, "POSTMAN – REQUEST THỰC THI CHI TIẾT")
    ws.cell(1, 1).font = Font(size=18, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor="17365D")
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, "Mỗi dòng là một HTTP request. Chỉ sửa các ô màu vàng từ Status đến Evidence; request/input lấy nguyên văn từ các cột bên trái.")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor="D9EAF7")
    ws.cell(2, 1).alignment = Alignment(wrap_text=True)
    ws.append(POSTMAN_HEADERS)
    for spec in specs:
        ws.append((
            spec["request_id"], spec["case_id"], spec["uc"], spec["title"], spec["method"], spec["url"],
            spec["actor"], spec["path_vars"], spec["query"], spec["headers"], spec["body"], spec["setup"],
            ", ".join(str(code) for code in spec["codes"]), spec["expected"], spec["tests"],
            "Not Run", "", "", "", "", "", "",
        ))
    thin = Side(style="thin", color="B7C9D6")
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    status_validation = DataValidation(type="list", formula1='"Not Run,Pass,Fail,Blocked,N/A"')
    ws.add_data_validation(status_validation)
    status_validation.add(f"P4:P{ws.max_row}")
    for row in range(4, ws.max_row + 1):
        # Input-validation formula: execution fields become mandatory after Not Run.
        formula = (
            f'=IF(P{row}="Not Run","Chưa chạy",IF(R{row}="","Thiếu Tester",'
            f'IF(S{row}="","Thiếu Ngày chạy",IF(Q{row}="","Thiếu Actual Result",'
            f'IF(AND(P{row}="Fail",T{row}=""),"Thiếu Defect ID",'
            f'IF(AND(OR(P{row}="Fail",P{row}="Blocked"),U{row}=""),"Thiếu Evidence","Hợp lệ"))))))'
        )
        ws.cell(row, 22, formula)
        for col in range(1, last_col + 1):
            cell = ws.cell(row, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.protection = Protection(locked=col not in {16, 17, 18, 19, 20, 21})
        for col in {16, 17, 18, 19, 20, 21}:
            ws.cell(row, col).fill = PatternFill("solid", fgColor="FFF2CC")
        ws.row_dimensions[row].height = 125
    ws.auto_filter.ref = f"A3:{get_column_letter(last_col)}{ws.max_row}"
    ws.freeze_panes = "F4"
    ws.sheet_view.showGridLines = False
    ws.conditional_formatting.add(f"V4:V{ws.max_row}", FormulaRule(formula=['$V4="Hợp lệ"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"V4:V{ws.max_row}", FormulaRule(formula=['AND($V4<>"Hợp lệ",$V4<>"Chưa chạy")'], fill=PatternFill("solid", fgColor="FFC7CE")))
    widths = [16, 15, 14, 42, 10, 62, 28, 36, 38, 42, 58, 62, 16, 60, 72, 13, 48, 20, 14, 18, 30, 20]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.protection.sheet = True
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = False
    wb.save(workbook_path)


def make_postman_url(raw_url: str) -> dict[str, object]:
    path_str, query_str = raw_url.split("?", 1) if "?" in raw_url else (raw_url, "")
    has_trailing_slash = path_str.endswith("/")
    rel_path = path_str
    if rel_path.startswith("{{base_url}}"):
        rel_path = rel_path[len("{{base_url}}"):]
    rel_path = rel_path.strip("/")
    path_segments = [seg for seg in rel_path.split("/") if seg]
    if has_trailing_slash:
        path_segments.append("")
    query_list = []
    if query_str:
        for item in query_str.split("&"):
            if not item:
                continue
            k, v = item.split("=", 1) if "=" in item else (item, "")
            query_list.append({"key": k, "value": v})
    url_obj: dict[str, object] = {
        "raw": raw_url,
        "host": ["{{base_url}}"],
        "path": path_segments,
    }
    if query_list:
        url_obj["query"] = query_list
    return url_obj


def postman_request(spec: dict[str, object]) -> dict[str, object]:
    header_rows = []
    auth = None
    for line in str(spec["headers"]).splitlines():
        key, value = line.split(":", 1)
        key, value = key.strip(), value.strip()
        if key.lower() == "authorization":
            token = "token_u4" if "token_u4" in str(spec["actor"]) else "token_u7" if "token_u7" in str(spec["actor"]) else "access_token"
            auth = {"type": "bearer", "bearer": [{"key": "token", "value": f"{{{{{token}}}}}", "type": "string"}]}
        else:
            header_rows.append({"key": key, "value": value, "type": "text"})
    request: dict[str, object] = {
        "method": spec["method"], "header": header_rows,
        "url": make_postman_url(str(spec["url"])),
        "description": f"Test case: {spec['case_id']}\n\n{spec['setup']}\n\nExpected: {spec['expected']}",
    }
    if auth:
        request["auth"] = auth
    if spec["body"] != "Không có body":
        request["body"] = {"mode": "raw", "raw": spec["body"], "options": {"raw": {"language": "json"}}}
    test_lines = str(spec["tests"]).splitlines()
    raw_url = str(spec["url"])
    capture_rules = [
        ("/auth/register/", "registration_challenge", "challenge"),
        ("/auth/verify-registration/", "access_token", "access"),
        ("/auth/verify-registration/", "refresh_token", "refresh"),
        ("/auth/login/", "access_token", "access"),
        ("/auth/login/", "refresh_token", "refresh"),
        ("/auth/forgot-password/", "password_reset_challenge", "challenge"),
        ("/join-requests/", "join_request_id", "id"),
        ("/ai-conversations/", "ai_conversation_id", "id"),
        ("/ai/query/", "ai_interaction_id", "id"),
        ("/profile/", "profile_version", "updated_at"),
    ]
    for marker, variable, field in capture_rules:
        if marker in raw_url:
            test_lines += [
                f"if (pm.response.code >= 200 && pm.response.code < 300) {{ const j=pm.response.json(); if (j.{field}) pm.environment.set('{variable}', j.{field}); }}"
            ]
    return {
        "name": f"{spec['request_id']} – {spec['title']}",
        "event": [{"listen": "test", "script": {"type": "text/javascript", "exec": test_lines}}],
        "request": request,
    }


def export_postman(output_dir: Path, specs: list[dict[str, object]], uc_names: dict[str, str]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    folders = []
    for uc, name in uc_names.items():
        folders.append({"name": f"{uc} – {name}", "item": [postman_request(spec) for spec in specs if spec["uc"] == uc]})
    collection = {
        "info": {
            "_postman_id": str(uuid.uuid4()),
            "name": "ERP Common Services – UC-COM-01 to UC-COM-12",
            "description": "Generated execution collection. Fill environment variables before running; never store production secrets.",
            "schema": "https://schema.getpostman.com/json/collection/v2.1.0/collection.json",
        },
        "item": folders,
    }
    variables = {
        "base_url": "http://127.0.0.1:8000", "access_token": "", "refresh_token": "", "token_u4": "", "token_u7": "",
        "login": "new@example.com", "password": "StrongPass!2026", "new_password": "NewStrongPass!2026", "new_password_confirm": "NewStrongPass!2026",
        "register_email": "qa.user.001@example.com", "register_phone": "0901000001",
        "register_password": "StrongPass!2026", "register_password_confirm": "StrongPass!2026", "changed_email": "qa.changed@example.com",
        "registration_challenge": "", "password_reset_challenge": "", "email_change_challenge": "", "otp_code": "",
        "c1_id": "", "c2_id": "", "b1_id": "", "b2_id": "", "join_request_id": "", "invitation_code": "",
        "notification_id": "", "work_item_id": "", "c2_work_item_id": "", "profile_version": "",
        "company_search_q": "AC", "search_q": "anh duong", "search_entity_type": "master_data.Party", "search_entity_id": "",
        "ai_conversation_id": "", "ai_interaction_id": "", "source_index": "0", "ai_question": "Phân tích tình hình hiện tại",
    }
    environment = {
        "id": str(uuid.uuid4()), "name": "ERP Common Services – Local/Staging Template",
        "values": [{"key": key, "value": value, "enabled": True, "type": "secret" if "token" in key or "password" in key or key == "otp_code" else "default"} for key, value in variables.items()],
        "_postman_variable_scope": "environment",
        "_postman_exported_using": "Codex generator",
    }
    (output_dir / "COMMON_SERVICES.postman_collection.json").write_text(json.dumps(collection, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "COMMON_SERVICES.local.postman_environment.json").write_text(json.dumps(environment, ensure_ascii=False, indent=2), encoding="utf-8")


def enrich_workbook_and_export_postman(workbook_path: Path, cases: dict[str, list[dict[str, str]]], uc_names: dict[str, str]) -> int:
    specs = request_specs(cases)
    add_postman_sheet(workbook_path, specs)
    export_postman(workbook_path.parent / "postman", specs, uc_names)
    return len(specs)
